"""
Sube las rutas del dia a Google My Maps (lo que antes se hacia a mano) y
lleva el registro de los links.

My Maps no tiene API: un "robot" (Playwright) maneja un Chrome aparte, con su
propio perfil en data/chrome-mymaps, donde la cuenta de Opaustro inicio
sesion a mano (el robot nunca escribe contrasenas). Por cada grupo de
camiones crea un mapa:

  titulo       "1-MAPA DE ENTREGAS LUNES 12 DE OCTUBRE"  (fecha de ENTREGA)
  descripcion  "Transportistas F150-F151-F153"
  capas        una por archivo MAPA (max. 10 por mapa), "Agrupar sitios por:
               vendedor" con un color distinto por vendedor, etiquetas con
               la Razon, y todas desmarcadas al abrir
  compartir    visible para cualquiera con el enlace, sin busqueda publica

Grupos: F1xx+F2xx, F3xx, FO (y cualquier otro prefijo aparte); si un grupo
pasa de 10 capas se parte en mapas parejos.

Las capas se suben desde una copia de cada MAPA con las ubicaciones
corregidas (las mismas que usa el mapa web). Los mapas con fecha de entrega
de hace mas de 30 dias se mandan a la papelera de Drive (solo en My Maps: el
mapa web guarda su historial completo).
"""
import json
import math
import os
import re
import shutil
import subprocess
import tempfile
import time
import urllib.request
from datetime import date, datetime, timedelta
from pathlib import Path
from urllib.parse import parse_qs, urlparse

CARPETA = Path(__file__).resolve().parent
PERFIL_CHROME = CARPETA / "data" / "chrome-mymaps"
ARCHIVO_REGISTRO = CARPETA / "data" / "mymaps_registro.json"
CAPTURA_ERROR = CARPETA / "data" / "mymaps_error.png"
# La cuenta de Google de los mapas se anota aqui (data/ nunca va a GitHub):
# {"cuenta": "...@gmail.com"}. Si no esta, se toma la primera con la que se
# inicie sesion en el Chrome del robot.
ARCHIVO_CONFIG = CARPETA / "data" / "mymaps_config.json"
PUERTO = 9333
INICIO = "https://www.google.com/maps/d/?hl=es"
MAX_CAPAS = 10
DIAS_HISTORIAL = 30
MINUTOS_PARA_INICIAR_SESION = 10

DIAS_LARGOS = ["LUNES", "MARTES", "MIERCOLES", "JUEVES", "VIERNES", "SABADO", "DOMINGO"]
MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO",
         "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]

# Mismos nombres que escribe generar_rutas.escribir_excel.
RE_ARCHIVO = re.compile(r"MAPA ([A-Z]{1,3})(\d{1,4})( CARGUE (\d+) - Liq [0-9A-Za-z]+)?\.xlsx")

# My Maps solo deja elegir entre sus 30 colores fijos. Este orden alterna
# tono y claridad, para que dos vendedores seguidos nunca se parezcan.
PALETA = [
    (255, 82, 82), (2, 136, 209), (15, 157, 88), (245, 124, 0), (156, 39, 176),
    (255, 234, 0), (0, 0, 0), (0, 151, 167), (194, 24, 91), (124, 179, 66),
    (26, 35, 126), (121, 85, 72), (189, 189, 189), (165, 39, 20), (103, 58, 183),
    (175, 180, 43), (0, 96, 100), (249, 168, 37), (136, 14, 79), (57, 73, 171),
    (9, 113, 56), (230, 81, 0), (117, 117, 117), (1, 87, 155), (129, 119, 23),
    (85, 139, 47), (255, 214, 0), (251, 192, 45), (78, 52, 46), (66, 66, 66),
]


class ErrorMyMaps(Exception):
    """Error con un mensaje ya pensado para mostrarse tal cual."""


class Cancelado(ErrorMyMaps):
    pass


class _SinCuadro(ErrorMyMaps):
    """My Maps no mostro un cuadro esperado al importar."""


class MapaNoExiste(ErrorMyMaps):
    """Al abrir el mapa, My Maps mando al inicio: esta en la papelera o se borro."""


class _ImportacionVacia(Exception):
    """My Maps cerro la importacion sin cargar nada (se reintenta una vez)."""


# ---------------------------------------------------------------------------
# Fechas, grupos y textos
# ---------------------------------------------------------------------------

def fecha_larga(d: date) -> str:
    return f"{DIAS_LARGOS[d.weekday()]} {d.day:02d} DE {MESES[d.month - 1]}"


def entrega_sugerida(fecha_liq: date) -> date:
    """El siguiente dia habil: viernes (o viernes + sabado) se entrega el lunes."""
    d = fecha_liq + timedelta(days=1)
    while d.weekday() >= 5:
        d += timedelta(days=1)
    return d


def opciones_de_entrega(fecha_liq: date, cuantas=8):
    """Dias que se pueden elegir como entrega (sin domingos), por si el
    siguiente dia habil no se trabaja."""
    dias, d = [], fecha_liq
    while len(dias) < cuantas:
        d += timedelta(days=1)
        if d.weekday() != 6:
            dias.append(d)
    return dias


def fecha_de_carpeta(carpeta: Path):
    """BASE\\MAPAS\\2026-10-09 VIERNES -> date(2026, 10, 9)."""
    try:
        return datetime.strptime(carpeta.name[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


# Mapas hechos a mano antes del robot: "1-MAPA DE ENTREGAS LUNES 12 DE OCTUBRE".
# El dia de la semana a veces esta mal escrito, asi que la fecha sale del
# numero y el mes; el año, de la fecha que muestra My Maps en su lista.
RE_TITULO = re.compile(r"^(\d+)\s*-\s*MAPA DE ENTREGAS\s+[A-Z]+\s+(\d{1,2})\s+DE\s+([A-Z]+)$")
MESES_CORTOS = {"ENE": 1, "FEB": 2, "MAR": 3, "ABR": 4, "MAY": 5, "JUN": 6, "JUL": 7, "AGO": 8, "SEP": 9,
                "SET": 9, "OCT": 10, "NOV": 11, "DIC": 12, "JAN": 1, "APR": 4, "AUG": 8, "DEC": 12}


def _sin_tildes(texto: str) -> str:
    import unicodedata
    s = unicodedata.normalize("NFKD", str(texto or ""))
    return " ".join("".join(c for c in s if not unicodedata.combining(c)).upper().split())


def fecha_de_lista(texto: str):
    """Fecha que muestra la lista de My Maps: '13 oct 2026' o 'Tue Oct 13 2026'."""
    t = _sin_tildes(texto).replace(".", "")
    m = re.search(r"(\d{1,2}) ([A-Z]{3})[A-Z]* (\d{4})", t) or re.search(r"([A-Z]{3})[A-Z]* (\d{1,2}) (\d{4})", t)
    if not m:
        return None
    a, b, anio = m.groups()
    dia, mes = (a, b) if a.isdigit() else (b, a)
    try:
        return date(int(anio), MESES_CORTOS[mes[:3]], int(dia))
    except (KeyError, ValueError):
        return None


def fecha_de_titulo(titulo: str, referencia: date):
    """(numero, fecha de entrega) de un titulo de mapa de entregas, o None.
    El año es el que deja la fecha mas cerca de `referencia`."""
    m = RE_TITULO.match(_sin_tildes(titulo))
    if not m:
        return None
    numero, dia, mes = int(m.group(1)), int(m.group(2)), m.group(3)
    if mes not in MESES:
        return None
    opciones = []
    for anio in (referencia.year - 1, referencia.year, referencia.year + 1):
        try:
            opciones.append(date(anio, MESES.index(mes) + 1, dia))
        except ValueError:
            pass
    if not opciones:
        return None
    return numero, min(opciones, key=lambda d: abs((d - referencia).days))


def registrar_manuales(registro: dict, encontrados, hoy: date) -> int:
    """Anota en el registro los mapas de entregas hechos a mano (para que
    salgan en la lista de links y se limpien a los 30 dias como los demas).
    encontrados: [{mid, titulo, fecha_lista}]. Devuelve cuantos se anotaron."""
    conocidos = {m.get("mid") for m in registro["mapas"]}
    nuevos = 0
    for e in encontrados:
        if not e.get("mid") or e["mid"] in conocidos:
            continue
        datos = fecha_de_titulo(e.get("titulo", ""), e.get("fecha_lista") or hoy)
        if not datos:
            continue
        numero, entrega = datos
        if entrega > (e.get("fecha_lista") or hoy) + timedelta(days=7):
            # La fecha de la lista (ultimo cambio) nunca es meses ANTES de la
            # entrega: un mapa a mano no se hace con tanta anticipacion, asi
            # que la entrega es del año anterior.
            try:
                entrega = entrega.replace(year=entrega.year - 1)
            except ValueError:              # 29 de febrero
                entrega = entrega.replace(year=entrega.year - 1, day=28)
        registro["mapas"].append({
            "mid": e["mid"], "titulo": e["titulo"], "descripcion": "", "transportistas": [], "capas": [],
            "numero": numero, "fecha_entrega": entrega.isoformat(), "origen": "manual",
            "url": f"https://www.google.com/maps/d/edit?mid={e['mid']}&usp=sharing",
            "creado": (e.get("fecha_lista") or hoy).isoformat(), "estado": "activo"})
        conocidos.add(e["mid"])
        nuevos += 1
    return nuevos


ARCHIVO_PUBLICADO = "_publicado.json"     # lo escribe generar_rutas.anotar_publicado


def archivos_mapa(carpeta: Path):
    """Los MAPA de la carpeta del dia. Si la ruta se publico, solo los que se
    publicaron (asi My Maps sube lo mismo que el mapa web, aunque despues se
    haya sacado algun Excel suelto con "Solo generar Excel")."""
    todos = sorted(p for p in carpeta.iterdir() if p.is_file() and RE_ARCHIVO.fullmatch(p.name))
    try:
        publicados = set(json.loads((carpeta / ARCHIVO_PUBLICADO).read_text(encoding="utf-8"))["archivos"])
    except (OSError, ValueError, KeyError, TypeError):
        return todos
    return [p for p in todos if p.name in publicados]


def cambiados_desde_publicar(carpeta: Path):
    """MAPA publicados que ya no son iguales a lo que se publico (o que
    faltan). Lista vacia si no hay anotacion de lo publicado."""
    from generar_rutas import huella_archivo
    try:
        huellas = json.loads((carpeta / ARCHIVO_PUBLICADO).read_text(encoding="utf-8")).get("huellas") or {}
    except (OSError, ValueError, AttributeError):
        return []
    cambiados = []
    for nombre, huella in sorted(huellas.items()):
        ruta = carpeta / nombre
        try:
            if huella_archivo(ruta) != huella:
                cambiados.append(nombre)
        except OSError:
            cambiados.append(nombre)
    return cambiados


def texto_cambiados(cambiados) -> str:
    return (f"Estos archivos cambiaron después de publicar la ruta (por ejemplo con \"Solo generar Excel\"): "
            f"{', '.join(cambiados)}.\nPara que My Maps tenga lo mismo que el mapa web, vuelve a \"Generar y "
            "publicar\" en \"Actualizar ruta\".")


def datos_archivo(ruta: Path):
    m = RE_ARCHIVO.fullmatch(ruta.name)
    letras, numero, cargue = m.group(1), m.group(2), m.group(4)
    return {"camion": letras + numero, "letras": letras, "numero": int(numero),
            "cargue": int(cargue) if cargue else 0}


def grupo_de(letras: str, numero: str):
    """(orden, nombre) del mapa al que va un camion."""
    if letras == "F" and numero[0] in "12":
        return (1, "F1-F2")
    if letras == "F":
        return (int(numero[0]), f"F{numero[0]}xx")
    if letras == "FO":
        return (50, "FO")
    return (90, letras)


def partir(camiones, maximo=MAX_CAPAS):
    """Reparte camiones (cada uno con 1+ capas) en mapas de hasta `maximo`
    capas, parejos, sin separar los cargues de un mismo camion."""
    total = sum(len(capas) for _, capas in camiones)
    if total <= maximo:
        return [camiones]
    objetivo = math.ceil(total / math.ceil(total / maximo))
    partes, actual, n = [], [], 0
    for camion in camiones:
        capas = len(camion[1])
        if actual and (n + capas > objetivo or n + capas > maximo):
            partes.append(actual)
            actual, n = [], 0
        actual.append(camion)
        n += capas
    if actual:
        partes.append(actual)
    return partes


def armar_mapas(archivos, fecha_entrega: date):
    """Lista de mapas a crear: numero, titulo, descripcion, transportistas y
    archivos (capas), en el orden F1-F2, F3, ..., FO, otros."""
    camiones = {}
    for ruta in archivos:
        d = datos_archivo(ruta)
        clave = (grupo_de(d["letras"], str(d["numero"])), d["letras"], d["numero"])
        camiones.setdefault(clave, (d["camion"], []))[1].append((d["cargue"], ruta))
    grupos = {}
    for clave in sorted(camiones):
        camion, capas = camiones[clave]
        grupos.setdefault(clave[0], []).append((camion, [r for _, r in sorted(capas)]))
    mapas = []
    for grupo in sorted(grupos):
        for parte in partir(grupos[grupo]):
            n = len(mapas) + 1
            ids = [camion for camion, _ in parte]
            mapas.append({
                "numero": n,
                "grupo": grupo[1],
                "titulo": f"{n}-MAPA DE ENTREGAS {fecha_larga(fecha_entrega)}",
                "descripcion": "Transportistas " + "-".join(ids),
                "transportistas": ids,
                "archivos": [r for _, capas in parte for r in capas],
            })
    return mapas


def texto_para_compartir(mapas) -> str:
    """Para pegar en WhatsApp: titulo, transportistas y link de cada mapa."""
    bloques = ["\n".join(x for x in (m.get("titulo"), m.get("descripcion"), m["url"]) if x)
               for m in mapas if m.get("url")]
    return "\n\n".join(bloques)


# ---------------------------------------------------------------------------
# Registro de mapas creados (data/mymaps_registro.json)
# ---------------------------------------------------------------------------

def leer_registro(ruta: Path = ARCHIVO_REGISTRO) -> dict:
    try:
        datos = json.loads(ruta.read_text(encoding="utf-8"))
        if isinstance(datos, dict) and isinstance(datos.get("mapas"), list):
            return datos
    except (OSError, ValueError):
        pass
    return {"mapas": []}


def guardar_registro(registro: dict, ruta: Path = ARCHIVO_REGISTRO) -> None:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    temporal = ruta.with_suffix(".tmp")
    temporal.write_text(json.dumps(registro, ensure_ascii=False, indent=1), encoding="utf-8")
    os.replace(temporal, ruta)


def activos(registro: dict, fecha_entrega: str = None):
    return [m for m in registro["mapas"] if m.get("estado") == "activo"
            and (fecha_entrega is None or m.get("fecha_entrega") == fecha_entrega)]


def por_mandar_a_papelera(registro: dict, hoy: date, reemplazar_fecha: str = None):
    """Activos con entrega de hace mas de 30 dias, los de la fecha que se
    reemplaza, y los que quedaron a medias en un intento anterior (o que no
    se pudieron mandar al reemplazarlos)."""
    limite = (hoy - timedelta(days=DIAS_HISTORIAL)).isoformat()
    salida = []
    for m in registro["mapas"]:
        if not m.get("mid"):
            continue
        if m.get("estado") in ("creando", "fallido", "reemplazado"):
            salida.append(m)
        elif m.get("estado") == "activo" and (m.get("fecha_entrega", "") < limite
                                              or m.get("fecha_entrega") == reemplazar_fecha):
            salida.append(m)
    return salida


def limpiar_registro(registro: dict, hoy: date) -> None:
    """Lo que ya esta en la papelera hace mas de 60 dias deja de anotarse."""
    limite = (hoy - timedelta(days=2 * DIAS_HISTORIAL)).isoformat()
    registro["mapas"] = [m for m in registro["mapas"]
                         if not (m.get("estado") in ("papelera", "sin_crear")
                                 and m.get("fecha_entrega", "") < limite)]


# ---------------------------------------------------------------------------
# Copias con ubicaciones corregidas
# ---------------------------------------------------------------------------

COLUMNAS_UBICACION = ("CodigoCliente", "longitud_x", "latitud_y")


def _encabezado(ws, ruta: Path):
    posiciones = {str(c.value or "").strip(): i for i, c in enumerate(next(ws.iter_rows(max_row=1)))}
    if not all(c in posiciones for c in COLUMNAS_UBICACION):
        raise ErrorMyMaps(f"\"{ruta.name}\" no tiene las columnas CodigoCliente, longitud_x y latitud_y.")
    return posiciones


def camion_de_archivo(ruta: Path) -> str:
    """'MAPA F152 CARGUE 1 - Liq 1557.xlsx' -> 'F152 CARGUE 1 - Liq 1557'."""
    return ruta.name[len("MAPA "):-len(".xlsx")]


def revisar_georef(archivos, overrides: dict):
    """Clientes de los MAPA sin georreferencia valida, con la MISMA regla del
    mapa web (generar_rutas.clientes_sin_georef). Lista vacia = todo bien."""
    import openpyxl
    from generar_rutas import clientes_sin_georef

    camiones = []
    for ruta in archivos:
        wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            pos = _encabezado(ws, ruta)
            filas = []
            for valores in ws.iter_rows(min_row=2, values_only=True):
                if not valores or all(v in (None, "") for v in valores):
                    continue
                filas.append({c: (valores[i] if i < len(valores) else None) for c, i in pos.items()})
        finally:
            wb.close()
        chofer = next((f.get("Chofer") for f in filas if f.get("Chofer")), "")
        camiones.append({"id": camion_de_archivo(ruta), "chofer": str(chofer or ""), "filas": filas})
    return clientes_sin_georef(camiones, overrides)


def preparar_copias(archivos, destino: Path, overrides: dict):
    """Copia cada MAPA (mismo nombre: es el nombre de la capa) con la
    ubicacion corregida de cada cliente (la misma que usa el mapa web).
    Antes hay que revisar con revisar_georef: aqui un cliente sin ubicacion
    valida es un error (un 0,0 aleja el mapa, y una sola celda vacia hace que
    My Maps ya no reconozca la columna como latitud).
    Devuelve (copias, cuantas corregidas)."""
    import openpyxl
    from generar_rutas import a_coordenada, codigo_texto, correccion_valida, problema_de_ubicacion

    copias, corregidas = [], 0
    for ruta in archivos:
        wb = openpyxl.load_workbook(ruta)
        try:
            ws = wb.worksheets[0]
            pos = _encabezado(ws, ruta)
            i_cod, i_lng, i_lat = (pos[c] for c in COLUMNAS_UBICACION)
            vacias = []
            for fila in ws.iter_rows(min_row=2):
                if all(c.value in (None, "") for c in fila):
                    vacias.append(fila[0].row)
                    continue
                ajuste = correccion_valida(overrides, fila[i_cod].value)
                if ajuste:
                    fila[i_lat].value, fila[i_lng].value = ajuste
                    corregidas += 1
                elif problema_de_ubicacion(fila[i_lat].value, fila[i_lng].value):
                    raise ErrorMyMaps(f"El cliente {codigo_texto(fila[i_cod].value)} de {ruta.name} no tiene "
                                      "georreferencia válida.")
                else:
                    fila[i_lat].value = a_coordenada(fila[i_lat].value)
                    fila[i_lng].value = a_coordenada(fila[i_lng].value)
            for numero in reversed(vacias):
                ws.delete_rows(numero)
            copia = destino / ruta.name
            wb.save(copia)
        finally:
            wb.close()
        copias.append(copia)
    return copias, corregidas


def _celda_huella(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float):
        return f"{valor:.6f}"
    return str(valor).strip()


def huella_capa(ruta: Path) -> str:
    """Resumen de TODO lo que se subio en una capa (ubicaciones corregidas,
    vendedor, razon, clientes...): si despues cambia (una correccion de
    ubicacion o la ruta vuelta a publicar con otros datos), esa capa de My
    Maps esta vieja y se puede sustituir sin cambiar el link."""
    import hashlib
    import openpyxl

    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        _encabezado(ws, ruta)
        filas, encabezado = [], None
        for valores in ws.iter_rows(values_only=True):
            if not valores or all(v in (None, "") for v in valores):
                continue
            texto_fila = "|".join(_celda_huella(v) for v in valores).rstrip("|")
            if encabezado is None:
                encabezado = texto_fila
            else:
                filas.append(texto_fila)
    finally:
        wb.close()
    return hashlib.sha1("\n".join([encabezado or ""] + sorted(filas)).encode("utf-8")).hexdigest()


def _grupo_de_capas(nombres):
    """Grupos de mapa (F1-F2, F3xx, FO...) de esos archivos MAPA."""
    grupos = set()
    for n in nombres:
        m = RE_ARCHIVO.fullmatch(n)
        if m:
            grupos.add(grupo_de(m.group(1), m.group(2))[1])
    return grupos


def actualizables(registro: dict, hoy: date):
    """Mapas que subio el robot, de hoy en adelante y con huellas: los unicos
    que se pueden actualizar sin cambiar el link."""
    hoy_iso = hoy.isoformat()
    return [m for m in registro["mapas"]
            if m.get("estado") == "activo" and m.get("origen") != "manual" and m.get("fecha_entrega", "") >= hoy_iso
            and m.get("huellas") and m.get("carpeta")]


def mapas_por_actualizar(registro: dict, overrides: dict, hoy: date, destino: Path):
    """Mapas de My Maps de hoy en adelante con cambios desde que se subieron
    (ubicaciones corregidas en el Excel de control, en la ventana o en el
    mapa, o la ruta vuelta a publicar con otros datos en los mismos
    camiones). Deja en `destino` las copias nuevas de las capas afectadas.
    Devuelve (planes, avisos): planes = [{mapa, capas: {nombre: copia}}]."""
    planes, avisos = [], []
    revisadas = set()
    for m in actualizables(registro, hoy):
        huellas, carpeta = m["huellas"], Path(m["carpeta"])
        if not carpeta.is_dir():
            avisos.append(f"{m['titulo']}: no encuentro su carpeta ({carpeta.name}); no se puede revisar.")
            continue
        if cambiados_desde_publicar(carpeta):
            avisos.append(f"{m['titulo']}: los archivos MAPA cambiaron después de publicar la ruta; vuelve a "
                          "«Generar y publicar» y después actualiza My Maps (el link no cambia).")
            continue
        # Solo se sustituyen capas: si cambiaron los camiones o los cargues
        # (otros archivos), hay que volver a subir la ruta.
        actuales = {p.name for p in archivos_mapa(carpeta)}
        propias = set(m.get("capas") or huellas)
        # Capas de los mapas de esa misma entrega (un grupo de mas de 10
        # camiones se parte en 2 mapas): no cuentan como camiones nuevos.
        de_la_entrega = {c for o in activos(registro, m.get("fecha_entrega"))
                         if o.get("carpeta") == m.get("carpeta") for c in (o.get("capas") or [])}
        nuevos = actuales - propias - de_la_entrega
        if not propias <= actuales or _grupo_de_capas(nuevos) & _grupo_de_capas(propias):
            avisos.append(f"{m['titulo']}: cambiaron los camiones o cargues de la ruta después de subirla; "
                          "vuelve a subirla a My Maps.")
            continue
        clave = (m["carpeta"], m.get("fecha_entrega"))
        if clave not in revisadas:
            # Un camion de un grupo que no tiene ningun mapa (p. ej. un FO
            # nuevo): ningun mapa lo trae, hay que volver a subir la ruta.
            revisadas.add(clave)
            grupos_con_mapa = _grupo_de_capas(de_la_entrega)
            sin_mapa = sorted(n for n in actuales - de_la_entrega if not _grupo_de_capas([n]) & grupos_con_mapa)
            if sin_mapa:
                avisos.append(f"Entrega del {fecha_larga(date.fromisoformat(m['fecha_entrega']))}: "
                              f"{', '.join(sin_mapa)} no está en ningún mapa de My Maps; vuelve a subir la ruta.")
        archivos = [carpeta / n for n in huellas]
        faltan = revisar_georef(archivos, overrides)
        if faltan:
            avisos.append(f"{m['titulo']}: {len(faltan)} clientes sin ubicación; corrígelos primero.")
            continue
        sub = destino / m["mid"]
        sub.mkdir(parents=True, exist_ok=True)
        copias, _ = preparar_copias(archivos, sub, overrides)
        capas = {c.name: c for c in copias if huella_capa(c) != huellas.get(c.name)}
        if capas:
            planes.append({"mapa": m, "capas": capas})
    return planes, avisos


# ---------------------------------------------------------------------------
# Robot
# ---------------------------------------------------------------------------

def cuenta_esperada():
    try:
        cuenta = json.loads(ARCHIVO_CONFIG.read_text(encoding="utf-8")).get("cuenta")
        return cuenta.strip().lower() if isinstance(cuenta, str) and "@" in cuenta else None
    except (OSError, ValueError, AttributeError):
        return None


def anotar_cuenta(cuenta: str) -> None:
    ARCHIVO_CONFIG.parent.mkdir(parents=True, exist_ok=True)
    ARCHIVO_CONFIG.write_text(json.dumps({"cuenta": cuenta}, indent=1), encoding="utf-8")


def ruta_chrome() -> Path:
    for variable in ("PROGRAMFILES", "PROGRAMFILES(X86)", "LOCALAPPDATA"):
        base = os.environ.get(variable)
        if base:
            exe = Path(base) / "Google" / "Chrome" / "Application" / "chrome.exe"
            if exe.exists():
                return exe
    raise ErrorMyMaps("No encontré Google Chrome instalado en esta computadora.")


def chrome_escuchando() -> bool:
    try:
        with urllib.request.urlopen(f"http://127.0.0.1:{PUERTO}/json/version", timeout=2):
            return True
    except Exception:
        return False


def esperar(condicion, segundos, paso=0.4) -> bool:
    fin = time.time() + segundos
    while True:
        try:
            if condicion():
                return True
        except Exception:
            pass
        if time.time() >= fin:
            return False
        time.sleep(paso)


class Robot:
    """Chrome del robot abierto con su perfil, y los pasos en My Maps.

    Se abre Chrome normal (con un puerto de depuracion local) y Playwright
    se conecta a el: asi es el mismo Chrome donde se inicio sesion a mano."""

    def __init__(self, log, cancelar=None):
        self.log = log
        self.cancelar = cancelar
        self.pw = self.navegador = self.pg = None
        self.mid_en_curso = None
        # Se llama con el mid apenas existe el mapa (para anotarlo aunque
        # despues se corte todo).
        self.al_crear = None

    def __enter__(self):
        from playwright.sync_api import sync_playwright

        if not chrome_escuchando():
            PERFIL_CHROME.mkdir(parents=True, exist_ok=True)
            subprocess.Popen([str(ruta_chrome()), f"--user-data-dir={PERFIL_CHROME}",
                              f"--remote-debugging-port={PUERTO}", "--no-first-run",
                              "--no-default-browser-check", "--lang=es", "--window-size=1300,880", "about:blank"])
            if not esperar(chrome_escuchando, 30, 0.5):
                raise ErrorMyMaps("El Chrome del robot no respondió.\nSi hay una ventana de Chrome del robot "
                                  "abierta (la de My Maps), ciérrala y vuelve a intentar.")
        self.pw = sync_playwright().start()
        try:
            self.navegador = self.pw.chromium.connect_over_cdp(f"http://127.0.0.1:{PUERTO}")
            # El robot trabaja en una pestana propia y nunca toca las demas
            # (puede haber otra cosa abierta en ese Chrome).
            self.pg = self.navegador.contexts[0].new_page()
            self.pg.set_default_timeout(30000)
        except Exception:
            self.pw.stop()
            raise
        return self

    def __exit__(self, *exc):
        # Se cierra la pestana del robot, y el Chrome del robot entero si no
        # le queda nada mas abierto (asi no queda abierto con la sesion).
        try:
            self.pg.close()
            vacias = ("about:blank", "chrome://newtab/", "chrome://new-tab-page/", "")
            if all(p.url in vacias for c in self.navegador.contexts for p in c.pages):
                self.navegador.new_browser_cdp_session().send("Browser.close")
            else:
                self.log("(El Chrome del robot queda abierto porque tiene otras pestañas.)")
        except Exception:
            pass
        try:
            self.pw.stop()
        except Exception:
            pass
        return False

    def revisar_cancelado(self):
        if self.cancelar is not None and self.cancelar.is_set():
            raise Cancelado("Se canceló la subida a My Maps.")

    # -- sesion -----------------------------------------------------------
    def cuenta_visible(self):
        """Correo de la cuenta con sesion iniciada en My Maps (None si no hay)."""
        etiqueta = self.pg.locator('[aria-label^="Cuenta de Google"]').first
        if not etiqueta.count():
            return None
        m = re.search(r"\(([^)]+@[^)]+)\)", etiqueta.get_attribute("aria-label") or "")
        return m.group(1).strip().lower() if m else None

    def asegurar_sesion(self):
        pg = self.pg
        pg.goto(INICIO, wait_until="domcontentloaded")
        esperar(lambda: self.cuenta_visible() or "accounts.google" in pg.url, 15)
        cuenta, esperada = self.cuenta_visible(), cuenta_esperada()
        if cuenta is None:
            self.log(f"Inicia sesión con {esperada or 'la cuenta de Google de los mapas'} en la ventana de Chrome "
                     f"que se abrió (te espero hasta {MINUTOS_PARA_INICIAR_SESION} minutos)...")
            pg.bring_to_front()
            fin = time.time() + MINUTOS_PARA_INICIAR_SESION * 60
            while time.time() < fin:
                self.revisar_cancelado()
                time.sleep(3)
                if "google.com/maps/d" in pg.url:
                    cuenta = self.cuenta_visible()
                    if cuenta:
                        break
                elif "accounts.google" not in pg.url and "google.com" in pg.url:
                    # Termino el inicio de sesion en otra pagina de Google.
                    pg.goto(INICIO, wait_until="domcontentloaded")
            if not cuenta:
                raise ErrorMyMaps("No se inició sesión en Google a tiempo. Vuelve a intentar cuando puedas "
                                  "iniciar sesión (puede pedir el código de verificación).")
        if esperada is None:
            anotar_cuenta(cuenta)
            self.log(f"Los mapas se crean en la cuenta {cuenta} (anotada en {ARCHIVO_CONFIG.name}).")
        elif cuenta != esperada:
            raise ErrorMyMaps(f"El Chrome del robot tiene abierta la cuenta {cuenta}, pero los mapas van en "
                              f"{esperada}.\nCierra esa sesión en el Chrome del robot e inicia con la correcta.")

    # -- crear un mapa ------------------------------------------------------
    def crear_mapa(self, titulo, descripcion, archivos):
        """Crea el mapa completo y devuelve {mid, url}. Si algo falla, el
        mapa a medias queda en self.mid_en_curso (para mandarlo a la papelera)."""
        pg = self.pg
        self.mid_en_curso = None
        pg.goto(INICIO, wait_until="domcontentloaded")
        boton = pg.locator('div[role="button"][aria-label="Crear un nuevo mapa"]').filter(visible=True).first
        boton.wait_for(state="visible", timeout=30000)
        boton.click()

        def en_editor():
            if "maps/d/edit" in pg.url and "mid=" in pg.url:
                return True
            # Aviso de Drive que a veces sale antes de crear el mapa.
            for texto in ("CREATE", "CREAR", "Crear"):
                aviso = pg.get_by_text(texto, exact=True).filter(visible=True)
                if aviso.count():
                    aviso.first.click()
                    break
            return False

        if not esperar(en_editor, 45, 1):
            raise ErrorMyMaps("My Maps no abrió el mapa nuevo.")
        self.mid_en_curso = parse_qs(urlparse(pg.url).query)["mid"][0]
        if self.al_crear:
            self.al_crear(self.mid_en_curso)
        if not esperar(lambda: self.capas(), 30):
            raise ErrorMyMaps("El mapa nuevo no terminó de abrir.")

        # Titulo y descripcion
        pg.get_by_text("Mapa sin nombre", exact=True).filter(visible=True).first.click()
        campo = pg.locator('input[type="text"][placeholder="Mapa sin nombre"]').filter(visible=True).first
        campo.wait_for(state="visible")
        campo.fill(titulo)
        pg.locator("textarea").filter(visible=True).first.fill(descripcion)
        pg.get_by_role("button", name="Guardar").filter(visible=True).first.click()
        if not esperar(lambda: pg.get_by_text(titulo, exact=True).filter(visible=True).count(), 15):
            raise ErrorMyMaps("No se pudo poner el título del mapa.")

        for i, archivo in enumerate(archivos):
            self.revisar_cancelado()
            self.log(f"   capa {i + 1}/{len(archivos)}: {archivo.name}")
            if i == 0:
                ly = self.capas()[0]
            else:
                antes = self.capas()
                pg.get_by_text("Añadir capa", exact=True).filter(visible=True).first.click()
                if not esperar(lambda: len(self.capas()) > len(antes), 20):
                    raise ErrorMyMaps("No se pudo añadir una capa nueva.")
                ly = [c for c in self.capas() if c not in antes][0]
            self.importar(ly, archivo)
            self.dar_estilo(ly)
            nombre = pg.locator(f"#{ly}-layer-header").inner_text().strip()
            if nombre != archivo.name:
                self.log(f"   Aviso: la capa quedó con el nombre \"{nombre}\".")

        self.revisar_cancelado()
        for ly in self.capas():
            casilla = pg.locator(f'#{ly}-layer-header [role="checkbox"]').first
            if casilla.get_attribute("aria-checked") == "true":
                casilla.click()
                esperar(lambda: casilla.get_attribute("aria-checked") == "false", 5)
        url = self.compartir()
        self.esperar_guardado()
        self.verificar(len(archivos))
        return {"mid": self.mid_en_curso, "url": url}

    def capas(self):
        return self.pg.evaluate("""() => Array.from(document.querySelectorAll('[id$="-layer-items-container"]'))
            .map(e => e.id.split('-')[0])""")

    IFRAME_SELECTOR = 'iframe[src*="docs.google.com/picker"]'

    def selector_de_archivos(self):
        """El cuadro "Selecciona un archivo que quieras importar" (un iframe
        de Google Picker) visible y ya con la opcion "Subir". Se busca por el
        iframe de la pagina y no por la lista de frames: esa lista a veces no
        trae la direccion del cuadro (pasa al volver a importar) y no se
        encontraba aunque estuviera en pantalla."""
        iframes = self.pg.locator(self.IFRAME_SELECTOR)
        try:
            for i in range(iframes.count()):
                elemento = iframes.nth(i)
                if not elemento.is_visible():
                    continue
                cuadro = elemento.content_frame
                if cuadro.get_by_text("Subir", exact=True).filter(visible=True).count():
                    return cuadro
        except Exception:
            pass
        return None

    def abrir_selector(self, ly):
        pg = self.pg
        if self.selector_de_archivos():
            return              # sigue abierto de un intento anterior: se usa ese
        enlace = pg.locator(f"#{ly}-layerview-import-link")
        if not enlace.count():
            enlace = pg.get_by_text("Importar", exact=True).filter(visible=True).last
        enlace.click()
        # La primera vez el cuadro tarda bastante (carga la lista de Drive).
        # Solo si ni siquiera existe su iframe, el clic se perdio y se repite.
        if esperar(self.selector_de_archivos, 90, 1):
            return
        if not pg.locator(self.IFRAME_SELECTOR).filter(visible=True).count():
            try:
                enlace.click(timeout=5000)
            except Exception:
                pass
            if esperar(self.selector_de_archivos, 60, 1):
                return
        raise ErrorMyMaps("No se abrió la ventana para importar el archivo.")

    def cuadro_visible(self, texto, segundos):
        cuadros = self.pg.locator('div[role="dialog"]').filter(has_text=texto).filter(visible=True)
        if not esperar(cuadros.count, segundos, 0.5):
            raise _SinCuadro(f"No apareció el cuadro \"{texto}...\" al importar.")
        return cuadros.first

    def importar(self, ly, archivo: Path):
        # A veces My Maps no termina una importacion y no dice nada (la capa
        # queda vacia, sin ningun cuadro abierto): se repite una vez.
        for intento in (1, 2):
            self.abrir_selector(ly)
            try:
                # Cuando funciona, los clientes aparecen en segundos: si a los 40 s
                # no hay nada (y no hay cuadros abiertos), se reintenta.
                self.subir_archivo(ly, archivo, esperar_segundos=40 if intento == 1 else 180)
                return
            except _ImportacionVacia as e:
                if intento == 2:
                    raise ErrorMyMaps(str(e))
                self.log(f"   My Maps no cargó {archivo.name}; se vuelve a importar.")

    def subir_archivo(self, ly, archivo: Path, esperar_segundos=180):
        pg = self.pg
        selector = self.selector_de_archivos()
        selector.get_by_text("Subir", exact=True).filter(visible=True).first.click()
        entrada = selector.locator("input[type=file]")
        entrada.wait_for(state="attached", timeout=20000)
        entrada.set_input_files(str(archivo))

        sitios = pg.locator(f"#{ly}-layer-items-container div.un1lmc-pbTTYe-ibnC6b-KDwhZb")
        try:
            self.marcar_columnas(archivo)
        except _SinCuadro as e:
            # A veces My Maps no muestra el cuadro de columnas (sin aviso): si
            # la capa sigue vacia, se vuelve a importar como cuando no carga.
            if sitios.count():
                raise
            raise _ImportacionVacia(f"No se pudo importar {archivo.name}: {e}")
        self.elegir_titulo()

        if not esperar(lambda: sitios.count() > 0 and not pg.locator('div[role="dialog"]').filter(visible=True).count(),
                       esperar_segundos, 1):
            abiertos = pg.locator('div[role="dialog"]').filter(visible=True)
            if abiertos.count():
                raise ErrorMyMaps(f"No se pudo importar {archivo.name}: {abiertos.first.inner_text()[:300]}")
            time.sleep(5)
            if sitios.count():
                return                  # termino justo ahora: no se reintenta (duplicaria los puntos)
            barra = pg.locator("#butterbar-wrap")
            aviso_google = barra.inner_text().strip() if barra.count() else ""
            raise _ImportacionVacia(f"No se pudo importar {archivo.name}: no aparecieron los clientes"
                                    + (f" (My Maps dijo: {aviso_google[:200]})" if aviso_google else ""))
        aviso = pg.locator(f"#{ly}-layerview-error-infobox-text")
        if aviso.count() and aviso.is_visible():
            self.log(f"   Aviso de My Maps: {aviso.inner_text().strip()}")

    def marcar_columnas(self, archivo: Path):
        """Cuadro "Selecciona columnas...": latitud_y = Latitud, longitud_x = Longitud.
        My Maps deja ocultos los cuadros de importaciones anteriores: hay que
        usar siempre el que esta visible."""
        dlg = self.cuadro_visible("Selecciona columnas", 90)
        anterior = None                 # burbuja usada para la columna anterior
        for columna, valor in (("latitud_y", "latitude"), ("longitud_x", "longitude")):
            id_etiqueta = dlg.locator('span[id^="upload-checkbox-label-"]', has_text=re.compile(rf"^{columna}$")) \
                .first.get_attribute("id")
            numero = id_etiqueta.rsplit("-", 1)[1]
            casilla = dlg.locator(f'[role="checkbox"][aria-labelledby="{id_etiqueta}"]')
            casilla.scroll_into_view_if_needed()
            if casilla.get_attribute("aria-checked") != "true":
                casilla.click()
            anterior = self.elegir_tipo_de_columna(valor, dlg, numero, anterior)
            if not anterior:
                raise ErrorMyMaps(f"No se pudo marcar la columna {columna} al importar {archivo.name}.")
        continuar = dlg.get_by_role("button", name="Continuar")
        if not esperar(lambda: continuar.is_enabled(), 5, 0.3):
            raise ErrorMyMaps(f"My Maps no aceptó las columnas de ubicación al importar {archivo.name}.")
        time.sleep(0.5)
        continuar.click()
        time.sleep(1)

    def elegir_tipo_de_columna(self, valor, dlg, numero, anterior):
        """Latitud/Longitud se elige en la burbuja de esa columna, que sale al
        marcarla (cada columna tiene la suya). A veces la burbuja de la columna
        anterior todavia no se cerro: se usa la que NO es la anterior y se
        confirma que la columna quedo "(Latitud)"/"(Longitud)". Devuelve la
        burbuja usada (o None si no se pudo)."""
        pg = self.pg
        anotada = dlg.locator(f"#upload-checkbox-mapped-{numero}")
        esperado = "(latitud)" if valor == "latitude" else "(longitud)"
        for _ in range(4):
            burbuja = None
            fin = time.time() + 6
            while time.time() < fin and not burbuja:
                visibles = pg.evaluate("""() => Array.from(document.querySelectorAll('[role="alertdialog"]'))
                    .filter(b => { const r = b.getBoundingClientRect(); return r.width > 0 && r.height > 0
                                   && b.querySelector('[role="radio"][data-value]'); })
                    .map(b => b.getAttribute('aria-describedby') || '')""")
                nuevas = [b for b in visibles if b and b != anterior]
                if len(nuevas) == 1:
                    burbuja = nuevas[0]
                else:
                    time.sleep(0.3)
            if burbuja:
                try:
                    pg.locator(f'[role="alertdialog"][aria-describedby="{burbuja}"] '
                               f'[role="radio"][data-value="{valor}"]').first.click(timeout=6000)
                except Exception as e:
                    if not type(e).__module__.startswith("playwright"):
                        raise
            if esperar(lambda: anotada.inner_text().strip().lower() == esperado, 4, 0.3):
                return burbuja or anterior or "?"
            time.sleep(1)
        return None

    def elegir_titulo(self):
        """Cuadro "Selecciona una columna...": el titulo de cada punto es la Razon."""
        dlg2 = self.cuadro_visible("Selecciona una columna", 30)
        time.sleep(0.5)
        dlg2.locator('[role="radio"][data-value="Razon"]').click()
        time.sleep(0.5)
        dlg2.get_by_role("button", name="Finalizar").click()

    def elegir_del_menu(self, texto):
        opcion = self.pg.locator("div.VIpgJd-j7LFlb-bN97Pc").filter(visible=True).filter(
            has_text=re.compile(rf"^{re.escape(texto)}$"))
        if not opcion.count():
            opcion = self.pg.get_by_text(texto, exact=True).filter(visible=True)
        opcion.last.click()

    def dar_estilo(self, ly):
        pg = self.pg
        pg.locator(f"#{ly}-layerview-stylepopup-link").click()
        popup = pg.locator("#layer-style-popup")
        popup.wait_for(state="visible")
        listas = popup.locator('[role="listbox"]').filter(visible=True)
        listas.first.click()
        self.elegir_del_menu("vendedor")
        if not esperar(lambda: listas.first.inner_text().strip() == "vendedor", 10):
            raise ErrorMyMaps("No se pudo agrupar los sitios por vendedor.")
        listas.last.click()
        self.elegir_del_menu("Razon")
        if not esperar(lambda: listas.last.inner_text().strip() == "Razon", 10):
            raise ErrorMyMaps("No se pudo poner las etiquetas con la Razon.")
        pg.keyboard.press("Escape")
        if not esperar(lambda: not popup.is_visible(), 3):
            popup.locator(".TvD9Pc, [aria-label='Cerrar']").first.click()

        categorias = pg.locator(f"#{ly}-layer-items-container div.un1lmc-pbTTYe-ibnC6b-JNdkSc")
        sin_color = []
        for i in range(categorias.count()):
            fila = categorias.nth(i)
            r, g, b = PALETA[i % len(PALETA)]
            if not self.color_de_fila(fila, r, g, b):
                sin_color.append(fila.inner_text().split("(")[0].strip() or str(i + 1))
        if sin_color:
            # El color es solo para distinguir vendedores: no vale perder el mapa por eso.
            self.log(f"   Aviso: no se pudo poner el color de {', '.join(sin_color)} (queda con el color "
                     "por defecto; se puede cambiar a mano).")

    def color_de_fila(self, fila, r, g, b) -> bool:
        """Pone el color a un vendedor de la capa. El balde de pintura solo se
        ve con el mouse encima de la fila, y con muchas capas el panel se
        desplaza: se trae la fila a la vista y se vuelve a pasar el mouse."""
        pg = self.pg
        balde = fila.locator("div.un1lmc-pbTTYe-ibnC6b-DyVDA")
        for _ in range(4):
            self.revisar_cancelado()
            try:
                fila.scroll_into_view_if_needed(timeout=5000)
                time.sleep(0.3)
                fila.hover(timeout=5000)
                balde.wait_for(state="visible", timeout=2500)
                balde.click(timeout=5000)
                pg.locator("#stylepopup-color").wait_for(state="visible", timeout=10000)
            except Exception as e:
                if type(e).__module__.startswith("playwright") or isinstance(e, TimeoutError):
                    pg.mouse.move(5, 5)          # sacar el mouse y volver a intentar
                    time.sleep(0.8)
                    continue
                raise
            pg.locator(f'#stylepopup-color [aria-label="RGB ({r}, {g}, {b})"]').click()
            hexa = f"{r:02x}{g:02x}{b:02x}"
            esperar(lambda: hexa in (fila.locator("[iconcode]").first.get_attribute("iconcode") or ""), 5)
            pg.locator("#stylepopup-close").click()
            return True
        return False

    def compartir(self) -> str:
        pg = self.pg
        pg.get_by_text("Compartir", exact=True).filter(visible=True).first.click()
        enlace, busqueda = pg.locator("#linksharing"), pg.locator("#discoverability")
        enlace.wait_for(state="visible", timeout=20000)
        if not enlace.is_checked():
            enlace.click()
            if not esperar(enlace.is_checked, 10):
                raise ErrorMyMaps("No se pudo activar \"Visible para cualquier persona con el enlace\".")
        if busqueda.is_checked():
            busqueda.click()
            esperar(lambda: not busqueda.is_checked(), 10)
        dialogo = pg.locator('div[role="dialog"]').filter(has=enlace).first
        url = ""
        for _ in range(20):
            campos = dialogo.locator("input[type=text], input:not([type])")
            for i in range(campos.count()):
                valor = campos.nth(i).input_value()
                if valor.startswith("https://www.google.com/maps/d/"):
                    url = valor
            if url:
                break
            time.sleep(0.5)
        if not url:
            url = f"https://www.google.com/maps/d/edit?mid={self.mid_en_curso}&usp=sharing"
        cerrar = dialogo.get_by_role("button", name="Cerrar").filter(visible=True)
        if cerrar.count():
            cerrar.first.click()
        else:
            pg.keyboard.press("Escape")
        return url

    def esperar_guardado(self):
        pg = self.pg
        time.sleep(2)
        esperar(lambda: not pg.get_by_text(re.compile("Guardando")).filter(visible=True).count(), 30, 1)
        time.sleep(1)

    def verificar(self, cuantas):
        """Vuelve a abrir el mapa y revisa que haya quedado guardado."""
        pg = self.pg
        pg.reload(wait_until="domcontentloaded")
        if not esperar(lambda: len(self.capas()) >= cuantas, 40, 1):
            raise ErrorMyMaps(f"Al volver a abrir el mapa no están las {cuantas} capas.")
        for ly in self.capas():
            if pg.locator(f'#{ly}-layer-header [role="checkbox"]').first.get_attribute("aria-checked") == "true":
                raise ErrorMyMaps("Al volver a abrir el mapa, alguna capa quedó marcada.")

    # -- mapas hechos a mano --------------------------------------------------
    def mapas_propios(self):
        """[{mid, titulo, fecha_lista}] de la lista "Propios" del inicio de My Maps."""
        pg = self.pg
        pg.goto(INICIO, wait_until="domcontentloaded")
        tarjetas = pg.locator("[data-mid][data-name]")
        if not esperar(tarjetas.count, 30, 1):
            return []
        pestana = pg.get_by_text(re.compile(r"^propios$", re.I)).filter(visible=True)
        if pestana.count():
            pestana.first.click()
            time.sleep(2)
        antes = -1
        for _ in range(40):                     # la lista carga de a poco al bajar
            n = tarjetas.count()
            if n == antes:
                break
            antes = n
            pg.evaluate("""() => { const t = Array.from(document.querySelectorAll('[data-mid][data-name]'))
                .filter(e => e.getBoundingClientRect().height > 0); if (t.length) t[t.length - 1].scrollIntoView(); }""")
            time.sleep(1.5)
        datos = pg.evaluate("""() => { const vistos = {}; for (const e of document.querySelectorAll('[data-mid][data-name]')) {
            const mid = e.getAttribute('data-mid'); if (vistos[mid]) continue;
            let card = e; for (let k = 0; k < 6 && card && !/\\d{4}/.test(card.innerText || ''); k++) card = card.parentElement;
            vistos[mid] = [e.getAttribute('data-name'), card ? ((card.innerText || '').split('\\n').filter(t => /\\d{4}/.test(t))[0] || '') : ''];
        } return vistos; }""")
        return [{"mid": mid, "titulo": nombre, "fecha_lista": fecha_de_lista(fecha)}
                for mid, (nombre, fecha) in datos.items() if mid and not mid.startswith("dummy")]

    # -- actualizar ubicaciones en un mapa ya subido -----------------------------
    def capa_por_nombre(self, nombre):
        for ly in self.capas():
            if self.pg.locator(f"#{ly}-layer-header").inner_text().strip() == nombre:
                return ly
        return None

    def orden_capas(self):
        return [self.pg.locator(f"#{ly}-layer-header").inner_text().strip() for ly in self.capas()]

    def abrir_opciones_capa(self, ly):
        pg = self.pg
        pg.keyboard.press("Escape")
        cabecera = pg.locator(f"#{ly}-layer-header")
        caja = cabecera.bounding_box()
        # El boton de opciones aparece con la capa elegida (clic en la cabecera).
        pg.mouse.click(caja["x"] + caja["width"] * 0.6, caja["y"] + caja["height"] / 2)
        time.sleep(0.6)
        cabecera.hover()
        opciones = pg.locator(f'#{ly}-layer-header [aria-label="Opciones de la capa"]')
        opciones.wait_for(state="visible", timeout=10000)
        opciones.click()
        time.sleep(0.6)

    def item_menu(self, texto, sin=None):
        items = self.pg.locator('[role="menuitem"]').filter(visible=True).filter(has_text=texto)
        if sin:
            items = items.filter(has_not_text=sin)
        items.first.wait_for(state="visible", timeout=10000)
        return items.first

    def sustituir_capa(self, nombre, archivo: Path):
        """Cambia todos los puntos de la capa por los del archivo ("Volver a
        importar y combinar" → "Volver a importar" → "Sustituir todos los
        elementos"): el mapa y su link siguen siendo los mismos."""
        pg = self.pg
        ly = self.capa_por_nombre(nombre)
        if not ly:
            raise ErrorMyMaps(f"En el mapa no está la capa {nombre}.")
        # Al terminar de sustituir, My Maps deja la capa MARCADA (y la pasa al
        # final): esa es la señal de fin. Por eso tiene que empezar desmarcada.
        # (Los ids de los puntos no sirven: cambian solos y, con muchas capas,
        # la lista de una capa fuera de la vista ni se dibuja.)
        if self.capa_marcada(ly):
            pg.locator(f'#{ly}-layer-header [role="checkbox"]').first.click()
            if not esperar(lambda: not self.capa_marcada(ly), 10, 0.5):
                raise ErrorMyMaps(f"No se pudo desmarcar la capa {nombre} antes de volver a importarla.")
            time.sleep(1)
        self.abrir_opciones_capa(ly)
        self.item_menu("Volver a importar y combinar").hover()
        time.sleep(0.6)
        self.item_menu("Volver a importar", sin="combinar").hover()
        time.sleep(0.6)
        self.item_menu("Sustituir todos los elementos").click()
        if not esperar(self.selector_de_archivos, 150, 1):
            raise ErrorMyMaps("No se abrió la ventana para volver a importar el archivo.")
        selector = self.selector_de_archivos()
        selector.get_by_text("Subir", exact=True).filter(visible=True).first.click()
        entrada = selector.locator("input[type=file]")
        entrada.wait_for(state="attached", timeout=20000)
        entrada.set_input_files(str(archivo))
        # Normalmente recuerda las columnas; si las vuelve a pedir, se responden igual.
        fin = time.time() + 180
        while time.time() < fin:
            time.sleep(1.5)
            if self.pg.locator('div[role="dialog"]').filter(has_text="Selecciona columnas").filter(visible=True).count():
                self.marcar_columnas(archivo)
                continue
            if self.pg.locator('div[role="dialog"]').filter(has_text="Selecciona una columna").filter(visible=True).count():
                self.elegir_titulo()
                continue
            if self.selector_de_archivos() or pg.locator('div[role="dialog"]').filter(visible=True).count():
                continue
            nuevo = self.capa_por_nombre(nombre)
            if nuevo and self.capa_marcada(nuevo):
                time.sleep(2)
                return self.capa_por_nombre(nombre)
        raise ErrorMyMaps(f"No terminó de volver a importar {nombre}.")

    def capa_marcada(self, ly) -> bool:
        return self.pg.locator(f'#{ly}-layer-header [role="checkbox"]').first.get_attribute("aria-checked") == "true"

    def mover_capa_arriba_de(self, ly_mover, ly_destino):
        """Arrastra una capa para dejarla encima de otra (como una persona)."""
        pg = self.pg
        o = pg.locator(f"#{ly_mover}-layer-header").bounding_box()
        d = pg.locator(f"#{ly_destino}-layer-header").bounding_box()
        x0, y0 = o["x"] + o["width"] * 0.35, o["y"] + o["height"] / 2
        pg.mouse.move(x0, y0)
        time.sleep(0.3)
        pg.mouse.down()
        time.sleep(0.3)
        pg.mouse.move(x0, y0 - 8, steps=4)
        time.sleep(0.3)
        pg.mouse.move(x0, d["y"] - 2, steps=30)
        time.sleep(0.6)
        pg.mouse.move(x0, d["y"] - 1, steps=2)
        time.sleep(0.4)
        pg.mouse.up()
        time.sleep(2)

    def ordenar_capas(self, orden):
        """Deja las capas en el orden dado (las sustituidas quedan al final)."""
        for i, nombre in enumerate(orden):
            for _ in range(3):
                actual = self.orden_capas()
                if nombre not in actual or actual.index(nombre) == i:
                    break
                self.mover_capa_arriba_de(self.capa_por_nombre(nombre), self.capa_por_nombre(actual[i]))

    def actualizar_mapa(self, mid, capas: dict, orden, tocadas=None):
        """Sustituye las capas indicadas ({nombre: copia}) y deja el mapa como
        siempre: estilo por vendedor, orden de capas y todas desmarcadas.
        tocadas: lista donde se anota cada capa antes de empezar a cambiarla
        (si algo falla, esas pueden haber quedado a medias)."""
        pg = self.pg
        self.mid_en_curso = mid
        pg.goto(f"https://www.google.com/maps/d/edit?mid={mid}&hl=es", wait_until="domcontentloaded")
        # Esperar a que se vean los nombres de las capas (la pagina carga en blanco un rato).
        if not esperar(lambda: set(capas) <= set(self.orden_capas()), 60, 1):
            if "/maps/d/home" in pg.url:
                # Se llego al mapa y My Maps mando al inicio (lo hace con los
                # mapas en la papelera o borrados).
                raise MapaNoExiste("El mapa ya no existe o está en la papelera; no se pudo actualizar.")
            faltan = sorted(set(capas) - set(self.orden_capas()))
            raise ErrorMyMaps(f"En el mapa no están las capas {', '.join(faltan)} (¿se borraron a mano?).")
        time.sleep(2)
        for nombre, archivo in capas.items():
            self.revisar_cancelado()
            self.log(f"   capa {nombre}")
            if tocadas is not None:
                tocadas.append(nombre)
            ly = self.sustituir_capa(nombre, archivo)
            self.dar_estilo(ly)
        # La capa sustituida pasa al final, pero a veces un rato despues: se
        # recarga el mapa antes de ordenar para ver el orden de verdad.
        self.esperar_guardado()
        presentes = [n for n in orden if n in self.orden_capas()]
        for intento in range(3):
            pg.reload(wait_until="domcontentloaded")
            esperar(lambda: set(presentes) <= set(self.orden_capas()), 60, 1)
            time.sleep(2)
            if self.orden_capas()[:len(presentes)] == presentes:
                break
            if intento == 2:
                self.log("   Aviso: no se pudo dejar las capas en su orden de siempre (los puntos sí quedaron bien).")
                break
            try:
                self.ordenar_capas(presentes)
            except Cancelado:
                raise
            except Exception as e:
                # Con muchas capas el panel se desplaza y arrastrar puede fallar;
                # el orden es solo de presentacion: se reintenta y si no, aviso.
                self.log(f"   (no se pudo mover una capa: {str(e).splitlines()[0][:100]})")
            self.esperar_guardado()
        for ly in self.capas():
            casilla = pg.locator(f'#{ly}-layer-header [role="checkbox"]').first
            if casilla.get_attribute("aria-checked") == "true":
                casilla.click()
                esperar(lambda: casilla.get_attribute("aria-checked") == "false", 5)
        self.esperar_guardado()
        self.verificar(len(orden))

    # -- papelera -----------------------------------------------------------
    def a_la_papelera(self, mid) -> None:
        pg = self.pg
        pg.goto(f"https://www.google.com/maps/d/edit?mid={mid}&hl=es", wait_until="domcontentloaded")
        menu = pg.locator("#map-action-menu")
        if not esperar(lambda: menu.count() or "/maps/d/home" in pg.url, 30, 1):
            raise ErrorMyMaps("No se pudo abrir el mapa para mandarlo a la papelera.")
        if not menu.count():
            # My Maps manda al inicio si el mapa ya esta en la papelera (o no existe).
            return
        menu.click()
        item = pg.locator("#mapmenu-trashmap").filter(visible=True).first
        item.wait_for(state="visible")
        item.click()
        dialogo = pg.locator('div[role="dialog"]').filter(has_text="papelera").filter(visible=True).first
        dialogo.wait_for(state="visible")
        dialogo.locator('[name="trash"]').click()
        if not esperar(lambda: "/maps/d/home" in pg.url or mid not in pg.url, 30, 1):
            raise ErrorMyMaps("My Maps no confirmó que el mapa fue a la papelera.")


# ---------------------------------------------------------------------------
# Todo el proceso
# ---------------------------------------------------------------------------

def subir(carpeta: Path, fecha_entrega: date, log, cancelar=None, reemplazar=False,
          registro_ruta: Path = ARCHIVO_REGISTRO, overrides=None, hoy: date = None):
    """Crea en My Maps los mapas de la carpeta del dia. Antes manda a la
    papelera los vencidos (y, si reemplazar, los de la misma entrega).
    Devuelve los mapas creados (con url)."""
    hoy = hoy or date.today()
    archivos = archivos_mapa(carpeta)
    if not archivos:
        raise ErrorMyMaps(f"No hay archivos MAPA en:\n{carpeta}")
    cambiados = cambiados_desde_publicar(carpeta)
    if cambiados:
        raise ErrorMyMaps(texto_cambiados(cambiados) + "\n(No se subió nada.)")
    if overrides is None:
        from generar_rutas import overrides_vigentes
        log("Revisando la georreferencia de los clientes...")
        overrides, al_dia = overrides_vigentes()
        if not al_dia:
            log("Aviso: sin conexión con GitHub; se usan las ubicaciones corregidas descargadas la última vez.")
    faltan = revisar_georef(archivos, overrides)
    if faltan:
        from generar_rutas import texto_sin_georef
        raise ErrorMyMaps(f"{len(faltan)} clientes no tienen georreferencia y todos deben tenerla:\n\n"
                          f"{texto_sin_georef(faltan)}\n\nCorrígelos en \"Actualizar ruta\" (lista en rojo: \"1. "
                          "Buscar en Google Maps\" → \"2. Pegar ubicación copiada\"), vuelve a \"Generar y publicar\" "
                          "y luego sube a My Maps.\n(No se subió nada.)")

    registro = leer_registro(registro_ruta)
    entrega = fecha_entrega.isoformat()
    if activos(registro, entrega) and not reemplazar:
        raise ErrorMyMaps(f"Ya hay mapas en My Maps para el {fecha_larga(fecha_entrega)}.")

    temporal = Path(tempfile.mkdtemp(prefix="mymaps_"))
    creados = []
    try:
        copias, corregidas = preparar_copias(archivos, temporal, overrides)
        if corregidas:
            log(f"{corregidas} clientes van con la ubicación corregida.")
        planes = armar_mapas(copias, fecha_entrega)

        with Robot(log, cancelar) as robot:
            log("Revisando la sesión de Google...")
            robot.asegurar_sesion()
            ya_conocidos = {m.get("mid") for m in activos(registro, entrega)}
            _anotar_manuales(robot, registro, hoy, log, registro_ruta)
            hechos_a_mano = [m for m in activos(registro, entrega) if m.get("mid") not in ya_conocidos]
            if hechos_a_mano:
                # No se sabia de ellos al preguntar si se reemplazaba esa fecha.
                raise ErrorMyMaps(
                    f"Encontré {len(hechos_a_mano)} mapas hechos a mano para el {fecha_larga(fecha_entrega)}:\n"
                    + "\n".join(f"• {m['titulo']}" for m in hechos_a_mano)
                    + "\n\nYa quedaron en la lista de links. Vuelve a darle a «Subir a My Maps» y te pregunto "
                      "si se reemplazan.\n(No se subió nada.)")
            # Los anteriores de esa entrega (si se reemplaza) van a la papelera
            # DESPUES de crear los nuevos: si algo falla a medias, los links de
            # antes siguen sirviendo.
            anteriores = activos(registro, entrega) if reemplazar else []
            _mandar_a_papelera(robot, registro, por_mandar_a_papelera(registro, hoy), log, registro_ruta)

            for plan in planes:
                robot.revisar_cancelado()
                log(f"Creando {plan['titulo']} ({len(plan['archivos'])} capas)...")
                entrada = {
                    "titulo": plan["titulo"], "descripcion": plan["descripcion"],
                    "transportistas": plan["transportistas"], "capas": [a.name for a in plan["archivos"]],
                    "numero": plan["numero"], "fecha_entrega": entrega,
                    "fecha_liq": (fecha_de_carpeta(carpeta) or hoy).isoformat(), "carpeta": str(carpeta),
                    # Huella de lo subido en cada capa: si despues se corrige una
                    # ubicacion, se sabe que capas hay que actualizar.
                    "huellas": {a.name: huella_capa(a) for a in plan["archivos"]},
                    "creado": datetime.now().isoformat(timespec="seconds"), "estado": "creando",
                }
                registro["mapas"].append(entrada)
                guardar_registro(registro, registro_ruta)

                def anotar_mid(mid, entrada=entrada):
                    entrada["mid"] = mid
                    guardar_registro(registro, registro_ruta)

                robot.al_crear = anotar_mid
                try:
                    resultado = robot.crear_mapa(plan["titulo"], plan["descripcion"],
                                                 plan["archivos"])
                except Exception:
                    try:
                        robot.pg.screenshot(path=str(CAPTURA_ERROR))
                        log(f"   (Captura de la pantalla del error en {CAPTURA_ERROR})")
                    except Exception:
                        pass
                    entrada["mid"] = robot.mid_en_curso
                    entrada["estado"] = "fallido" if robot.mid_en_curso else "sin_crear"
                    guardar_registro(registro, registro_ruta)
                    if robot.mid_en_curso:
                        try:
                            robot.a_la_papelera(robot.mid_en_curso)
                            entrada["estado"] = "papelera"
                            guardar_registro(registro, registro_ruta)
                        except Exception:
                            pass
                    raise
                entrada.update(resultado, estado="activo")
                guardar_registro(registro, registro_ruta)
                creados.append(entrada)
                log(f"   Listo: {resultado['url']}")
            if anteriores:
                log(f"Ya están los nuevos; los {len(anteriores)} mapas anteriores del "
                    f"{fecha_larga(fecha_entrega)} van a la papelera...")
                _, fallidos = _mandar_a_papelera(robot, registro, anteriores, log, registro_ruta)
                for m in fallidos:
                    # Ya no cuenta como activo (no sale en el mensaje del dia ni
                    # en los links); se vuelve a intentar la proxima vez.
                    m["estado"] = "reemplazado"
                if fallidos:
                    guardar_registro(registro, registro_ruta)
                    log("   Aviso: esos se mandan a la papelera la próxima vez; ya no salen en la lista de links.")
        limpiar_registro(registro, hoy)
        guardar_registro(registro, registro_ruta)
    finally:
        shutil.rmtree(temporal, ignore_errors=True)
    return creados


def _anotar_manuales(robot, registro, hoy, log, registro_ruta):
    """Suma al registro los mapas de entregas hechos a mano (para la lista
    de links y la limpieza de 30 dias). Si falla, se sigue sin ellos."""
    try:
        n = registrar_manuales(registro, robot.mapas_propios(), hoy)
        if n:
            log(f"Se anotaron {n} mapas hechos a mano (se limpian a los {DIAS_HISTORIAL} días de su entrega).")
            guardar_registro(registro, registro_ruta)
    except Cancelado:
        raise
    except Exception as e:
        log(f"   Aviso: no se pudo revisar la lista de mapas hechos a mano ({str(e).splitlines()[0][:120]}).")


def _mandar_a_papelera(robot, registro, viejos, log, registro_ruta):
    """Devuelve (enviados, fallidos)."""
    if viejos:
        log(f"Mandando a la papelera {len(viejos)} mapas (vencidos o reemplazados)...")
    enviados, fallidos = [], []
    for m in viejos:
        robot.revisar_cancelado()
        try:
            robot.a_la_papelera(m["mid"])
            m["estado"], m["papelera"] = "papelera", datetime.now().isoformat(timespec="seconds")
            enviados.append(m)
            log(f"   a la papelera: {m.get('titulo')}")
        except Cancelado:
            raise
        except Exception as e:
            fallidos.append(m)
            log(f"   Aviso: no se pudo mandar a la papelera \"{m.get('titulo')}\" ({str(e).splitlines()[0][:150]}).")
        guardar_registro(registro, registro_ruta)
    return enviados, fallidos


def limpiar(log, cancelar=None, registro_ruta: Path = ARCHIVO_REGISTRO, hoy: date = None):
    """Solo la limpieza: anota los mapas hechos a mano y manda a la papelera
    los de entrega de hace mas de 30 dias (y los que quedaron a medias).
    Devuelve los que se mandaron a la papelera; si alguno no se pudo,
    ErrorMyMaps al final (se reintenta la proxima vez)."""
    hoy = hoy or date.today()
    registro = leer_registro(registro_ruta)
    with Robot(log, cancelar) as robot:
        log("Revisando la sesión de Google...")
        robot.asegurar_sesion()
        _anotar_manuales(robot, registro, hoy, log, registro_ruta)
        enviados, fallidos = _mandar_a_papelera(robot, registro, por_mandar_a_papelera(registro, hoy), log,
                                                registro_ruta)
    limpiar_registro(registro, hoy)
    guardar_registro(registro, registro_ruta)
    if fallidos:
        listos = ("\n\nSí se mandaron:\n" + "\n".join(f"• {m.get('titulo')}" for m in enviados)) if enviados else ""
        raise ErrorMyMaps(f"No se pudieron mandar a la papelera {len(fallidos)} mapas (se intenta de nuevo la "
                          "próxima vez):\n" + "\n".join(f"• {m.get('titulo')}" for m in fallidos) + listos)
    return enviados


def actualizar_ubicaciones(log, cancelar=None, registro_ruta: Path = ARCHIVO_REGISTRO, overrides=None,
                           hoy: date = None):
    """Pasa a My Maps los cambios (ubicaciones corregidas o la ruta vuelta a
    publicar): en cada mapa de hoy en adelante, sustituye solo las capas que
    cambiaron (el link no cambia). Devuelve [(titulo, [capas])] actualizados.
    Si algun mapa falla, se sigue con los demas y al final ErrorMyMaps con
    cuales fallaron. Si no se pudo actualizar nada por un aviso (ruta sin
    publicar, camiones cambiados...), ErrorMyMaps con los avisos."""
    hoy = hoy or date.today()
    if not actualizables(leer_registro(registro_ruta), hoy):
        return []
    if overrides is None:
        from generar_rutas import overrides_vigentes
        overrides, al_dia = overrides_vigentes()
        if not al_dia:
            # Con lo descargado la ultima vez se podria volver a poner una
            # ubicacion vieja en My Maps: mejor no tocar nada.
            raise ErrorMyMaps("Sin conexión con GitHub no se pueden revisar las ubicaciones corregidas. "
                              "Vuelve a intentar con internet.")
    registro = leer_registro(registro_ruta)
    temporal = Path(tempfile.mkdtemp(prefix="mymaps_act_"))
    hechos, fallidos = [], []
    try:
        planes, avisos = mapas_por_actualizar(registro, overrides, hoy, temporal)
        for a in avisos:
            log("Aviso: " + a)
        if not planes:
            if avisos:
                raise ErrorMyMaps("No se actualizó nada en My Maps:\n" + "\n".join(f"• {a}" for a in avisos))
            return hechos
        por_mid = {m.get("mid"): m for m in registro["mapas"]}
        with Robot(log, cancelar) as robot:
            log("Revisando la sesión de Google...")
            robot.asegurar_sesion()
            for plan in planes:
                robot.revisar_cancelado()
                m = por_mid[plan["mapa"]["mid"]]
                log(f"Actualizando {m['titulo']} ({len(plan['capas'])} capas)...")
                tocadas = []
                try:
                    robot.actualizar_mapa(m["mid"], plan["capas"], m.get("capas") or list(plan["capas"]), tocadas)
                except Exception as e:
                    # Las capas que se empezaron a cambiar pueden haber quedado a
                    # medias (o ya nuevas, pero sin ordenar): sin huella, se
                    # vuelven a ofrecer y se rehacen.
                    for nombre in tocadas:
                        m.setdefault("huellas", {})[nombre] = ""
                    guardar_registro(registro, registro_ruta)
                    if isinstance(e, Cancelado):
                        raise
                    try:
                        robot.pg.screenshot(path=str(CAPTURA_ERROR))
                        log(f"   (Captura de la pantalla del error en {CAPTURA_ERROR})")
                    except Exception:
                        pass
                    motivo = str(e).splitlines()[0][:200] if str(e) else type(e).__name__
                    if isinstance(e, MapaNoExiste):
                        # Lo mandaron a la papelera (o lo borraron) a mano: ya no se ofrece.
                        m["estado"], m["papelera"] = "papelera", datetime.now().isoformat(timespec="seconds")
                        guardar_registro(registro, registro_ruta)
                        motivo = "ya no existe o está en la papelera (se deja de ofrecer)"
                    log(f"   No se pudo: {motivo}")
                    fallidos.append(f"{m['titulo']}: {motivo}")
                    continue
                for nombre, copia in plan["capas"].items():
                    m.setdefault("huellas", {})[nombre] = huella_capa(copia)
                m["actualizado"] = datetime.now().isoformat(timespec="seconds")
                guardar_registro(registro, registro_ruta)
                hechos.append((m["titulo"], sorted(plan["capas"])))
                log("   Listo (mismo link).")
    finally:
        shutil.rmtree(temporal, ignore_errors=True)
    if fallidos:
        listos = ("\n\nSí se actualizaron:\n" + "\n".join(f"• {t}" for t, _ in hechos)) if hechos else ""
        raise ErrorMyMaps("No se pudieron actualizar:\n" + "\n".join(f"• {f}" for f in fallidos) + listos)
    return hechos
