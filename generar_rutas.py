"""
Ventana para actualizar la ruta del dia directo desde el archivo base
"RUTAS DIARIAS CAMIONES - MY MAPS.xlsx" (hoja SABANA DE VENTA), sin pasar
por la tabla dinamica de "dia" ni por RESUMEN:

1. Lee la hoja SABANA DE VENTA (aunque el Excel este abierto).
2. Propone la fecha de liquidacion mas reciente (si es sabado, junta
   viernes + sabado) y muestra los camiones de esa fecha para marcar
   cuales van al mapa. Recuerda lo desmarcado para la proxima vez.
3. Genera los "MAPA Fxxx.xlsx" (iguales a los que salian de la tabla
   dinamica) en la carpeta del archivo base, y data/latest.json.
4. Si se pide, publica en el mapa (lo mismo que publicar_datos.bat).

Se abre con doble clic en "ACTUALIZAR RUTA.bat".

Prueba sin ventana (no publica ni toca data/):
    python generar_rutas.py --prueba "<archivo base>" "<carpeta de salida>"
"""
import ctypes
import html
import json
import os
import queue
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import traceback
import unicodedata
import xml.etree.ElementTree as ET
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font

    from archivar_dia import archivar, guardar_en_historial
    from convertir import COLUMNAS_ESPERADAS, asignar_colores, cliente_desde_fila, texto
except Exception:
    # Abierto con doble clic (pythonw) no hay consola: sin esto, un error
    # aqui haria que la ventana simplemente no abra, sin decir por que.
    _detalle = traceback.format_exc()
    try:
        _log = Path(__file__).resolve().parent / "data" / "generador_error.log"
        _log.parent.mkdir(exist_ok=True)
        with open(_log, "a", encoding="utf-8") as _f:
            _f.write(f"--- {datetime.now():%Y-%m-%d %H:%M:%S}\n{_detalle}\n")
        if os.name == "nt":
            ctypes.windll.user32.MessageBoxW(None, "No se pudo abrir Actualizar Ruta:\n\n"
                                             + _detalle.strip().splitlines()[-1], "Actualizar Ruta", 0x10)
    except Exception:
        pass
    raise

CARPETA = Path(__file__).resolve().parent
SALIDA_JSON = CARPETA / "data" / "latest.json"
ARCHIVO_CONFIG = CARPETA / "data" / "generador_config.json"
ARCHIVO_ERRORES = CARPETA / "data" / "generador_error.log"
REPO_DATOS = Path(r"C:\MAPA-RUTAS-DATOS")
CARPETA_MAPAS_DIARIOS = (Path.home() / "OneDrive" / "Documentos"
                         / "2AVANCES DE RUTA Y MAPAS DE ENTREGA" / "MAPAS DIARIOS")
NOMBRE_BASE = "RUTAS DIARIAS CAMIONES - MY MAPS.xlsx"
HOJA_BASE = "SABANA DE VENTA"
# Maestro fijo: un solo archivo base de siempre, y los MAPA de cada dia en
# BASE\MAPAS\<fecha> (BASE/ esta en .gitignore: nunca se sube a GitHub).
CARPETA_BASE = CARPETA / "BASE"
BASE_FIJA = CARPETA_BASE / NOMBRE_BASE
CARPETA_MAPAS_FIJA = CARPETA_BASE / "MAPAS"
DIAS_LARGOS = ["LUNES", "MARTES", "MIERCOLES", "JUEVES", "VIERNES", "SABADO", "DOMINGO"]

# Lo mismo que hace la tabla dinamica de la hoja "dia": agrupa por estos
# campos (en este orden, que tambien es el orden de las filas) ...
CAMPOS_GRUPO = ["Fecha Liq", "Chofer", "Numero Liq", "vendedor", "CodigoCliente",
                "Razon", "Direccion", "canal", "longitud_x", "latitud_y"]
# ... y suma estos (las columnas "Suma de ..." de los MAPA).
CAMPOS_SUMA = ["Peso", "VentaNeta", "Total"]
assert COLUMNAS_ESPERADAS == CAMPOS_GRUPO + ["Suma de " + c for c in CAMPOS_SUMA]

DIAS = ["lun", "mar", "mié", "jue", "vie", "sáb", "dom"]
FECHAS_A_MOSTRAR = 6


class ErrorRuta(Exception):
    """Error con un mensaje ya pensado para mostrarse tal cual."""


def normalizar(valor) -> str:
    # Nunca comparar nombres de hoja/encabezados con tilde tal cual: una
    # tilde copiada de otro lado puede venir en otra forma unicode.
    s = unicodedata.normalize("NFKD", str(valor or ""))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return " ".join(s.upper().split())


def es_numero(valor) -> bool:
    return isinstance(valor, (int, float)) and not isinstance(valor, bool)


def formato_fecha(d: date) -> str:
    return f"{DIAS[d.weekday()]} {d:%d/%m/%Y}"


def formato_fechas(fechas) -> str:
    return " + ".join(formato_fecha(d) for d in sorted(fechas))


# ---------------------------------------------------------------------------
# Lectura rapida del archivo base
#
# openpyxl tarda mas de un minuto con la SABANA (55 mil filas x 57
# columnas). Aqui se lee el XML de la hoja directamente y solo se
# decodifican las 13 columnas que hacen falta: unos segundos.
# ---------------------------------------------------------------------------

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
NS_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
RE_CELDA = re.compile(rb'<c r="([A-Z]{1,3})(\d+)"([^>]*?)(?:/>|>(.*?)</c>)', re.S)
RE_VALOR = re.compile(rb"<v>([^<]*)</v>")
RE_TEXTO = re.compile(rb"<t[^>]*>([^<]*)</t>")
RE_ENTERO = re.compile(r"-?\d+")
FILAS_PARA_ENCABEZADO = 30


def copiar_sin_bloqueo(origen: Path) -> Path:
    """Copia el archivo a la carpeta temporal para poder leerlo aunque este
    abierto en Excel (Excel no deja abrirlo directo desde Python, pero la
    copia de Windows si funciona)."""
    descriptor, nombre = tempfile.mkstemp(prefix="mrd_", suffix=".xlsx")
    os.close(descriptor)
    destino = Path(nombre)
    if os.name == "nt" and ctypes.windll.kernel32.CopyFileW(str(origen), str(destino), False):
        return destino
    shutil.copyfile(origen, destino)
    return destino


def hojas_del_libro(z: zipfile.ZipFile):
    libro = ET.fromstring(z.read("xl/workbook.xml"))
    pr = libro.find(NS + "workbookPr")
    fecha1904 = pr is not None and pr.get("date1904") in ("1", "true")
    rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    destinos = {r.get("Id"): r.get("Target", "") for r in rels}
    hojas = []
    for hoja in libro.iter(NS + "sheet"):
        parte = destinos.get(hoja.get(NS_REL + "id"), "").lstrip("/")
        if not parte.startswith("xl/"):
            parte = "xl/" + parte
        hojas.append((hoja.get("name"), parte))
    return hojas, fecha1904


def textos_compartidos(z: zipfile.ZipFile):
    if "xl/sharedStrings.xml" not in z.namelist():
        return []
    textos = []
    with z.open("xl/sharedStrings.xml") as f:
        for _, el in ET.iterparse(f):
            if el.tag == NS + "si":
                t = el.find(NS + "t")
                if t is not None:
                    textos.append(t.text or "")
                else:
                    textos.append("".join(r.findtext(NS + "t") or "" for r in el.findall(NS + "r")))
                el.clear()
    return textos


def valor_celda(atributos: bytes, contenido, textos):
    """Mismo tipo que devolveria openpyxl (int si no tiene decimales)."""
    if contenido is None:
        return None
    if b't="s"' in atributos:
        v = RE_VALOR.search(contenido)
        return textos[int(v.group(1))] if v else None
    if b't="inlineStr"' in atributos:
        return html.unescape(b"".join(RE_TEXTO.findall(contenido)).decode("utf-8"))
    v = RE_VALOR.search(contenido)
    if v is None:
        return None
    crudo = v.group(1).decode("utf-8")
    if b't="str"' in atributos:
        return html.unescape(crudo)
    if b't="e"' in atributos:
        return None
    if b't="b"' in atributos:
        return crudo == "1"
    if RE_ENTERO.fullmatch(crudo):
        return int(crudo)
    try:
        return float(crudo)
    except ValueError:
        return crudo


def columnas_si_es_encabezado(celdas: dict, buscados: dict):
    """Si la fila tiene todos los encabezados buscados, devuelve
    {letra de columna: campo}; si no, None."""
    por_nombre = {normalizar(v): letra for letra, v in celdas.items() if isinstance(v, str)}
    if not all(n in por_nombre for n in buscados):
        return None
    return {por_nombre[n]: campo for n, campo in buscados.items()}


def filas_de_hoja(datos: bytes, textos, campos):
    """Filas (dict campo -> valor) de una hoja, buscando el encabezado en
    las primeras filas. None si la hoja no tiene esos encabezados."""
    buscados = {normalizar(c): c for c in campos}

    # 1) Encontrar la fila de encabezados (y donde empiezan los datos).
    columnas = None
    inicio_datos = len(datos)
    fila_actual = None
    celdas = {}
    vistas = 0
    for m in RE_CELDA.finditer(datos):
        if m.group(2) != fila_actual:
            if celdas:
                columnas = columnas_si_es_encabezado(celdas, buscados)
                if columnas is not None:
                    inicio_datos = m.start()
                    break
            vistas += 1
            if vistas > FILAS_PARA_ENCABEZADO:
                return None
            fila_actual = m.group(2)
            celdas = {}
        if m.group(4) is not None:
            celdas[m.group(1)] = valor_celda(m.group(3), m.group(4), textos)
    if columnas is None and celdas:
        columnas = columnas_si_es_encabezado(celdas, buscados)
    if columnas is None:
        return None

    # 2) Leer solo las columnas que hacen falta: la busqueda salta solita
    #    las otras 40 y pico, que es lo que hace esto rapido.
    letras = b"|".join(re.escape(letra) for letra in columnas)
    re_util = re.compile(rb'<c r="(' + letras + rb')(\d+)"([^>]*?)(?:/>|>(.*?)</c>)', re.S)
    filas = []
    fila_actual = None
    for m in re_util.finditer(datos, inicio_datos):
        contenido = m.group(4)
        if contenido is None:
            continue
        if m.group(2) != fila_actual:
            fila_actual = m.group(2)
            celdas = {}
            filas.append(celdas)
        celdas[columnas[m.group(1)]] = valor_celda(m.group(3), contenido, textos)
    return filas


def filas_con_openpyxl(ruta: Path, nombre_hoja: str, campos):
    """Camino lento de respaldo, por si el archivo no lo guardo Excel y la
    lectura rapida no reconoce su formato."""
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        ws = wb[nombre_hoja]
        buscados = {normalizar(c): c for c in campos}
        columnas = None
        filas = []
        for i, fila in enumerate(ws.iter_rows(values_only=True)):
            if columnas is None:
                if i >= FILAS_PARA_ENCABEZADO:
                    return None
                celdas = {j: v for j, v in enumerate(fila) if v is not None}
                columnas = columnas_si_es_encabezado(celdas, buscados)
                continue
            filas.append({campo: fila[j] for j, campo in columnas.items() if j < len(fila)})
        return filas if columnas is not None else None
    finally:
        wb.close()


def a_fecha(valor, base: date):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if es_numero(valor) and valor > 0:
        try:
            return base + timedelta(days=int(valor))
        except (OverflowError, ValueError):
            return None         # un numero enorme (ej. 09102026 tecleado sin barras)
    if isinstance(valor, str):
        for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S",
                        "%d/%m/%Y %H:%M:%S"):
            try:
                return datetime.strptime(valor.strip(), formato).date()
            except ValueError:
                pass
    return None


def leer_base(ruta: Path):
    """Filas de la hoja SABANA DE VENTA (o de la primera hoja que tenga
    esas columnas, por ejemplo la descarga del sistema tal cual), con
    'Fecha Liq' ya convertida a fecha."""
    campos = CAMPOS_GRUPO + CAMPOS_SUMA
    if not ruta.exists():
        raise ErrorRuta(f"No existe el archivo:\n{ruta}")
    copia = copiar_sin_bloqueo(ruta)
    try:
        try:
            z = zipfile.ZipFile(copia)
        except zipfile.BadZipFile:
            raise ErrorRuta("Ese archivo no es un Excel .xlsx (si es .xls o .csv, ábrelo y "
                            "guárdalo como .xlsx).")
        with z:
            hojas, fecha1904 = hojas_del_libro(z)
            principal = normalizar(HOJA_BASE)
            hojas.sort(key=lambda h: normalizar(h[0]) != principal)
            textos = None
            filas = None
            for nombre, parte in hojas:
                if parte not in z.namelist():
                    continue
                if textos is None:
                    textos = textos_compartidos(z)
                filas = filas_de_hoja(z.read(parte), textos, campos)
                if filas is None and normalizar(nombre) == principal:
                    filas = filas_con_openpyxl(copia, nombre, campos)
                if filas is not None:
                    break
        if filas is None:
            raise ErrorRuta("No encontré la hoja SABANA DE VENTA (ni otra hoja con las columnas "
                            + ", ".join(campos) + ").")
    finally:
        try:
            copia.unlink()
        except OSError:
            pass

    base = date(1904, 1, 1) if fecha1904 else date(1899, 12, 30)
    utiles = []
    for fila in filas:
        fila["Fecha Liq"] = a_fecha(fila.get("Fecha Liq"), base)
        if fila["Fecha Liq"] and fila.get("Chofer"):
            utiles.append(fila)
    if not utiles:
        raise ErrorRuta("La hoja no tiene filas con Fecha Liq y Chofer.")
    return utiles


# ---------------------------------------------------------------------------
# Armado de la ruta (lo que antes hacia la tabla dinamica + RESUMEN)
# ---------------------------------------------------------------------------

RE_CAMION = re.compile(r"^\s*([A-Za-z]{1,3})\s*-?\s*(\d{1,4})")


def id_camion(chofer: str) -> str:
    """'F150-PEREZ JUAN' -> 'F150', 'FO54- LOPEZ ANA' -> 'FO54'."""
    m = RE_CAMION.match(chofer)
    if m:
        return (m.group(1) + m.group(2)).upper()
    return chofer.split("-")[0].strip().upper() or "SIN-CAMION"


def nombre_chofer(chofer: str) -> str:
    resto = RE_CAMION.sub("", chofer, count=1).strip(" -")
    return resto or chofer


def fechas_sugeridas(fechas) -> set:
    """La fecha mas reciente; si es sabado, tambien el viernes anterior."""
    ultima = max(fechas)
    sugeridas = {ultima}
    if ultima.weekday() == 5:
        viernes = ultima - timedelta(days=1)
        if viernes in fechas:
            sugeridas.add(viernes)
    return sugeridas


def orden_valor(valor):
    # Como ordena la tabla dinamica: numeros, luego textos (sin importar
    # mayusculas), y los vacios al final.
    if valor is None or valor == "":
        return (2, 0, "")
    if isinstance(valor, date):
        return (0, valor.toordinal(), "")
    if es_numero(valor):
        return (0, valor, "")
    return (1, 0, str(valor).casefold())


def agrupar(filas, fechas):
    """Camiones (ordenados por id) con sus filas ya agrupadas y sumadas,
    solo de las fechas indicadas."""
    grupos = {}
    movidos = {}    # clave -> ruta de origen, para clientes pasados a otra ruta
    for fila in filas:
        if fila["Fecha Liq"] not in fechas:
            continue
        clave = tuple(fila.get(c) for c in CAMPOS_GRUPO)
        sumas = grupos.get(clave)
        if sumas is None:
            sumas = grupos[clave] = [0.0] * len(CAMPOS_SUMA)
        for i, campo in enumerate(CAMPOS_SUMA):
            v = fila.get(campo)
            if es_numero(v):
                sumas[i] += v
        if fila.get("_movido_de"):
            movidos[clave] = (fila["_movido_de"], fila.get("_cargue_destino"))

    camiones = {}
    # Los clientes que vienen de otra ruta van al final de la suya (asi no
    # cambia el orden ni los colores de vendedor de los clientes propios).
    for clave in sorted(grupos, key=lambda k: (k in movidos, tuple(orden_valor(v) for v in k))):
        fila = dict(zip(CAMPOS_GRUPO, clave))
        for i, campo in enumerate(CAMPOS_SUMA):
            fila["Suma de " + campo] = round(grupos[clave][i], 4)
        if clave in movidos:
            fila["_movido_de"], cargue = movidos[clave]
            if cargue:
                fila["_cargue_destino"] = cargue
        chofer = texto(fila["Chofer"])
        cid = id_camion(chofer)
        camion = camiones.setdefault(cid, {"id": cid, "chofer": chofer, "filas": []})
        camion["filas"].append(fila)

    for camion in camiones.values():
        completar_camion(camion)
    return [camiones[cid] for cid in sorted(camiones)]


def completar_camion(camion: dict) -> dict:
    camion["total"] = sum(f["Suma de Total"] for f in camion["filas"])
    # Los cargues son los de los clientes propios: un cliente pasado desde
    # otra ruta trae la liquidacion de esa otra ruta y no es un 2do cargue.
    propias = [f for f in camion["filas"] if not f.get("_movido_de")]
    camion["liquidaciones"] = sorted({f["Numero Liq"] for f in propias}, key=orden_valor)
    camion["movidos"] = sum(1 for f in camion["filas"] if f.get("_movido_de"))
    return camion


def separar_por_liquidacion(camion: dict):
    """Un mismo camion con 2 cargues (2 numeros de liquidacion) -> un
    'camion' por cargue, numerados por N de liquidacion de menor a mayor:
    'F325 CARGUE 1 - Liq 238', 'F325 CARGUE 2 - Liq 239'. Los clientes
    pasados desde otra ruta van en el cargue que indico la regla ('Hacia
    F325 CARGUE 2 - Liq 239') o, si no indico ninguno, en el CARGUE 1."""
    liquidaciones = camion["liquidaciones"]
    if len(liquidaciones) < 2:
        return [camion]
    texto_liqs = [str(liq) for liq in liquidaciones]

    def pieza_de(f):
        if f.get("_movido_de"):
            destino = str(f.get("_cargue_destino") or "")
            return texto_liqs.index(destino) if destino in texto_liqs else 0
        return texto_liqs.index(str(f["Numero Liq"])) if str(f["Numero Liq"]) in texto_liqs else 0

    piezas = []
    for i, liq in enumerate(liquidaciones):
        filas = [f for f in camion["filas"] if pieza_de(f) == i]
        piezas.append(completar_camion({"id": f"{camion['id']} CARGUE {i + 1} - Liq {liq}",
                                        "chofer": camion["chofer"], "filas": filas}))
    return piezas


def armar_seleccion(camiones, incluidos: dict, separados: dict):
    seleccion = []
    for camion in camiones:
        if not incluidos.get(camion["id"], True):
            continue
        if separados.get(camion["id"]):
            seleccion.extend(separar_por_liquidacion(camion))
        else:
            seleccion.append(camion)
    return seleccion


def escribir_excel(camion: dict, carpeta: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RUTA"
    # Si trae clientes pasados desde otra ruta, una columna extra al final
    # dice de cual (las 13 primeras quedan iguales que siempre).
    con_movidos = any(f.get("_movido_de") for f in camion["filas"])
    ws.append(COLUMNAS_ESPERADAS + (["Movido desde"] if con_movidos else []))
    for celda in ws[1]:
        celda.font = Font(bold=True)
    for fila in camion["filas"]:
        ws.append([fila[c] for c in COLUMNAS_ESPERADAS]
                  + ([fila.get("_movido_de", "")] if con_movidos else []))
    if con_movidos:
        ws.column_dimensions["N"].width = 14
    for fila in ws.iter_rows(min_row=2):
        fila[0].number_format = "dd/mm/yyyy"
        fila[11].number_format = '"$"#,##0.00'
        fila[12].number_format = '"$"#,##0.00'
    for letra, ancho in zip("ABCDEFGHIJKLM", [12, 24, 11, 10, 14, 36, 40, 12, 13, 13, 12, 13, 13]):
        ws.column_dimensions[letra].width = ancho
    ws.freeze_panes = "A2"
    ruta = carpeta / f"MAPA {camion['id']}.xlsx"
    wb.save(ruta)
    return ruta


# Ecuador continental + Galapagos, con margen. Lo que cae fuera (0,0, signo
# perdido, lat/lng cruzadas) no se publica: un solo punto asi aleja el mapa
# del chofer a escala de continente y le hace gastar datos.
LAT_MIN, LAT_MAX = -6.0, 2.5
LNG_MIN, LNG_MAX = -92.5, -75.0


def a_coordenada(valor):
    if es_numero(valor):
        return float(valor)
    if isinstance(valor, str):
        try:
            return float(valor.strip().replace(",", "."))
        except ValueError:
            return None
    return None


def en_ecuador(lat, lng) -> bool:
    return LAT_MIN <= lat <= LAT_MAX and LNG_MIN <= lng <= LNG_MAX


def armar_json(camiones, fechas, overrides=None):
    """Mismo formato que produce convertir.py desde los MAPA. Devuelve
    tambien los avisos de clientes que no pueden salir en el mapa.

    overrides: correcciones de ubicacion {codigo: {lat, lng}}. Un cliente
    sin coordenada valida en el sistema pero con correccion sale en el mapa
    en el punto corregido (para los demas, el mapa ya aplica la correccion
    solo al dibujar)."""
    overrides = overrides or {}
    salida = []
    sin_ubicacion = []
    avisos = {"sin_coordenadas": [], "fuera_de_ecuador": [], "camiones_vacios": []}
    for camion in camiones:
        clientes = []
        for fila in camion["filas"]:
            lat, lng = a_coordenada(fila["latitud_y"]), a_coordenada(fila["longitud_x"])
            problema = None
            if lat is None or lng is None:
                problema = "sin_coordenadas"
            elif not en_ecuador(lat, lng):
                problema = "fuera_de_ecuador"
            if problema:
                ajuste = overrides.get(str(fila["CodigoCliente"]).strip())
                if isinstance(ajuste, dict) and es_numero(ajuste.get("lat")) and es_numero(ajuste.get("lng")) \
                        and en_ecuador(ajuste["lat"], ajuste["lng"]):
                    lat, lng = float(ajuste["lat"]), float(ajuste["lng"])
                else:
                    avisos[problema].append((camion["id"], fila["CodigoCliente"]))
                    sin_ubicacion.append({
                        "codigo": fila["CodigoCliente"], "nombre": texto(fila["Razon"]), "ruta": camion["id"],
                        "vendedor": texto(fila["vendedor"]), "direccion": texto(fila["Direccion"]),
                        "motivo": "sin coordenadas" if problema == "sin_coordenadas"
                        else "coordenada en 0,0 o fuera de Ecuador"})
                    continue
            cliente = cliente_desde_fila({**fila, "latitud_y": lat, "longitud_x": lng})
            if problema:
                # Su lat/lng es la correccion (el sistema no trae una valida):
                # el Excel de control no debe tomarla como coordenada del sistema.
                cliente["sin_coord_sistema"] = True
            if fila.get("_movido_de"):
                cliente["movido_de"] = fila["_movido_de"]
            clientes.append(cliente)
        if not clientes:
            avisos["camiones_vacios"].append(camion["id"])
            continue
        asignar_colores(clientes)
        salida.append({"id": camion["id"], "chofer": camion["chofer"], "clientes": clientes})
    datos = {
        "fecha": max(fechas).strftime("%Y-%m-%d"),
        "generado": datetime.now().isoformat(timespec="seconds"),
        "camiones": salida,
    }
    if sin_ubicacion:
        datos["sin_ubicacion"] = sin_ubicacion
    return datos, avisos


def texto_avisos(avisos) -> str:
    """Resumen de avisos para mostrar al usuario ('' si no hay)."""
    lineas = []
    for clave, explicacion in (("sin_coordenadas", "sin coordenadas"),
                               ("fuera_de_ecuador", "con coordenadas en 0,0 o fuera de Ecuador")):
        casos = avisos[clave]
        if casos:
            ejemplos = ", ".join(f"{cid} cliente {codigo}" for cid, codigo in casos[:8])
            mas = f" y {len(casos) - 8} más" if len(casos) > 8 else ""
            lineas.append(f"Aviso: {len(casos)} clientes {explicacion} no salen en el mapa "
                          f"(sí están en los Excel): {ejemplos}{mas}.")
    if avisos["camiones_vacios"]:
        lineas.append(f"OJO: {', '.join(avisos['camiones_vacios'])} no tiene ningún cliente con "
                      "ubicación válida, así que NO sale en el mapa.")
    return "\n".join(lineas)


# ---------------------------------------------------------------------------
# Publicacion (lo mismo que publicar_datos.bat)
# ---------------------------------------------------------------------------

ERRORES_DE_RED = ("could not resolve host", "unable to access", "failed to connect",
                  "timed out", "connection was reset", "network is unreachable",
                  "could not read from remote")
NO_PUBLICADO = "No se publicó nada; el mapa sigue con la ruta anterior."


def correr_git(argumentos, log, tiempo=60):
    """Corre git en el repo de datos y devuelve (codigo, salida). Con
    tiempo limite para que una conexion colgada no deje la ventana en
    "Publicando..." para siempre."""
    try:
        r = subprocess.run(["git", "-c", "http.lowSpeedLimit=1000", "-c", "http.lowSpeedTime=30",
                            *argumentos],
                           cwd=REPO_DATOS, capture_output=True, text=True, encoding="utf-8",
                           errors="replace", timeout=tiempo,
                           creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    except subprocess.TimeoutExpired:
        raise ErrorRuta("GitHub no respondió a tiempo. Revisa la conexión a internet y vuelve a "
                        "darle a \"Generar y publicar\".")
    salida = ((r.stdout or "") + (r.stderr or "")).strip()
    if salida and log:
        log(salida)
    return r.returncode, salida


def es_error_de_red(salida: str) -> bool:
    salida = salida.lower()
    return any(t in salida for t in ERRORES_DE_RED)


def rebase_a_medias() -> bool:
    return (REPO_DATOS / ".git" / "rebase-merge").exists() or (REPO_DATOS / ".git" / "rebase-apply").exists()


def preparar_repo(log):
    """Deja el repo de datos limpio antes de publicar. Si una publicacion
    anterior se corto a medias, descarta sus restos (latest.json e
    historial se vuelven a generar ahora mismo); cualquier otro cambio no
    se toca y se avisa."""
    if (REPO_DATOS / ".git" / "index.lock").exists():
        raise ErrorRuta("Otro programa está usando la carpeta de datos (C:\\MAPA-RUTAS-DATOS), o se "
                        "quedó trabado un intento anterior.\nCierra cualquier otro programa que la use y "
                        f"vuelve a intentar. Si sigue igual, avisa (archivo .git\\index.lock).\n{NO_PUBLICADO}")
    if rebase_a_medias():
        log("Había una combinación anterior a medias; se deshace.")
        correr_git(["rebase", "--abort"], log)
    codigo, salida = correr_git(["status", "--porcelain"], None)
    if codigo != 0:
        raise ErrorRuta(f"No pude revisar la carpeta de datos (git status).\n{NO_PUBLICADO}")
    # Cada linea es "<estado> <ruta>" (ej. " M latest.json", "?? historial/x.json").
    sucios = [linea.strip().split(None, 1)[1].strip('"') for linea in salida.splitlines()
              if len(linea.strip().split(None, 1)) == 2]
    propios = ("latest.json", "overrides.json")
    ajenos = [p for p in sucios if p not in propios and not p.startswith("historial/")]
    if ajenos:
        raise ErrorRuta("La carpeta de datos tiene cambios que no hizo este programa: "
                        f"{', '.join(ajenos[:5])}.\n{NO_PUBLICADO} Avisa antes de reintentar.")
    if sucios:
        log("Se descartan restos de una publicación anterior que no terminó (se vuelven a generar).")
        rutas = [p for p in propios if (REPO_DATOS / p).exists()] + ["historial"]
        correr_git(["reset", "-q", "--", *rutas], log)
        correr_git(["checkout", "--", *rutas], log)
        correr_git(["clean", "-fdq", "--", "historial"], log)


def traer_de_github(log):
    """git pull --rebase. Las correcciones del modo admin del mapa se suben
    directo a GitHub, asi que hay que traerlas antes de subir nada. Si
    choca con un intento anterior que no llego a subirse (solo puede chocar
    en overrides.json), ese intento se descarta: todo lo que se sube se
    vuelve a armar ahora mismo sobre lo ultimo de GitHub."""
    codigo, salida = correr_git(["pull", "--rebase", "origin", "main"], log, tiempo=300)
    if codigo == 0:
        return
    chocaba = rebase_a_medias()
    if chocaba:
        correr_git(["rebase", "--abort"], log)
    if es_error_de_red(salida):
        raise ErrorRuta(f"No hay conexión con GitHub. Revisa internet y vuelve a intentar.\n{NO_PUBLICADO}")
    if chocaba:
        log("Un intento anterior chocaba con lo de GitHub; se descarta y se rehace sobre lo último.")
        if correr_git(["reset", "-q", "--hard", "origin/main"], log)[0] == 0:
            return
    raise ErrorRuta(f"No se pudo combinar con los cambios de GitHub.\n{NO_PUBLICADO} "
                    "Avisa antes de reintentar.")


def descartar_intentos_viejos(log, con_ruta):
    """Si un intento anterior se guardo (commit) pero no llego a GitHub (se
    corto el internet al subir), se descarta: lo que se sube se vuelve a
    armar ahora sobre lo ultimo de GitHub, y las ubicaciones del Excel siguen
    pendientes, asi que se comparan siempre contra lo mismo que la foto.
    Excepcion: si ese intento traia una ruta y ahora solo se aplican
    ubicaciones, no se descarta (se pide publicar la ruta primero)."""
    codigo, salida = correr_git(["fetch", "-q", "origin"], log, tiempo=300)
    if codigo != 0:
        if es_error_de_red(salida):
            raise ErrorRuta(f"No hay conexión con GitHub. Revisa internet y vuelve a intentar.\n{NO_PUBLICADO}")
        raise ErrorRuta(f"No se pudo consultar GitHub (git fetch).\n{NO_PUBLICADO} Avisa.")
    if not commits_sin_subir():
        return
    _, nombres = correr_git(["diff", "--name-only", "origin/main", "HEAD"], None)
    trae_ruta = any(n.strip() == "latest.json" or n.strip().startswith("historial/") for n in nombres.splitlines())
    if trae_ruta and not con_ruta:
        raise ErrorRuta("Hay una ruta de una publicación anterior que no llegó a subirse (se cortó el internet). "
                        "Dale primero a \"Generar y publicar en el mapa\"; las ubicaciones se aplican en esa misma "
                        "subida.\n(No se subió nada.)")
    log("Se descarta un intento anterior que no llegó a GitHub; todo se vuelve a armar ahora.")
    if correr_git(["reset", "-q", "--hard", "origin/main"], log)[0] != 0:
        raise ErrorRuta(f"No pude dejar la carpeta de datos igual a GitHub.\n{NO_PUBLICADO} Avisa.")


def commits_sin_subir() -> int:
    codigo, salida = correr_git(["rev-list", "--count", "origin/main..HEAD"], None)
    return int(salida) if codigo == 0 and salida.isdigit() else 1


def armar_commit(log, con_ruta, ubicaciones):
    """Deja en la carpeta de datos lo que hay que subir y lo guarda (commit).
    Se puede repetir sin problema (si ya estaba, no cambia nada).
    Devuelve los estados de las ubicaciones del Excel de control."""
    rutas, partes, estados = [], [], {}
    if con_ruta:
        log(archivar(REPO_DATOS, SALIDA_JSON))
        log(guardar_en_historial(REPO_DATOS, SALIDA_JSON))
        shutil.copyfile(SALIDA_JSON, REPO_DATOS / "latest.json")
        rutas += ["latest.json"] + (["historial"] if (REPO_DATOS / "historial").is_dir() else [])
        partes.append("Ruta del dia")
    if ubicaciones:
        import control_rutas
        estados, cambiado = control_rutas.aplicar_ubicaciones_en_archivo(REPO_DATOS / "overrides.json",
                                                                         ubicaciones)
        if cambiado:
            rutas.append("overrides.json")
            partes.append("Corregir ubicaciones (Excel de control)")
    if not rutas:
        return estados
    if correr_git(["add", *rutas], log)[0] != 0:
        raise ErrorRuta("No se pudo preparar la subida (git add). ¿Hay otro programa usando la "
                        f"carpeta de datos?\n{NO_PUBLICADO}")
    codigo, _ = correr_git(["diff", "--cached", "--quiet"], None)
    if codigo == 1:
        if correr_git(["commit", "-m", " + ".join(partes)], log)[0] != 0:
            raise ErrorRuta(f"No se pudo guardar el cambio (git commit).\n{NO_PUBLICADO} Avisa.")
    elif codigo != 0:
        raise ErrorRuta(f"No pude revisar los cambios (git diff).\n{NO_PUBLICADO}")
    return estados


def overrides_locales(cambios_pendientes=()):
    """Correcciones de ubicacion de la carpeta de datos, mas las que estan
    por aplicarse desde el Excel de control."""
    try:
        actual = json.loads((REPO_DATOS / "overrides.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        actual = {}
    for c in cambios_pendientes:
        if c["accion"] == "quitar":
            actual.pop(c["codigo"], None)
        else:
            actual[c["codigo"]] = {"lat": c["lat"], "lng": c["lng"]}
    return actual


def fecha_publicada():
    try:
        with open(REPO_DATOS / "latest.json", encoding="utf-8") as f:
            fecha = json.load(f).get("fecha")
        datetime.strptime(fecha, "%Y-%m-%d")
        return fecha
    except (OSError, ValueError, TypeError):
        return None


def publicar(log, con_ruta=True, ubicaciones=None):
    """Sube al repo privado la ruta (data/latest.json) y/o los cambios de
    ubicacion del Excel de control. Solo termina sin error si GitHub quedo
    con todo; si algo falla, lanza ErrorRuta (nunca dice "Listo" sin haber
    publicado). Devuelve los estados de las ubicaciones."""
    if not (REPO_DATOS / ".git").is_dir():
        raise ErrorRuta(f"No encuentro la carpeta del repo de datos ({REPO_DATOS}).")

    preparar_repo(log)
    log("Combinando con GitHub...")
    descartar_intentos_viejos(log, con_ruta)
    traer_de_github(log)

    for intento in (1, 2, 3):
        estados = armar_commit(log, con_ruta, ubicaciones)
        # Se sube tambien si no hubo commit nuevo, por si quedo uno de un
        # intento anterior que no llego a GitHub.
        if not commits_sin_subir():
            log("En GitHub ya estaba exactamente esto.")
            return estados
        log("Subiendo a GitHub...")
        codigo, salida = correr_git(["push", "origin", "main"], log, tiempo=300)
        if codigo == 0:
            if commits_sin_subir():
                raise ErrorRuta("GitHub no confirmó la subida. Vuelve a intentar.")
            return estados
        if es_error_de_red(salida):
            raise ErrorRuta("No hay conexión con GitHub. Revisa internet y vuelve a intentar "
                            "(quedó todo listo para subir).")
        # Rechazado: GitHub tiene algo mas nuevo (lo comun: justo se corrigio
        # un punto en el mapa). Se trae y se vuelve a armar encima.
        log("GitHub tenía un cambio más nuevo; se combina y se reintenta.")
        traer_de_github(log)
    raise ErrorRuta("No se pudo subir a GitHub. Vuelve a intentar; si sigue fallando, avisa.")


# ---------------------------------------------------------------------------
# Configuracion recordada entre usos (data/ no se sube a GitHub)
# ---------------------------------------------------------------------------

def cargar_config() -> dict:
    try:
        with open(ARCHIVO_CONFIG, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


def guardar_config(config: dict) -> None:
    ARCHIVO_CONFIG.parent.mkdir(exist_ok=True)
    ARCHIVO_CONFIG.write_text(json.dumps(config, ensure_ascii=False, indent=2), encoding="utf-8")


def buscar_base_reciente(config: dict):
    """El maestro fijo (BASE\\...) si existe. Si no, el archivo base mas
    reciente entre el ultimo usado y los de las carpetas del dia en
    OneDrive (MAPAS DIARIOS\\<año>\\<mes>\\<dia>)."""
    if BASE_FIJA.is_file():
        return BASE_FIJA
    candidatos = []
    if config.get("archivo_base"):
        candidatos.append(Path(config["archivo_base"]))
    if CARPETA_MAPAS_DIARIOS.is_dir():
        candidatos.extend(CARPETA_MAPAS_DIARIOS.glob(f"*/*/*/{NOMBRE_BASE}"))
    existentes = [c for c in candidatos if c.is_file()]
    return max(existentes, key=lambda c: c.stat().st_mtime) if existentes else None


def es_base_fija(archivo_base: Path) -> bool:
    try:
        return archivo_base.resolve() == BASE_FIJA.resolve()
    except OSError:
        return False


def carpeta_de_salida(archivo_base: Path, fechas) -> Path:
    """Donde van los MAPA: con el maestro fijo, una carpeta por fecha
    (BASE\\MAPAS\\2026-10-08 JUEVES); con otro archivo, junto a el (como
    se hacia con la carpeta del dia)."""
    if not es_base_fija(archivo_base):
        return archivo_base.parent
    ultima = max(fechas)
    return CARPETA_MAPAS_FIJA / f"{ultima:%Y-%m-%d} {DIAS_LARGOS[ultima.weekday()]}"


# Exactamente los nombres que escribe escribir_excel (distingue mayusculas:
# una copia del usuario como "Mapa F150 corregido.xlsx" nunca se toca).
RE_MAPA_GENERADO = re.compile(r"MAPA [A-Z]{1,3}\d{1,4}( CARGUE \d+ - Liq [0-9A-Za-z]+)?\.xlsx")


def revisar_cerrados(rutas) -> None:
    """Antes de escribir o borrar nada: si alguno esta abierto en Excel,
    avisar sin haber cambiado ningun archivo."""
    for ruta in rutas:
        if not ruta.exists():
            continue
        try:
            with open(ruta, "r+b"):
                pass
        except PermissionError:
            raise ErrorRuta(f"Cierra \"{ruta.name}\" en Excel y vuelve a intentar.\n"
                            "(No se cambió ningún archivo.)")


def mapas_viejos(carpeta: Path):
    """MAPA de una generacion anterior del mismo dia, solo en la carpeta
    por fecha del maestro fijo (que es exclusiva de este programa)."""
    if not carpeta.is_dir() or CARPETA_MAPAS_FIJA.resolve() not in carpeta.resolve().parents:
        return []
    return [p for p in carpeta.iterdir() if p.is_file() and RE_MAPA_GENERADO.fullmatch(p.name)]


# ---------------------------------------------------------------------------
# Ventana
# ---------------------------------------------------------------------------

def abrir_ventana():
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    try:
        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        pass

    class Ventana:
        def __init__(self, raiz):
            self.raiz = raiz
            self.cola = queue.Queue()
            self.config = cargar_config()
            self.archivo_base = None
            self.filas = []
            self.vars_fechas = {}
            self.camiones = []
            self.incluidos = {}
            self.separados = {}
            self.editables = {"ubicaciones": [], "cambios": [], "foto": None}
            self.reglas = []
            self.estados_cambios = []
            self.cambios_ubic = []
            self.estados_ubic = {}
            self.carpeta_ultima = None
            self.nota_sabado = False
            self.ocupado = False
            # Medidas pensadas a 100% de escala de Windows; se agrandan solas
            # si la pantalla esta al 125%/150% para que nada quede apretado.
            self.escala = max(1.0, raiz.winfo_fpixels("1i") / 96)
            self.construir()
            raiz.after(100, self.revisar_cola)

            base = buscar_base_reciente(self.config)
            if base:
                self.cargar_base(base)
            else:
                self.lbl_base.config(text="No encontré el archivo base.")
                self.lbl_base_info.config(text="Elige el archivo con \"Cambiar archivo...\".")
                self.ocupar(False)

        # -- armado de la ventana --------------------------------------
        def px(self, n):
            return int(n * self.escala)

        def construir(self):
            raiz = self.raiz
            raiz.title("Actualizar ruta del día")
            alto = min(self.px(800), raiz.winfo_screenheight() - self.px(70))
            raiz.geometry(f"{self.px(900)}x{alto}+{self.px(40)}+{self.px(10)}")
            raiz.minsize(self.px(760), self.px(600))

            estilo = ttk.Style()
            if "vista" in estilo.theme_names():
                estilo.theme_use("vista")
            estilo.configure("Treeview", rowheight=self.px(24))
            estilo.configure("Accion.TButton", font=("Segoe UI", 10, "bold"), padding=(12, 6))
            estilo.configure("Titulo.TLabel", font=("Segoe UI", 10, "bold"))
            estilo.configure("Ayuda.TLabel", foreground="#666666")

            cont = ttk.Frame(raiz, padding=12)
            cont.pack(fill="both", expand=True)

            m1 = ttk.LabelFrame(cont, text=" 1. Archivo base ", padding=8)
            m1.pack(fill="x")
            m1.columnconfigure(0, weight=1)
            self.lbl_base = ttk.Label(m1, text="", style="Titulo.TLabel")
            self.lbl_base.grid(row=0, column=0, sticky="w")
            self.lbl_base_info = ttk.Label(m1, text="", style="Ayuda.TLabel", wraplength=self.px(600))
            self.lbl_base_info.grid(row=1, column=0, sticky="w", pady=(2, 0))
            ttk.Button(m1, text="Abrir en Excel", command=self.abrir_base_en_excel).grid(
                row=0, column=1, rowspan=2, padx=(8, 0))
            self.btn_cambiar = ttk.Button(m1, text="Cambiar archivo...", command=self.elegir_base)
            self.btn_cambiar.grid(row=0, column=2, rowspan=2, padx=(6, 0))
            self.btn_releer = ttk.Button(m1, text="Volver a leer", command=self.releer)
            self.btn_releer.grid(row=0, column=3, rowspan=2, padx=(6, 0))

            m2 = ttk.LabelFrame(cont, text=" 2. Fecha de liquidación ", padding=8)
            m2.pack(fill="x", pady=(10, 0))
            self.marco_fechas = ttk.Frame(m2)
            self.marco_fechas.pack(fill="x")
            self.lbl_nota_fechas = ttk.Label(m2, text="", style="Ayuda.TLabel")
            self.lbl_nota_fechas.pack(anchor="w", pady=(4, 0))

            m3 = ttk.LabelFrame(cont, text=" 3. Camiones que van al mapa ", padding=8)
            m3.pack(fill="both", expand=True, pady=(10, 0))
            barra = ttk.Frame(m3)
            barra.pack(fill="x")
            self.btn_todos = ttk.Button(barra, text="Marcar todos", command=lambda: self.marcar_todos(True))
            self.btn_todos.pack(side="left")
            self.btn_ninguno = ttk.Button(barra, text="Desmarcar todos", command=lambda: self.marcar_todos(False))
            self.btn_ninguno.pack(side="left", padx=(6, 0))
            self.lbl_resumen = ttk.Label(barra, text="", style="Titulo.TLabel")
            self.lbl_resumen.pack(side="right")

            marco_lista = ttk.Frame(m3)
            marco_lista.pack(fill="both", expand=True, pady=(6, 0))
            columnas = ("incluir", "camion", "chofer", "liq", "clientes", "total", "cargues")
            self.lista = ttk.Treeview(marco_lista, columns=columnas, show="headings", selectmode="none")
            for col, titulo, ancho, ancla in [("incluir", "", 36, "center"), ("camion", "Camión", 64, "w"),
                                              ("chofer", "Chofer", 230, "w"), ("liq", "N° Liq.", 90, "w"),
                                              ("clientes", "Clientes", 64, "e"), ("total", "Total", 96, "e"),
                                              ("cargues", "Cargues", 130, "center")]:
                self.lista.heading(col, text=titulo, anchor=ancla)
                self.lista.column(col, width=self.px(ancho), anchor=ancla, stretch=(col == "chofer"))
            self.lista.tag_configure("fuera", foreground="#9a9a9a")
            self.lista.tag_configure("nuevo", background="#fff4cc")
            self.lista.tag_configure("varios", background="#e6f0ff")
            barra_v = ttk.Scrollbar(marco_lista, orient="vertical", command=self.lista.yview)
            self.lista.configure(yscrollcommand=barra_v.set)
            self.lista.pack(side="left", fill="both", expand=True)
            barra_v.pack(side="right", fill="y")
            self.lista.bind("<Button-1>", self.click_lista)
            # Sin esto, un doble clic marca y desmarca (parece que no hizo nada).
            self.lista.bind("<Double-Button-1>", lambda evento: "break")
            ttk.Label(m3, text="Toca un camión para marcarlo o desmarcarlo; lo que desmarques se recuerda "
                               "para la próxima vez (los camiones nuevos salen en amarillo). Los que tienen "
                               "2 cargues (2 N° de liquidación) salen en celeste: toca su columna \"Cargues\" "
                               "para elegir unificado o separados.",
                      style="Ayuda.TLabel", wraplength=self.px(820)).pack(anchor="w", pady=(6, 0))

            m4 = ttk.Frame(cont)
            m4.pack(fill="x", pady=(10, 0))
            self.btn_publicar = ttk.Button(m4, text="Generar y publicar en el mapa", style="Accion.TButton",
                                           command=lambda: self.generar(publicar_mapa=True))
            self.btn_publicar.pack(side="right")
            self.btn_excel = ttk.Button(m4, text="Solo generar Excel",
                                        command=lambda: self.generar(publicar_mapa=False))
            self.btn_excel.pack(side="right", padx=(0, 8))
            self.btn_carpeta = ttk.Button(m4, text="Carpeta de los MAPA", command=self.abrir_carpeta)
            self.btn_carpeta.pack(side="left")
            self.btn_control = ttk.Button(m4, text="Excel de control", command=self.abrir_control)
            self.btn_control.pack(side="left", padx=(8, 0))
            self.btn_ubicaciones = ttk.Button(m4, text="Aplicar ubicaciones", command=self.aplicar_ubicaciones)
            self.btn_ubicaciones.pack(side="left", padx=(8, 0))
            ttk.Button(m4, text="My Maps", command=self.abrir_mymaps).pack(side="left", padx=(8, 0))

            m5 = ttk.Frame(cont)
            m5.pack(fill="x", pady=(10, 0))
            self.barra_progreso = ttk.Progressbar(m5, mode="indeterminate", length=120)
            self.lbl_estado = ttk.Label(m5, text="", style="Ayuda.TLabel")
            self.lbl_estado.pack(side="left")
            self.txt = tk.Text(cont, height=5, wrap="word", font=("Consolas", 9), state="disabled",
                               background="#f6f6f6", relief="flat", padx=6, pady=4)
            self.txt.pack(fill="x", pady=(4, 0))

        # -- utilidades de la ventana ----------------------------------
        def log(self, mensaje):
            self.txt.config(state="normal")
            self.txt.insert("end", mensaje.rstrip() + "\n")
            self.txt.see("end")
            self.txt.config(state="disabled")

        def ocupar(self, ocupado, mensaje=""):
            estado = "disabled" if ocupado else "normal"
            for b in (self.btn_cambiar, self.btn_releer, self.btn_todos, self.btn_ninguno,
                      self.btn_publicar, self.btn_excel, self.btn_control, self.btn_ubicaciones):
                b.config(state=estado)
            for chk in self.marco_fechas.winfo_children():
                chk.config(state=estado)
            if not ocupado and not self.camiones:
                self.btn_publicar.config(state="disabled")
                self.btn_excel.config(state="disabled")
            self.lbl_estado.config(text=mensaje)
            if ocupado:
                self.barra_progreso.pack(side="right")
                self.barra_progreso.start(12)
                self.raiz.config(cursor="watch")
            else:
                self.barra_progreso.stop()
                self.barra_progreso.pack_forget()
                self.raiz.config(cursor="")
            self.ocupado = ocupado

        def en_segundo_plano(self, funcion, *argumentos):
            def correr():
                try:
                    funcion(*argumentos)
                except ErrorRuta as e:
                    self.cola.put(("error", str(e)))
                except PermissionError as e:
                    self.cola.put(("error", f"No pude guardar:\n{e.filename}\n\n¿Está abierto en Excel? "
                                            "Ciérralo y vuelve a intentar."))
                except Exception:
                    detalle = traceback.format_exc()
                    registrar_error(detalle)
                    self.cola.put(("error", "Algo falló inesperadamente:\n\n"
                                            + detalle.strip().splitlines()[-1]))
            threading.Thread(target=correr, daemon=True).start()

        def revisar_cola(self):
            # Un error al procesar un mensaje nunca debe cortar este ciclo (si
            # no, la ventana queda trabada en "Leyendo..." con todo apagado).
            try:
                while True:
                    try:
                        tipo, *datos = self.cola.get_nowait()
                    except queue.Empty:
                        break
                    try:
                        self.procesar_mensaje(tipo, datos)
                    except Exception:
                        detalle = traceback.format_exc()
                        registrar_error(detalle)
                        self.ocupar(False, "Hubo un problema, revisa el mensaje.")
                        messagebox.showerror("No se pudo completar", "Algo falló inesperadamente:\n\n"
                                             + detalle.strip().splitlines()[-1], parent=self.raiz)
            finally:
                self.raiz.after(100, self.revisar_cola)

        def procesar_mensaje(self, tipo, datos):
            if tipo == "log":
                self.log(datos[0])
            elif tipo == "leido":
                self.base_leida(*datos)
            elif tipo == "control":
                self.control_leido(*datos)
            elif tipo == "confirmar_ubicaciones":
                self.confirmar_ubicaciones(*datos)
            elif tipo == "abrir":
                self.ocupar(False, datos[1])
                try:
                    os.startfile(datos[0])
                except OSError as e:
                    messagebox.showerror("Excel de control", f"No pude abrir el archivo:\n{datos[0]}\n\n{e}",
                                         parent=self.raiz)
            elif tipo == "fin":
                self.ocupar(False, datos[0].splitlines()[0])
                self.log(datos[0])
                self.actualizar_nota_fechas()
                messagebox.showinfo("Listo", datos[0], parent=self.raiz)
                if len(datos) > 1 and datos[1]:
                    # Recien publicada la ruta: lo siguiente es subirla a My Maps.
                    self.abrir_mymaps(datos[1])
            elif tipo == "error":
                self.ocupar(False, "Hubo un problema, revisa el mensaje.")
                self.log("ERROR: " + datos[0])
                messagebox.showerror("No se pudo completar", datos[0], parent=self.raiz)

        # -- archivo base ----------------------------------------------
        def elegir_base(self):
            inicial = self.archivo_base.parent if self.archivo_base else CARPETA_MAPAS_DIARIOS
            ruta = filedialog.askopenfilename(parent=self.raiz, title="Elige el archivo base",
                                              initialdir=str(inicial),
                                              filetypes=[("Excel", "*.xlsx *.xlsm"), ("Todos", "*.*")])
            if ruta:
                self.cargar_base(Path(ruta))

        def releer(self):
            if self.archivo_base:
                self.cargar_base(self.archivo_base)

        def cargar_base(self, ruta: Path):
            self.archivo_base = ruta
            self.lbl_base.config(text=ruta.name)
            carpeta = ruta.parent
            if CARPETA_MAPAS_DIARIOS in carpeta.parents:
                carpeta = Path("…") / carpeta.relative_to(CARPETA_MAPAS_DIARIOS)
            try:
                guardado = datetime.fromtimestamp(ruta.stat().st_mtime)
                self.lbl_base_info.config(text=f"Carpeta: {carpeta}\nGuardado por última vez: "
                                               f"{formato_fecha(guardado.date())} {guardado:%H:%M}  "
                                               "— si acabas de pegar datos nuevos, guarda el Excel y dale "
                                               "a \"Volver a leer\".")
            except OSError:
                self.lbl_base_info.config(text=f"Carpeta: {carpeta}")
            # Nada de la lectura anterior debe quedar: si esta lectura falla,
            # no se puede publicar con datos viejos sin darse cuenta.
            self.filas = []
            self.camiones = []
            self.separados = {}
            self.vars_fechas = {}
            for chk in self.marco_fechas.winfo_children():
                chk.destroy()
            self.lbl_nota_fechas.config(text="")
            self.lbl_resumen.config(text="")
            self.lista.delete(*self.lista.get_children())
            self.ocupar(True, "Leyendo el archivo base...")
            self.en_segundo_plano(self._leer, ruta)

        def _leer(self, ruta):
            inicio = time.time()
            filas = leer_base(ruta)
            # Primero el Excel de control: asi, cuando se habilita "Publicar",
            # los cambios de ruta ya estan aplicados en la lista.
            self._leer_control()
            self.cola.put(("leido", filas, time.time() - inicio))

        def _leer_control(self):
            """Lo que el usuario escribio en el Excel de control (cambios de
            ruta y ubicaciones). Si falla, se sigue sin el."""
            self.cola.put(("control", self.leer_control_seguro()))

        def leer_control_seguro(self):
            """Lee el Excel de control; si no se puede (contraseña, daño,
            encabezado cambiado...), avisa y sigue SIN sus cambios, marcando
            el archivo como ilegible para no insistir en cada clic."""
            import control_rutas
            try:
                return control_rutas.leer_editables()
            except Exception as e:
                if not isinstance(e, ErrorRuta):
                    registrar_error(traceback.format_exc())
                motivo = str(e) if isinstance(e, ErrorRuta) else "no se pudo leer"
                self.cola.put(("log", f"Aviso: Excel de control: {motivo} Se sigue sin sus cambios."))
                return {"ubicaciones": [], "cambios": [], "foto": None, "ilegible": motivo,
                        "mtime": control_rutas.mtime_control()}

        def control_leido(self, editables):
            import control_rutas
            # Todo se calcula antes de guardarlo: si algo falla, queda lo
            # anterior entero (y no un estado a medias con la fecha nueva).
            reglas = control_rutas.reglas_desde(editables)
            cambios_ubic, estados_ubic = control_rutas.cambios_de_ubicacion(editables)
            self.editables, self.reglas = editables, reglas
            self.cambios_ubic, self.estados_ubic = cambios_ubic, estados_ubic
            activas = sum(1 for r in self.reglas if r["activa"])
            if activas or self.cambios_ubic:
                self.log(f"Excel de control: {activas} cambios de ruta activos, "
                         f"{len(self.cambios_ubic)} ubicaciones por aplicar.")
            if self.vars_fechas:
                self.refrescar_camiones()

        def base_leida(self, filas, segundos):
            self.filas = filas
            fechas = sorted({f["Fecha Liq"] for f in filas}, reverse=True)
            sugeridas = fechas_sugeridas(fechas)
            camiones_por_fecha = {}
            for f in filas:
                camiones_por_fecha.setdefault(f["Fecha Liq"], set()).add(id_camion(texto(f["Chofer"])))

            for chk in self.marco_fechas.winfo_children():
                chk.destroy()
            self.vars_fechas = {}
            for i, d in enumerate(fechas[:FECHAS_A_MOSTRAR]):
                var = tk.BooleanVar(value=d in sugeridas)
                self.vars_fechas[d] = var
                chk = ttk.Checkbutton(self.marco_fechas, variable=var, command=self.refrescar_camiones,
                                      text=f"{formato_fecha(d)}  ({len(camiones_por_fecha[d])} camiones)")
                chk.grid(row=i // 3, column=i % 3, sticky="w", padx=(0, 18), pady=1)

            self.nota_sabado = len(sugeridas) > 1
            self.actualizar_nota_fechas()

            self.log(f"Leídas {len(filas):,} filas de la base en {segundos:.0f} s.")
            self.refrescar_camiones()
            self.ocupar(False, "Revisa la fecha y los camiones, y dale a \"Generar y publicar en el mapa\".")

        def actualizar_nota_fechas(self):
            if not self.vars_fechas:
                return
            notas = []
            if self.nota_sabado:
                notas.append("Se juntó viernes + sábado porque la última fecha es sábado.")
            publicada = fecha_publicada()
            if publicada:
                d = datetime.strptime(publicada, "%Y-%m-%d").date()
                notas.append(f"En el mapa está publicada la ruta del {formato_fecha(d)}.")
            self.lbl_nota_fechas.config(text="  ".join(notas))

        # -- camiones ----------------------------------------------------
        def fechas_marcadas(self):
            return {d for d, var in self.vars_fechas.items() if var.get()}

        def refrescar_camiones(self):
            fechas = self.fechas_marcadas()
            filas, self.estados_cambios = self.filas, []
            if fechas and self.reglas:
                import control_rutas
                filas, self.estados_cambios = control_rutas.aplicar_movimientos(self.filas, self.reglas, fechas)
            self.camiones = agrupar(filas, fechas) if fechas else []
            excluidos = set(self.config.get("excluidos", []))
            # Si ya se publico esta fecha con algun camion en cargues
            # separados, se recuerda (asi una republicacion del mismo dia
            # mantiene los mismos nombres de ruta, ej. "F325 CARGUE 1 - Liq 238").
            separados_guardados = set(self.config.get("separados", {}).get(
                max(fechas).isoformat() if fechas else "", []))
            for camion in self.camiones:
                if camion["id"] not in self.incluidos:
                    self.incluidos[camion["id"]] = camion["id"] not in excluidos
                if camion["id"] not in self.separados:
                    self.separados[camion["id"]] = camion["id"] in separados_guardados
            self.lista.delete(*self.lista.get_children())
            for camion in self.camiones:
                self.lista.insert("", "end", iid=camion["id"])
            self.pintar_lista()
            if not self.ocupado:
                estado = "normal" if self.camiones else "disabled"
                self.btn_publicar.config(state=estado)
                self.btn_excel.config(state=estado)

        def pintar_lista(self):
            conocidos = set(self.config.get("conocidos", []))
            marcados = clientes = 0
            total = 0.0
            for camion in self.camiones:
                cid = camion["id"]
                dentro = self.incluidos[cid]
                nuevo = bool(conocidos) and cid not in conocidos
                varios = len(camion["liquidaciones"]) > 1
                nombre = (nombre_chofer(camion["chofer"]) + ("   (nuevo)" if nuevo else "")
                          + (f"   (+{camion['movidos']} de otra ruta)" if camion.get("movidos") else ""))
                if not varios:
                    cargues = "1"
                elif self.separados.get(cid):
                    cargues = f"{len(camion['liquidaciones'])} · SEPARADOS"
                else:
                    cargues = f"{len(camion['liquidaciones'])} · Unificado"
                etiquetas = ((() if dentro else ("fuera",)) + (("nuevo",) if nuevo else ())
                             + (("varios",) if varios else ()))
                self.lista.item(cid, tags=etiquetas, values=(
                    "☑" if dentro else "☐", cid, nombre,
                    ", ".join(str(liq) for liq in camion["liquidaciones"]),
                    len(camion["filas"]), f"${camion['total']:,.2f}", cargues))
                if dentro:
                    marcados += 1
                    clientes += len(camion["filas"])
                    total += camion["total"]
            texto_resumen = f"{marcados} de {len(self.camiones)} camiones · {clientes} clientes · ${total:,.2f}"
            aplicados = sum(1 for nivel, _ in self.estados_cambios if nivel == "ok")
            problemas = sum(1 for nivel, _ in self.estados_cambios if nivel == "error")
            if aplicados or problemas:
                texto_resumen += f" · {aplicados} cambios de ruta"
                if problemas:
                    texto_resumen += f" ({problemas} con problema)"
            self.lbl_resumen.config(text=texto_resumen)

        def click_lista(self, evento):
            if self.ocupado:
                return
            if self.lista.identify_region(evento.x, evento.y) != "cell":
                return
            cid = self.lista.identify_row(evento.y)
            if not cid:
                return
            columna = self.lista.identify_column(evento.x)
            camion = next(c for c in self.camiones if c["id"] == cid)
            if columna == "#7" and len(camion["liquidaciones"]) > 1:
                self.separados[cid] = not self.separados.get(cid, False)
            else:
                self.incluidos[cid] = not self.incluidos[cid]
            self.pintar_lista()

        def marcar_todos(self, valor):
            for camion in self.camiones:
                self.incluidos[camion["id"]] = valor
            self.pintar_lista()

        # -- generar / publicar -----------------------------------------
        def control_cambio(self) -> bool:
            """Si el Excel de control se guardo despues de leerlo (por ejemplo
            se agrego un cambio de ruta sin darle a "Volver a leer"), se vuelve
            a leer y se pide revisar la lista antes de seguir."""
            import control_rutas
            if self.editables.get("mtime") == control_rutas.mtime_control():
                return False
            editables = self.leer_control_seguro()
            self.control_leido(editables)
            if editables.get("ilegible"):
                messagebox.showwarning("Excel de control", f"El Excel de control cambió, pero: "
                                       f"{editables['ilegible']}\n\nSe sigue SIN sus cambios de ruta ni ubicaciones. "
                                       "Revisa la lista y vuelve a darle al botón.", parent=self.raiz)
            else:
                messagebox.showinfo("Excel de control", "El Excel de control cambió desde que se leyó, así que "
                                    "se volvió a leer con tus cambios.\n\nRevisa la lista de camiones y vuelve a "
                                    "darle al botón.", parent=self.raiz)
            return True

        def generar(self, publicar_mapa):
            if self.control_cambio():
                return
            seleccion = armar_seleccion(self.camiones, self.incluidos, self.separados)
            if not seleccion:
                messagebox.showwarning("Falta elegir", "Marca al menos un camión.", parent=self.raiz)
                return
            fechas = self.fechas_marcadas()
            if publicar_mapa:
                clientes = sum(len(c["filas"]) for c in seleccion)
                pregunta = (f"Se va a publicar en el mapa la ruta de:\n{formato_fechas(fechas)}\n\n"
                            f"{len(seleccion)} rutas, {clientes} clientes.")
                partidos = [c["id"] for c in self.camiones if self.incluidos[c["id"]]
                            and self.separados.get(c["id"]) and len(c["liquidaciones"]) > 1]
                if partidos:
                    pregunta += f"\nCargues separados: {', '.join(partidos)}."
                fuera = [c["id"] for c in self.camiones if not self.incluidos[c["id"]]]
                if fuera:
                    pregunta += f"\nNo van: {', '.join(fuera)}."
                aplicados = [t for nivel, t in self.estados_cambios if nivel == "ok"]
                problemas = [t for nivel, t in self.estados_cambios if nivel == "error"]
                if aplicados:
                    pregunta += f"\nCambios de ruta del Excel de control: {len(aplicados)}."
                if problemas:
                    pregunta += (f"\nOJO: {len(problemas)} cambios de ruta NO se pudieron aplicar: "
                                 + "; ".join(problemas[:3]) + ("..." if len(problemas) > 3 else "."))
                reciben_fuera = [c["id"] for c in self.camiones if c.get("movidos") and not self.incluidos[c["id"]]]
                if reciben_fuera:
                    pregunta += (f"\nOJO: {', '.join(reciben_fuera)} recibe clientes de otra ruta pero NO está "
                                 "marcado: esos clientes no saldrán en el mapa.")
                if self.cambios_ubic:
                    pregunta += f"\nUbicaciones del Excel de control por aplicar: {len(self.cambios_ubic)}."
                problemas_ubic = list(self.estados_ubic.values())
                if problemas_ubic:
                    pregunta += (f"\nOJO: {len(problemas_ubic)} filas de UBICACIONES tienen un problema y no se "
                                 "aplicarán: " + "; ".join(problemas_ubic[:2])
                                 + ("..." if len(problemas_ubic) > 2 else "."))
                if self.editables.get("ilegible"):
                    pregunta += ("\nOJO: no se pudo leer el Excel de control; se publica SIN sus cambios de ruta "
                                 "ni ubicaciones.")
                publicada = fecha_publicada()
                if publicada and max(fechas).isoformat() < publicada:
                    pregunta += (f"\n\nOJO: en el mapa ya está publicada una fecha más nueva ({publicada}). "
                                 "Esta la reemplazaría.")
                if not messagebox.askyesno("Publicar ruta", pregunta + "\n\n¿Continuar?", parent=self.raiz):
                    return

                # Lo elegido para el mapa se recuerda solo al publicar (no al
                # sacar un Excel suelto de un camion).
                excluidos = set(self.config.get("excluidos", []))
                for camion in self.camiones:
                    if self.incluidos[camion["id"]]:
                        excluidos.discard(camion["id"])
                    else:
                        excluidos.add(camion["id"])
                self.config["excluidos"] = sorted(excluidos)
                self.config["conocidos"] = sorted(set(self.config.get("conocidos", []))
                                                  | {c["id"] for c in self.camiones})
                separados = self.config.get("separados", {})
                separados[max(fechas).isoformat()] = sorted(
                    cid for cid, si in self.separados.items() if si)
                self.config["separados"] = {f: separados[f] for f in sorted(separados)[-15:]}
            self.config["archivo_base"] = str(self.archivo_base)
            guardar_config(self.config)

            carpeta = carpeta_de_salida(self.archivo_base, fechas)
            self.carpeta_ultima = carpeta
            self.ocupar(True, "Publicando..." if publicar_mapa else "Generando Excel...")
            self.en_segundo_plano(self._generar, seleccion, fechas, carpeta, publicar_mapa,
                                  list(self.cambios_ubic), self.editables,
                                  {i: t for i, (_, t) in enumerate(self.estados_cambios) if t})

        def _generar(self, seleccion, fechas, carpeta, publicar_mapa, cambios_ubic=(), editables=None,
                     estados_cambios=None):
            def log(mensaje):
                self.cola.put(("log", mensaje))

            carpeta.mkdir(parents=True, exist_ok=True)
            nuevos = [carpeta / f"MAPA {camion['id']}.xlsx" for camion in seleccion]
            # Al publicar, los MAPA de camiones que ya no van se quitan de la
            # carpeta del dia; con "Solo generar Excel" no se borra nada.
            viejos = [p for p in mapas_viejos(carpeta) if p not in nuevos] if publicar_mapa else []
            revisar_cerrados(nuevos + viejos)

            datos, avisos = armar_json(seleccion, fechas, overrides_locales(cambios_ubic))
            if publicar_mapa and not datos["camiones"]:
                raise ErrorRuta("Ningún camión tiene clientes con ubicación válida.\n" + NO_PUBLICADO)

            for viejo in viejos:
                viejo.unlink()
            for camion in seleccion:
                escribir_excel(camion, carpeta)
            log(f"{len(seleccion)} archivos MAPA guardados en {carpeta}")

            SALIDA_JSON.parent.mkdir(exist_ok=True)
            SALIDA_JSON.write_text(json.dumps(datos, ensure_ascii=False), encoding="utf-8")
            aviso = texto_avisos(avisos)
            if aviso:
                log(aviso)

            if publicar_mapa:
                estados_ubic = publicar(log, con_ruta=True, ubicaciones=cambios_ubic)
                fin = "Listo. Los choferes ya ven la ruta actualizada en el mapa."
                fin += self._actualizar_control(log, editables, estados_cambios, estados_ubic)
            else:
                fin = f"Listo. Se generaron {len(seleccion)} archivos MAPA en:\n{carpeta}"
            self.cola.put(("fin", fin + ("\n\n" + aviso if aviso else ""), carpeta if publicar_mapa else None))

        def _actualizar_control(self, log, editables=None, estados_cambios=None, estados_ubic=None):
            """Vuelve a armar el Excel de control despues de publicar. Devuelve
            una frase para el mensaje final ('' si salio bien)."""
            import control_rutas
            try:
                control_rutas.construir_control(log, editables, estados_cambios, estados_ubic)
            except ErrorRuta as e:
                log(f"Aviso: {e}")
                return f"\n\nEl Excel de control no se pudo actualizar: {e}"
            except Exception:
                registrar_error(traceback.format_exc())
                log("Aviso: no se pudo actualizar el Excel de control (detalle en data/generador_error.log).")
                return "\n\nOJO: el Excel de control no se pudo actualizar."
            finally:
                self._leer_control()
            # Lo que no se aplico: conflictos al aplicar + filas con problema
            # (enlace corto, coordenada mal pegada, codigo repetido...).
            problemas = [e for e in (estados_ubic or {}).values() if not e.startswith(control_rutas.RESUELTOS)]
            if editables:
                problemas += list(control_rutas.cambios_de_ubicacion(editables)[1].values())
            if problemas:
                return (f"\n\nOJO: {len(problemas)} filas de UBICACIONES no se aplicaron "
                        "(mira la columna Estado en el Excel de control).")
            return ""

        # -- Excel de control ---------------------------------------------
        def abrir_control(self):
            self.ocupar(True, "Armando el Excel de control...")
            self.en_segundo_plano(self._abrir_control)

        def _abrir_control(self):
            import control_rutas

            def log(mensaje):
                self.cola.put(("log", mensaje))

            try:
                ruta = control_rutas.construir_control(log)
                mensaje = "Excel de control listo."
            except ErrorRuta as e:
                if not control_rutas.ARCHIVO_CONTROL.exists():
                    raise
                ruta = control_rutas.ARCHIVO_CONTROL
                mensaje = f"Se abre sin actualizar: {e}"
                log(mensaje)
            self._leer_control()
            self.cola.put(("abrir", ruta, mensaje))

        def aplicar_ubicaciones(self):
            self.ocupar(True, "Leyendo el Excel de control...")
            self.en_segundo_plano(self._preparar_ubicaciones)

        def _preparar_ubicaciones(self):
            import control_rutas
            editables = control_rutas.leer_editables()
            cambios, estados = control_rutas.cambios_de_ubicacion(editables)
            self.cola.put(("confirmar_ubicaciones", editables, cambios, estados))

        def confirmar_ubicaciones(self, editables, cambios, estados):
            self.ocupar(False)
            self.control_leido(editables)
            errores = list(estados.values())
            if not cambios:
                texto_info = "No hay ubicaciones por aplicar en la hoja UBICACIONES del Excel de control."
                if errores:
                    texto_info += (f"\n\n{len(errores)} filas tienen un problema: " + "; ".join(errores[:3])
                                   + "\n(Abre el Excel de control para verlas.)")
                messagebox.showinfo("Aplicar ubicaciones", texto_info, parent=self.raiz)
                return
            poner = sum(1 for c in cambios if c["accion"] == "poner")
            quitar = len(cambios) - poner
            pregunta = (f"Se van a aplicar en el mapa {poner} ubicaciones corregidas"
                        + (f" y quitar {quitar}" if quitar else "") + ".\n"
                        "Los choferes las ven al volver a abrir el mapa (no hace falta volver a publicar la ruta).")
            if errores:
                pregunta += f"\n\nOJO: {len(errores)} filas tienen un problema y no se aplicarán."
            if not messagebox.askyesno("Aplicar ubicaciones", pregunta + "\n\n¿Continuar?", parent=self.raiz):
                return
            self.ocupar(True, "Aplicando ubicaciones...")
            self.en_segundo_plano(self._aplicar_ubicaciones, editables, cambios)

        def _aplicar_ubicaciones(self, editables, cambios):
            def log(mensaje):
                self.cola.put(("log", mensaje))

            estados = publicar(log, con_ruta=False, ubicaciones=cambios)
            aplicadas = sum(1 for e in estados.values() if e.startswith(("Aplicada", "Quitada")))
            fin = f"Listo. Se aplicaron {aplicadas} de {len(cambios)} cambios de ubicación."
            fin += self._actualizar_control(log, editables, None, estados)
            self.cola.put(("fin", fin))

        def abrir_carpeta(self):
            if self.carpeta_ultima and self.carpeta_ultima.is_dir():
                os.startfile(self.carpeta_ultima)
            elif self.archivo_base and es_base_fija(self.archivo_base):
                CARPETA_MAPAS_FIJA.mkdir(parents=True, exist_ok=True)
                os.startfile(CARPETA_MAPAS_FIJA)
            elif self.archivo_base:
                os.startfile(self.archivo_base.parent)

        def abrir_mymaps(self, carpeta=None):
            try:
                import ventana_mymaps
            except Exception:
                registrar_error(traceback.format_exc())
                messagebox.showerror("My Maps", "No se pudo abrir la ventana de My Maps (detalle en "
                                     "data/generador_error.log).", parent=self.raiz)
                return
            ventana_mymaps.abrir(self.raiz, carpeta or self.carpeta_ultima)

        def abrir_base_en_excel(self):
            if self.archivo_base and self.archivo_base.is_file():
                os.startfile(self.archivo_base)

    raiz = tk.Tk()

    def error_de_ventana(tipo, valor, tb):
        detalle = "".join(traceback.format_exception(tipo, valor, tb))
        registrar_error(detalle)
        messagebox.showerror("Error", detalle.strip().splitlines()[-1], parent=raiz)

    raiz.report_callback_exception = error_de_ventana
    Ventana(raiz)
    raiz.mainloop()


def registrar_error(detalle: str) -> None:
    try:
        ARCHIVO_ERRORES.parent.mkdir(exist_ok=True)
        with open(ARCHIVO_ERRORES, "a", encoding="utf-8") as f:
            f.write(f"--- {datetime.now():%Y-%m-%d %H:%M:%S}\n{detalle}\n")
    except OSError:
        pass


def prueba(archivo_base: Path, carpeta_salida: Path):
    inicio = time.time()
    filas = leer_base(archivo_base)
    print(f"Leidas {len(filas)} filas en {time.time() - inicio:.1f} s")
    fechas = sorted({f["Fecha Liq"] for f in filas})
    print("Ultimas fechas:", ", ".join(formato_fecha(d) for d in fechas[-FECHAS_A_MOSTRAR:]))
    sugeridas = fechas_sugeridas(fechas)
    print("Sugeridas:", formato_fechas(sugeridas))
    camiones = agrupar(filas, sugeridas)
    carpeta_salida.mkdir(parents=True, exist_ok=True)
    for camion in camiones:
        escribir_excel(camion, carpeta_salida)
        print(f"  {camion['id']:<6} {len(camion['filas']):>4} clientes  ${camion['total']:>10,.2f}")
    datos, avisos = armar_json(camiones, sugeridas)
    (carpeta_salida / "latest.json").write_text(json.dumps(datos, ensure_ascii=False), encoding="utf-8")
    print(texto_avisos(avisos) or "Sin avisos de coordenadas.")
    print(f"{len(camiones)} camiones, total {time.time() - inicio:.1f} s")


def main():
    if len(sys.argv) >= 2 and sys.argv[1] == "--prueba":
        prueba(Path(sys.argv[2]), Path(sys.argv[3]) if len(sys.argv) > 3 else Path(tempfile.mkdtemp()))
        return
    try:
        abrir_ventana()
    except Exception:
        registrar_error(traceback.format_exc())
        raise


if __name__ == "__main__":
    # Se usa a si mismo importado como 'generar_rutas', para que
    # control_rutas.py (que hace 'import generar_rutas') comparta este mismo
    # modulo y no una copia aparte (si no, sus errores no se reconocerian).
    import generar_rutas
    generar_rutas.main()
