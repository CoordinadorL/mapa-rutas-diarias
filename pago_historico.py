"""
Pago a transportes: importar a la web los cortes ANTERIORES a la web
(los Excel "CAMIONES-AYUDANTES-FRECUENCIAS - MARZO-ABRIL 2026.xlsx" ... que
se armaron a mano), para que las graficas de la pestana Cortes tengan
historia.

Solo LEE los Excel (con openpyxl, sobre una copia temporal: el original de
OneDrive no se abre ni se toca). Cada corte se sube a la web como CERRADO
(solo lectura) con los valores que tiene el Excel: valor de ruta y valor a
cancelar se guardan tal cual (la tarifa de ese momento), no se recalculan.

Reglas (los archivos viejos no tienen todos el mismo formato):
 - Pestanas de dia: las que terminan en "dd-mm" (LNS26-08, "LUNES 25-05",
   "MRLS 26-08"...). El encabezado se busca por titulo en las primeras
   filas (KGS o PESO, TARIFA TOTAL o RECONOCER TARIFA TOTAL = no sancionar...).
 - El corte sale del nombre del archivo (MARZO-ABRIL 2026 = 26/03 al 25/04).
 - Fila con fecha DESPUES del corte: es del corte siguiente (ya esta en el
   otro archivo), no se sube.
 - Fila con fecha ANTES del corte (un pago pendiente que se pago en este
   corte): se sube al primer dia del corte, con "fecha original dd/mm" en
   observaciones.
 - Codigo de camion que no es un codigo (".", "-", basura): no se sube.
 - Si las pestanas no tienen KGS (MARZO-ABRIL) se usa el CONSOLIDADO, que
   si lo tiene, siempre que tenga las mismas filas y totales.
 - Antes de subir se comparan los totales con el CONSOLIDADO del archivo:
   si no cuadran, no se deja importar (se muestra la diferencia).

Prueba (solo lee):  python pago_historico.py [carpeta]
"""
import re
import shutil
import sys
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path

import pago_excel
from pago_excel import corte_de, corte_de_archivo, normalizar

# Titulo normalizado del encabezado -> campo
CAMPOS = {
    "FECHA": "fecha", "AJUSTE": "ajuste", "LINEA": "linea", "COD DE FRECUENCIA": "cod_frec", "RUTA": "ruta",
    "VALOR DE RUTA": "valor_ruta", "CODIGO DE CAMION": "cod_camion", "NOMBRE": "nombre",
    "CHOFER INTERNO": "chofer_adicional", "CHOFER INTERNO / ADICIONAL": "chofer_adicional",
    "KGS": "kgs", "PESO": "kgs", "N° DE FACTURAS": "facturas", "N DE FACTURAS": "facturas", "NO DE FACTURAS": "facturas",
    "MONTO EN DOLARES $": "monto", "MONTO $": "monto", "RETORNO TOTAL EN $": "retorno", "RETORNO $": "retorno",
    "PAGO": "pago", "TARIFA TOTAL": "sin_sancion", "RECONOCER TARIFA TOTAL": "sin_sancion", "NO SANCIONAR": "sin_sancion",
    "VALOR A CANCELAR": "valor_cancelar", "NUMERO DE AYUDANTES": "n_ayudantes", "N° AYUDANTES": "n_ayudantes",
    "NOMBRES DE AYUDANTES": "ayudantes", "OBSERVACIONES": "observaciones",
}
RE_HOJA = re.compile(r"(\d{1,2})-(\d{1,2})\s*$")
RE_CODIGO = re.compile(r"^[A-Z]{1,3}\d{2,4}$")      # F150, FO51, FO01
SUMAS = ["monto", "retorno", "valor_ruta", "valor_cancelar"]


def num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return 0.0


def texto(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split())


def _columnas(fila):
    m = {}
    for i, v in enumerate(fila):
        k = CAMPOS.get(normalizar(v))
        if k and k not in m:
            m[k] = i
    return m


def _fecha(v, hoja, anio):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    m = RE_HOJA.search(hoja)
    if m:
        try:
            return date(anio, int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    return None


def _leer_hoja(ws, hoja, anio):
    """Filas de una pestana: (lista de dicts, columnas encontradas)."""
    filas = list(ws.iter_rows(values_only=True))
    enc = next((i for i, f in enumerate(filas[:10])
                if "cod_camion" in _columnas(f) and "fecha" in _columnas(f)), None)
    if enc is None:
        return [], set()
    col = _columnas(filas[enc])
    salida = []
    for n, f in enumerate(filas[enc + 1:], start=1):
        def g(k):
            i = col.get(k)
            return f[i] if i is not None and i < len(f) else None
        cod = texto(g("cod_camion")).upper()
        if not cod or cod in ("TOTAL", "TOTALES"):
            continue
        d = _fecha(g("fecha"), hoja, anio)
        o = {k: g(k) for k in col}
        o.update(fecha=d.isoformat() if d else "", cod_camion=cod, _hoja=hoja, _n=n)
        salida.append(o)
    return salida, set(col)


def _totales(filas):
    t = {"filas": len(filas), "dias": len({f["fecha"] for f in filas}),
         "transportistas": len({f["cod_camion"] for f in filas}),
         "kgs": round(sum(num(f.get("kgs")) for f in filas), 2)}
    for k in SUMAS:
        t[k] = round(sum(num(f.get(k)) for f in filas), 2)
    return t


def _dd_mm(iso):
    return f"{iso[8:10]}/{iso[5:7]}"


def leer_archivo(ruta: Path) -> dict:
    """Lee un Excel de corte viejo y deja listas las filas para la web.
    Devuelve {archivo, corte, filas, totales, consolidado, avisos, cuadra, problema}."""
    import openpyxl
    cid = corte_de_archivo(ruta)
    if not cid:
        raise ValueError(f"No se reconoce el corte en el nombre del archivo: {ruta.name}")
    corte = corte_de(date(int(cid[:4]), int(cid[5:7]), pago_excel.DIA_CORTE))
    desde, hasta = corte["desde"].isoformat(), corte["hasta"].isoformat()
    anio = corte["hasta"].year
    # Copia temporal: openpyxl lee sin abrir el original (que puede estar
    # abierto en Excel o sincronizandose con OneDrive).
    with tempfile.TemporaryDirectory() as tmp:
        copia = Path(tmp) / "corte.xlsx"
        shutil.copy2(ruta, copia)
        wb = openpyxl.load_workbook(copia, read_only=True, data_only=True)
        try:
            dias, cols_dias = [], set()
            for nombre in wb.sheetnames:
                m = RE_HOJA.search(nombre)
                if not m or normalizar(nombre) in ("BASE",):
                    continue
                fs, cols = _leer_hoja(wb[nombre], nombre, anio)
                dias += fs
                cols_dias |= cols
            cons = []
            if pago_excel.HOJA_CONSOLIDADO in wb.sheetnames:
                cons, _ = _leer_hoja(wb[pago_excel.HOJA_CONSOLIDADO], pago_excel.HOJA_CONSOLIDADO, anio)
        finally:
            wb.close()

    avisos = []
    basura = lambda f: not RE_CODIGO.match(f["cod_camion"])
    sin_fecha = lambda f: not f["fecha"]

    def limpiar(filas, donde):
        malas = [f for f in filas if basura(f)]
        if malas and donde:
            avisos.append(f"{len(malas)} fila(s) sin código de camión válido en {donde} (no se suben): "
                          + ", ".join(sorted({f['_hoja'] for f in malas})))
        filas = [f for f in filas if not basura(f)]
        nf = [f for f in filas if sin_fecha(f)]
        if nf and donde:
            avisos.append(f"{len(nf)} fila(s) sin fecha en {donde} (no se suben).")
        return [f for f in filas if not sin_fecha(f)]

    dias = limpiar(dias, "las pestañas")
    cons = limpiar(cons, "")
    despues = [f for f in dias if f["fecha"] > hasta]
    if despues:
        avisos.append(f"{len(despues)} fila(s) con fecha después del {_dd_mm(hasta)} (son del corte siguiente, "
                      f"ya están en ese archivo): {', '.join(sorted({_dd_mm(f['fecha']) for f in despues}))}.")
    dias = [f for f in dias if f["fecha"] <= hasta]
    cons = [f for f in cons if f["fecha"] <= hasta]

    # Sin KGS en las pestanas: se toman las filas del CONSOLIDADO (mismas
    # filas y totales, pero con el peso).
    if "kgs" not in cols_dias and cons and any(num(f.get("kgs")) for f in cons):
        a, b = _totales(dias), _totales(cons)
        if all(a[k] == b[k] for k in ["filas"] + SUMAS):
            avisos.append("Las pestañas no tienen KGS: se tomaron las filas del CONSOLIDADO (mismas filas y totales).")
            dias = cons
        else:
            avisos.append("Las pestañas no tienen KGS y el CONSOLIDADO no coincide: se suben sin peso.")

    # Pagos pendientes de antes del corte: al primer dia, con la fecha original.
    for f in dias:
        if f["fecha"] < desde:
            obs = texto(f.get("observaciones"))
            f["observaciones"] = (f"fecha original {_dd_mm(f['fecha'])}" + (" - " + obs if obs else ""))
            avisos.append(f"{f['cod_camion']} del {_dd_mm(f['fecha'])} ({f['_hoja']}): es de antes del corte, "
                          f"se pasa al {_dd_mm(desde)} con \"fecha original {_dd_mm(f['fecha'])}\".")
            f["fecha"] = desde

    tot = _totales(dias)
    tot_cons = _totales(cons) if cons else None
    problema = ""
    if tot_cons:
        dif = [f"{k}: pestañas {tot[k]:,.2f} / CONSOLIDADO {tot_cons[k]:,.2f}"
               for k in ["filas"] + SUMAS if abs(tot[k] - tot_cons[k]) > 0.01]
        if dif:
            problema = "Los totales no cuadran con el CONSOLIDADO: " + "; ".join(dif)
    else:
        avisos.append("El archivo no tiene CONSOLIDADO: no se pudieron comparar los totales.")

    filas = []
    for f in sorted(dias, key=lambda f: (f["fecha"], f["_hoja"], f["_n"])):
        ajuste = num(f.get("ajuste")) if texto(f.get("ajuste")) else ""
        filas.append({
            "fecha_entrega": f["fecha"], "orden": f["_n"],
            "clave_carga": f"HIST|{cid}|{f['_hoja']}|{f['_n']}",
            "cod_camion": f["cod_camion"], "nombre_carga": texto(f.get("nombre")),
            "ajuste": int(ajuste) if ajuste in (0, 1, 2, 3) else "",
            "linea": texto(f.get("linea")).upper(), "cod_frec": texto(f.get("cod_frec")).upper(),
            "chofer_adicional": texto(f.get("chofer_adicional")),
            "kgs": round(num(f.get("kgs")), 2), "facturas": num(f.get("facturas")),
            "monto": round(num(f.get("monto")), 2), "retorno": round(num(f.get("retorno")), 2),
            "sin_sancion": "SI" if normalizar(f.get("sin_sancion")) == "SI" else "",
            "n_ayudantes": num(f.get("n_ayudantes")) if texto(f.get("n_ayudantes")) else "",
            "ayudantes": texto(f.get("ayudantes")), "observaciones": texto(f.get("observaciones")),
            "ruta_hist": texto(f.get("ruta")), "valor_ruta_hist": round(num(f.get("valor_ruta")), 2),
            "valor_cancelar_hist": round(num(f.get("valor_cancelar")), 2),
        })
    return {"archivo": ruta.name, "ruta": str(ruta), "corte": {**corte, "desde": desde, "hasta": hasta},
            "filas": filas, "totales": tot, "consolidado": tot_cons, "avisos": avisos,
            "cuadra": not problema, "problema": problema}


def archivos_anteriores(carpeta=None):
    """Excel de cortes anteriores al primero que se sincroniza con la web
    (del mas viejo al mas nuevo, uno por corte)."""
    import pago_web
    config = pago_web.cargar_config()
    todos = pago_excel.archivos_bimestrales(carpeta)
    primer = config.get("primer_corte") or (corte_de_archivo(todos[0]) if todos else None)
    por_corte = {}
    for a in todos:                         # del mas nuevo al mas viejo: gana el mas nuevo
        cid = corte_de_archivo(a)
        if cid and (not primer or cid < primer) and cid not in por_corte:
            por_corte[cid] = a
    return [por_corte[c] for c in sorted(por_corte)]


def importar(leido: dict) -> dict:
    """Sube a la web un corte ya leido (leer_archivo)."""
    import pago_web
    if not leido["cuadra"]:
        raise pago_web.ErrorPago(leido["problema"])
    c = leido["corte"]
    return pago_web.llamar("pago_importar_corte", {
        "corte_id": c["id"], "archivo": leido["archivo"], "filas": leido["filas"],
    }, timeout=240)


if __name__ == "__main__":
    carpeta = Path(sys.argv[1]) if len(sys.argv) > 1 else None
    rutas = sorted(carpeta.glob(pago_excel.PATRON_ARCHIVO), key=corte_de_archivo) if carpeta else archivos_anteriores()
    for r in rutas:
        x = leer_archivo(r)
        t = x["totales"]
        print(f"{x['corte']['nombre']:<24} {t['dias']:>3} días {t['filas']:>4} filas {t['transportistas']:>3} transp  "
              f"monto {t['monto']:>13,.2f}  retorno {t['retorno']:>10,.2f}  kgs {t['kgs']:>12,.2f}  "
              f"v.ruta {t['valor_ruta']:>10,.2f}  a cancelar {t['valor_cancelar']:>10,.2f}  "
              f"{'OK' if x['cuadra'] else 'NO CUADRA'}")
        for a in x["avisos"]:
            print("     -", a)
        if x["problema"]:
            print("     !!", x["problema"])
