"""
Excel de control: BASE\\CONTROL DE RUTAS.xlsx (solo en esta PC; BASE\\ no se
sube a GitHub).

Hojas:
- RUTAS: detalle por cliente de los ultimos 60 dias publicados.
- RESUMEN: por fecha y ruta, de todo el historico.
- REVISAR UBICACIÓN: clientes con ubicacion sospechosa en la ruta vigente.
- UBICACIONES (editable): correcciones de ubicacion. Son las mismas que
  hace el arrastre del modo admin del mapa (overrides.json, globales por
  cliente): aqui se ven todas, se cambian, se agregan o se quitan.
- CAMBIOS DE RUTA (editable): pasar un cliente (o un camion entero) a la
  ruta de otro chofer, solo ese dia o siempre.

El historico se arma SIEMPRE con lo ya publicado (historial/*.json del repo
privado), nunca con el maestro, que se vacia cada dia. La ventana
(generar_rutas.py) aplica las hojas editables al publicar y luego vuelve a
armar este Excel con el Estado de cada fila.
"""
import json
import math
import os
import re
import statistics
import subprocess
from datetime import date, datetime, time, timedelta
from pathlib import Path
from urllib.parse import unquote
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import generar_rutas as g
from convertir import texto

ARCHIVO_CONTROL = g.CARPETA_BASE / "CONTROL DE RUTAS.xlsx"
RESPALDO_EDITABLES = g.CARPETA / "data" / "control_editables.json"
DIAS_DETALLE = 60
FILAS_EDITABLES = 1000          # filas con listas desplegables en las hojas editables
KM_LEJOS_DE_SU_RUTA = 25
M_CAMBIO_DEL_SISTEMA = 300
M_SISTEMA_YA_CORRIGIO = 50

HOJA_LEEME = "LEEME"
HOJA_RUTAS = "RUTAS"
HOJA_RESUMEN = "RESUMEN"
HOJA_REVISAR = "REVISAR UBICACIÓN"
HOJA_UBIC = "UBICACIONES"
HOJA_CAMBIOS = "CAMBIOS DE RUTA"
HOJA_LISTAS = "_LISTAS"
HOJA_FOTO = "_FOTO"

TIPO_DIA = "SOLO ESE DÍA"
TIPO_SIEMPRE = "SIEMPRE"
QUITAR = "QUITAR"

COLS_RUTAS = ["Fecha", "Día", "Estado", "Ruta", "Chofer", "Vendedor", "Código", "Cliente", "Dirección",
              "Canal", "Peso", "Venta", "Total", "Lat sistema", "Lng sistema", "Lat mapa", "Lng mapa",
              "Ubicación corregida", "Movido desde", "Ver en Maps"]
ANCHOS_RUTAS = [11, 6, 11, 22, 26, 10, 11, 34, 40, 14, 9, 10, 10, 11, 11, 11, 11, 11, 12, 9]
COLS_RESUMEN = ["Fecha", "Día", "Estado", "Ruta", "Chofer", "Clientes", "Vendedores", "Peso", "Venta",
                "Total", "Clientes movidos", "Con ubicación corregida"]
COLS_REVISAR = ["Motivo", "Ruta", "Código", "Cliente", "Dirección", "Vendedor",
                "Coordenada que usa el mapa", "Ver en Maps", "Tiene corrección"]
COLS_UBIC = ["Código", "Cliente", "Ruta (última vez)", "Coordenada del sistema", "Corrección vigente",
             "Nueva coordenada", "Acción", "Estado", "Ver sistema", "Ver corrección", "Distancia (m)", "Nota",
             "Dirección", "Vendedor"]
ANCHOS_UBIC = [11, 34, 22, 24, 24, 26, 10, 60, 9, 9, 11, 30, 40, 12]
EDITABLES_UBIC = {"Código", "Nueva coordenada", "Acción", "Nota"}
COLS_CAMBIOS = ["Tipo", "Fecha de la ruta", "Código", "Desde ruta", "Vendedor (opcional)", "Hacia ruta",
                "Activa", "Nota", "Estado", "Está hoy en"]
EDITABLES_CAMBIOS = {"Tipo", "Fecha de la ruta", "Código", "Desde ruta", "Vendedor (opcional)",
                     "Hacia ruta", "Activa", "Nota"}
# Columnas que definen una fila (si estan todas vacias, la fila no cuenta).
CLAVES_UBIC = ["Código", "Nueva coordenada", "Acción"]
CLAVES_CAMBIOS = ["Tipo", "Fecha de la ruta", "Código", "Desde ruta", "Vendedor (opcional)", "Hacia ruta"]

NEGRITA = Font(bold=True)
FONDO_TITULO = PatternFill("solid", fgColor="DDDDDD")
FONDO_EDITABLE = PatternFill("solid", fgColor="FFF2B3")
FONDO_SIN_GEO = PatternFill("solid", fgColor="F8CBAD")

# Filas que el programa pone arriba de UBICACIONES para cada cliente de la
# ruta que se va a publicar sin georreferencia. Se reconocen por el Estado:
# mientras no tengan 'Nueva coordenada' no son un error del usuario, y al
# volver a armar el Excel se regeneran desde la lista de ese momento.
MARCA_SIN_GEO = "SIN GEORREFERENCIA"
ESTADO_SIN_GEO = (f"{MARCA_SIN_GEO}: pega en 'Nueva coordenada' su ubicación (en Google Maps: clic derecho sobre "
                  "el local → clic en los números) y guarda")


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------

def normalizar_codigo(valor) -> str:
    """Codigo de cliente como texto ('1234567'), venga como numero o texto."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    s = str(valor).strip()
    if re.fullmatch(r"\d+\.0+", s):
        s = s.split(".")[0]
    return s


def id_base(valor) -> str:
    """'F325 CARGUE 1 - Liq 238', 'f325', 'F325-NOMBRE' -> 'F325'."""
    s = texto(valor)
    return g.id_camion(s) if s else ""


def distancia_m(lat1, lng1, lat2, lng2) -> float:
    r = 6371000.0
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp, dl = p2 - p1, math.radians(lng2 - lng1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * r * math.asin(math.sqrt(min(1.0, a)))


def en_ecuador(lat, lng) -> bool:
    return g.LAT_MIN <= lat <= g.LAT_MAX and g.LNG_MIN <= lng <= g.LNG_MAX


def enlace_maps(lat, lng) -> str:
    return f"https://www.google.com/maps?q={lat:.6f},{lng:.6f}"


def formula_enlace(lat, lng, rotulo="Ver"):
    return f'=HYPERLINK("{enlace_maps(lat, lng)}","{rotulo}")'


def texto_coord(lat, lng) -> str:
    return f"{lat:.6f}, {lng:.6f}"


_NUM = r"(-?\d{1,3}(?:\.\d+)?)"


# Solo los caracteres de un enlace corto de verdad: asi no se arrastra una
# comilla tipografica o un parentesis pegado al final al copiarlo.
RE_ENLACE_CORTO = re.compile(r"(?<![\w.])(?:https?://)?(?:maps\.app\.goo\.gl|goo\.gl/maps)/[A-Za-z0-9_-]+"
                             r"(?:\?[A-Za-z0-9_=&%.~-]*)?", re.IGNORECASE)
_ENLACES_RESUELTOS = {}
# Enlaces que de por si no sirven (Google dice que no existen): no se reintentan.
_ENLACES_INVALIDOS = {}
# Enlaces que no se pudieron abrir (sin internet, Google no respondio):
# {corto: (cuando, clave de la lectura, motivo)}. La misma lectura (el mismo
# guardado del Excel) no los reintenta durante SEGUNDOS_FALLIDO, para no
# esperar el tiempo de espera en cada vuelta.
_ENLACES_FALLIDOS = {}
SEGUNDOS_FALLIDO = 60
NO_LLEVA = "ese enlace corto no lleva a un punto de Google Maps"
MAL_COPIADO = "el enlace corto está mal copiado; cópialo otra vez desde el celular o pega los números"
SIN_INTERNET = ("no pude abrir el enlace corto (¿hay internet?); también puedes hacer clic "
                "derecho sobre el punto en Google Maps y copiar los números")
GOOGLE_NO_RESPONDE = ("Google no respondió al abrir el enlace corto ({}); intenta de nuevo en un momento "
                      "o pega los números")
NAVEGADOR = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
             "Chrome/130.0 Safari/537.36")


def es_enlace_corto(valor) -> bool:
    # Tambien dentro de otro enlace codificado (Safe Links de Outlook, l.facebook.com...).
    return bool(RE_ENLACE_CORTO.search(unquote(texto(valor))))


def resolver_enlace_corto(valor, tiempo=12, abrir=True, recordar_fallo=None) -> str:
    """El enlace largo al que lleva un enlace corto de Google Maps (el que se
    comparte desde el celular: maps.app.goo.gl/...). Solo sigue las
    redirecciones, sin descargar la pagina. Necesita internet; si no se puede,
    ValueError con un mensaje para el usuario. Se recuerda lo ya resuelto.
    abrir=False: solo lo ya sabido (no sale a internet; para no trabar la
    ventana). recordar_fallo: clave de la lectura (la fecha del Excel): si
    con esa misma clave ya fallo la conexion hace poco, no se reintenta."""
    import urllib.error
    import urllib.request
    from urllib.parse import parse_qs, urljoin, urlparse

    m = RE_ENLACE_CORTO.search(unquote(texto(valor)))
    if not m:
        return texto(valor)
    corto = m.group(0).rstrip(".,;)")
    if not corto.lower().startswith("http"):
        corto = "https://" + corto
    if corto in _ENLACES_RESUELTOS:
        return _ENLACES_RESUELTOS[corto]
    if corto in _ENLACES_INVALIDOS:
        raise ValueError(_ENLACES_INVALIDOS[corto])
    anterior = _ENLACES_FALLIDOS.get(corto)
    if not abrir:
        if anterior:
            raise ValueError(f"{anterior[2]} (al guardar otra vez el Excel se vuelve a intentar)")
        raise ValueError("todavía no se abrió el enlace corto (se abre solo en unos segundos si hay internet); "
                         "vuelve a intentar en un momento, o pega los números (clic derecho sobre el punto en "
                         "Google Maps)")
    ahora = datetime.now().timestamp()
    if (recordar_fallo is not None and anterior and anterior[1] == recordar_fallo
            and ahora - anterior[0] < SEGUNDOS_FALLIDO):
        raise ValueError(anterior[2])

    def fallo(motivo, definitivo=False):
        if definitivo:
            _ENLACES_INVALIDOS[corto] = motivo
        else:
            _ENLACES_FALLIDOS[corto] = (ahora, recordar_fallo, motivo)
        return ValueError(motivo)

    class _NoSeguir(urllib.request.HTTPRedirectHandler):
        def redirect_request(self, *args, **kwargs):
            return None

    abridor = urllib.request.build_opener(_NoSeguir)
    url = corto
    for _ in range(6):
        pedido = urllib.request.Request(url, headers={"User-Agent": NAVEGADOR, "Accept-Language": "es"})
        try:
            with abridor.open(pedido, timeout=tiempo):
                break                                   # ya no redirige: es el destino
        except urllib.error.HTTPError as e:
            if e.code in (301, 302, 303, 307, 308) and e.headers.get("Location"):
                url = urljoin(url, e.headers["Location"])
            elif e.code in (400, 404, 410):
                raise fallo(NO_LLEVA, definitivo=True)
            else:
                # 429 (demasiadas consultas), 5xx, 403 de un proxy...: pasajero.
                raise fallo(GOOGLE_NO_RESPONDE.format(f"código {e.code}"))
        except ValueError:
            # Caracteres que no van en una direccion web (enlace mal copiado).
            raise fallo(MAL_COPIADO, definitivo=True)
        except Exception:
            raise fallo(SIN_INTERNET)
        # Pagina de consentimiento de Google: el destino va en ?continue=
        if "consent." in urlparse(url).netloc:
            seguir = parse_qs(urlparse(url).query).get("continue")
            if seguir:
                url = seguir[0]
                break
        if not es_enlace_corto(url):
            break
    _ENLACES_RESUELTOS[corto] = url
    _ENLACES_FALLIDOS.pop(corto, None)
    return url


def parsear_coordenada(valor, abrir_enlaces=True):
    """(lat, lng) desde lo que se pegue: '-2.90012, -79.00456' (clic derecho
    en Google Maps -> copiar), '-2,90012; -79,00456', un enlace largo de
    Google Maps o uno corto de los que comparte el celular (se abre para ver
    a donde lleva; con abrir_enlaces=False solo si ya se abrio antes).
    Lanza ValueError con un mensaje para el usuario."""
    s = unquote(texto(valor))
    if not s:
        raise ValueError("vacía")
    if es_enlace_corto(s):
        s = unquote(resolver_enlace_corto(s, abrir=abrir_enlaces))
        if es_enlace_corto(s):
            raise ValueError("ese enlace corto no lleva a un punto de Google Maps")
    par = None
    patrones = [rf"!3d{_NUM}!4d{_NUM}", rf"[?&](?:q|query|ll|destination|daddr)={_NUM}\s*,\s*{_NUM}",
                # Pin compartido / lugar por coordenadas: la coordenada va en la ruta.
                rf"/maps/(?:search|place)/{_NUM}\s*,\s*\+?\s*{_NUM}"]
    # En un enlace de indicaciones (/dir/) o de busqueda (/search/) lo que va
    # despues de @ es el centro de la pantalla, no el cliente.
    if "/dir/" not in s and "/search/" not in s:
        patrones.append(rf"@{_NUM},{_NUM}")
    for patron in patrones:
        m = re.search(patron, s)
        if m:
            par = (float(m.group(1)), float(m.group(2)))
            break
    if par is None and ("/dir/" in s or "/search/" in s):
        raise ValueError("ese enlace es de indicaciones o de una búsqueda y no dice dónde está el cliente; "
                         "en Google Maps haz clic derecho sobre el punto exacto y copia los números")
    if par is None:
        m = re.fullmatch(r"\s*(-?\d+),(\d+)\s*[;\s]\s*(-?\d+),(\d+)\s*", s)       # decimales con coma
        if m:
            par = (float(f"{m.group(1)}.{m.group(2)}"), float(f"{m.group(3)}.{m.group(4)}"))
    if par is None:
        m = re.fullmatch(rf"\s*\(?\s*{_NUM}\s*[,;\s]\s*{_NUM}\s*\)?\s*", s)
        if m:
            par = (float(m.group(1)), float(m.group(2)))
    if par is None:
        raise ValueError("no entiendo esa coordenada; pégala como '-2.900123, -79.004567'")
    lat, lng = par
    if not en_ecuador(lat, lng) and en_ecuador(lng, lat):
        lat, lng = lng, lat                     # venian al reves (lng, lat)
    if not en_ecuador(lat, lng):
        raise ValueError("queda fuera de Ecuador")
    return lat, lng


def a_fecha(valor):
    return g.a_fecha(valor, date(1899, 12, 30))


# ---------------------------------------------------------------------------
# Lo publicado (repo privado), sin tocar los archivos de la carpeta
# ---------------------------------------------------------------------------

def leer_blobs(ref: str, rutas):
    """Contenido de varios archivos del repo en `ref` con una sola llamada a git."""
    if not rutas:
        return {}
    entrada = "".join(f"{ref}:{p}\n" for p in rutas).encode("utf-8")
    try:
        r = subprocess.run(["git", "cat-file", "--batch"], cwd=g.REPO_DATOS, input=entrada,
                           capture_output=True, timeout=120,
                           creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    except (OSError, subprocess.TimeoutExpired):
        raise g.ErrorRuta("No pude leer lo publicado desde la carpeta de datos (git).")
    datos, pos, salida = r.stdout, 0, {}
    for ruta in rutas:
        fin = datos.find(b"\n", pos)
        if fin < 0:
            break
        cabecera = datos[pos:fin].split()
        pos = fin + 1
        if len(cabecera) == 3 and cabecera[1] == b"blob":
            tam = int(cabecera[2])
            salida[ruta] = datos[pos:pos + tam]
            pos += tam + 1
    return salida


def datos_del_repo(log=None) -> dict:
    """Todas las rutas publicadas + las correcciones de ubicacion vigentes."""
    sin_red = False
    try:
        with g.BLOQUEO_GIT:
            if g.correr_git(["fetch", "-q", "origin"], None, tiempo=120)[0] != 0:
                sin_red = True
    except g.ErrorRuta:
        sin_red = True
    if sin_red and log:
        log("Sin conexión con GitHub: el Excel de control usa lo último descargado.")
    ref = "origin/main"
    codigo, salida = g.correr_git(["ls-tree", "--name-only", ref, "historial/"], None)
    if codigo != 0:
        ref = "HEAD"
        codigo, salida = g.correr_git(["ls-tree", "--name-only", ref, "historial/"], None)
    nombres = [n for n in salida.splitlines() if re.fullmatch(r"historial/\d{4}-\d{2}-\d{2}\.json", n.strip())]
    blobs = leer_blobs(ref, ["latest.json", "overrides.json", *nombres])
    rutas = {}
    for n in nombres:
        try:
            d = json.loads(blobs[n])
            rutas[d.get("fecha") or n[10:20]] = d
        except (KeyError, ValueError):
            pass
    vigente = None
    if "latest.json" in blobs:
        try:
            latest = json.loads(blobs["latest.json"])
            vigente = latest.get("fecha")
            if vigente:
                rutas[vigente] = latest
        except ValueError:
            pass
    try:
        overrides = json.loads(blobs.get("overrides.json") or b"{}")
    except ValueError:
        overrides = {}
    return {"rutas": rutas, "vigente": vigente, "overrides": overrides, "sin_red": sin_red}


# ---------------------------------------------------------------------------
# Hojas editables: lectura, respaldo y reglas
# ---------------------------------------------------------------------------

def _filas_de_hoja(ws, nombre_hoja, columnas, claves):
    """Filas como dict por encabezado. El encabezado se busca por nombre en
    las primeras filas (por si se inserto un titulo arriba). Una fila cuenta
    solo si tiene algo en `claves` (las columnas que escribe el usuario):
    asi una fila cuyas celdas amarillas se vaciaron con Supr no queda como
    regla fantasma. Si falta una columna clave, ErrorRuta (y no se toca el
    Excel ni el respaldo)."""
    filas_hoja = ws.iter_rows(values_only=True)
    posicion, faltan = {}, list(claves)
    for _ in range(10):
        try:
            encabezado = next(filas_hoja)
        except StopIteration:
            break
        nombres = {g.normalizar(h): i for i, h in enumerate(encabezado) if texto(h)}
        posicion = {c: nombres[g.normalizar(c)] for c in columnas if g.normalizar(c) in nombres}
        faltan = [c for c in claves if c not in posicion]
        if not faltan:
            break
    if faltan:
        raise g.ErrorRuta(f"En la hoja {nombre_hoja} del Excel de control no encuentro la columna "
                          f"'{faltan[0]}' (¿se cambió o se borró el encabezado?). No se tocó nada: "
                          "vuelve a poner el encabezado como estaba.")
    filas = []
    for fila in filas_hoja:
        d = {c: (fila[i] if i < len(fila) else None) for c, i in posicion.items()}
        if any(texto(d.get(c)) for c in claves):
            filas.append(d)
    return filas


def _serializable(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, (date, time, timedelta)):
        return v.isoformat() if hasattr(v, "isoformat") else str(v)
    return v


COLUMNAS_CON_COORDENADA = ("Nueva coordenada", "Corrección vigente", "Coordenada del sistema")


def enlaces_sin_abrir(editables: dict) -> bool:
    """Si quedaron enlaces cortos sin abrir en UBICACIONES."""
    return any(es_enlace_corto(f.get(c)) for f in (editables or {}).get("ubicaciones", [])
               for c in COLUMNAS_CON_COORDENADA)


def leer_editables(ruta: Path = None, abrir_enlaces=True) -> dict:
    """Lo que el usuario escribio en UBICACIONES y CAMBIOS DE RUTA, y la
    foto de las correcciones con que se armo el Excel. Si el Excel no
    existe, se toma del respaldo (data/control_editables.json). Si existe
    pero no se puede leer, ErrorRuta con un mensaje para el usuario.
    abrir_enlaces=False: los enlaces cortos solo se cambian si ya se
    abrieron antes (no sale a internet: para leer desde la ventana)."""
    ruta = ruta or ARCHIVO_CONTROL
    vacio = {"ubicaciones": [], "cambios": [], "foto": None, "origen": "nuevo"}
    if not ruta.exists():
        try:
            with open(RESPALDO_EDITABLES, encoding="utf-8") as f:
                respaldo = json.load(f)
            respaldo["origen"] = "respaldo"
            return respaldo
        except (OSError, ValueError):
            return vacio
    # La fecha de modificacion se toma ANTES de leer: si se vuelve a guardar
    # mientras se lee, queda distinta y se vuelve a leer (si no, el segundo
    # guardado quedaba como "ya leido" con el contenido del primero).
    mtime = mtime_control(ruta)
    copia = g.copiar_sin_bloqueo(ruta)
    try:
        try:
            wb = openpyxl.load_workbook(copia, read_only=True, data_only=True)
        except Exception:
            raise g.ErrorRuta("No pude abrir el Excel de control (¿tiene contraseña o está dañado?). "
                              "Quítale la contraseña, o cámbiale el nombre para que se vuelva a armar uno nuevo "
                              "(tus cambios de ruta se recuperan del respaldo).")
        try:
            hojas = {g.normalizar(n): n for n in wb.sheetnames}
            for hoja in (HOJA_UBIC, HOJA_CAMBIOS):
                if g.normalizar(hoja) not in hojas:
                    raise g.ErrorRuta(f"En el Excel de control no encuentro la hoja '{hoja}' (¿se le cambió el "
                                      "nombre o se borró?). No se tocó nada: vuelve a ponerle ese nombre.")
            ubic = _filas_de_hoja(wb[hojas[g.normalizar(HOJA_UBIC)]], HOJA_UBIC, COLS_UBIC, CLAVES_UBIC)
            cambios = _filas_de_hoja(wb[hojas[g.normalizar(HOJA_CAMBIOS)]], HOJA_CAMBIOS, COLS_CAMBIOS,
                                     CLAVES_CAMBIOS)
            foto = None
            if g.normalizar(HOJA_FOTO) in hojas:
                foto = {}
                for fila in wb[hojas[g.normalizar(HOJA_FOTO)]].iter_rows(min_row=2, values_only=True):
                    if fila and fila[0] not in (None, "") and len(fila) >= 3 \
                            and g.es_numero(fila[1]) and g.es_numero(fila[2]):
                        foto[normalizar_codigo(fila[0])] = [float(fila[1]), float(fila[2])]
        finally:
            wb.close()
    finally:
        try:
            copia.unlink()
        except OSError:
            pass
    # Enlaces cortos de Google Maps (los que comparte el celular): se abren
    # aqui, que casi siempre corre en segundo plano, para no trabar la ventana.
    for f in ubic:
        for columna in COLUMNAS_CON_COORDENADA:
            if es_enlace_corto(f.get(columna)):
                try:
                    f[columna] = resolver_enlace_corto(f.get(columna), abrir=abrir_enlaces, recordar_fallo=mtime)
                except ValueError:
                    pass                        # se avisa al leer la coordenada
    editables = {"ubicaciones": ubic, "cambios": cambios, "foto": foto, "origen": "excel", "mtime": mtime}
    guardar_respaldo(editables)
    return editables


def mtime_control(ruta: Path = None):
    """Fecha de modificacion del Excel de control (None si no existe)."""
    try:
        return (ruta or ARCHIVO_CONTROL).stat().st_mtime
    except OSError:
        return None


def guardar_respaldo(editables: dict) -> None:
    try:
        RESPALDO_EDITABLES.parent.mkdir(exist_ok=True)
        limpio = {k: v for k, v in editables.items() if k not in ("origen", "mtime", "ilegible")}
        limpio["ubicaciones"] = [{c: _serializable(v) for c, v in f.items()} for f in limpio.get("ubicaciones", [])]
        limpio["cambios"] = [{c: _serializable(v) for c, v in f.items()} for f in limpio.get("cambios", [])]
        # Atomico: la ventana tambien lo escribe al releer el Excel en segundo plano.
        temporal = RESPALDO_EDITABLES.with_name(f"{RESPALDO_EDITABLES.stem}.{os.getpid()}.{id(limpio)}.tmp")
        temporal.write_text(json.dumps(limpio, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
        os.replace(temporal, RESPALDO_EDITABLES)
    except Exception:
        pass


RE_LIQ = re.compile(r"\bLiq\.?\s*([0-9A-Za-z]+)", re.IGNORECASE)


def ruta_valida(valor) -> bool:
    """El texto empieza con un codigo de camion: 'F150', 'FO53', 'F150-NOMBRE',
    'F325 CARGUE 2 - Liq 239'. Un nombre de chofer o 'F2OO' (con letras O)
    no lo son."""
    s = texto(valor)
    m = g.RE_CAMION.match(s)
    return bool(m) and (m.end() == len(s) or s[m.end()] in " -")


def _liquidacion_en(valor) -> str:
    """'F325 CARGUE 2 - Liq 239' -> '239' (cuando se indica un cargue)."""
    m = RE_LIQ.search(texto(valor))
    return m.group(1) if m else ""


def reglas_desde(editables: dict):
    """Las filas de CAMBIOS DE RUTA como reglas (una por fila, mismo orden)."""
    reglas = []
    for i, f in enumerate(editables.get("cambios", [])):
        tipo = g.normalizar(f.get("Tipo"))
        reglas.append({
            "fila": i,
            "tipo": {"SOLO ESE DIA": TIPO_DIA, "SIEMPRE": TIPO_SIEMPRE}.get(tipo, tipo),
            "fecha": a_fecha(f.get("Fecha de la ruta")),
            "fecha_texto": texto(f.get("Fecha de la ruta")),
            "codigo": "TODOS" if g.normalizar(f.get("Código")) == "TODOS" else normalizar_codigo(f.get("Código")),
            "desde": id_base(f.get("Desde ruta")),
            "desde_texto": texto(f.get("Desde ruta")),
            "desde_liq": _liquidacion_en(f.get("Desde ruta")),
            "vendedor": g.normalizar(f.get("Vendedor (opcional)")),
            "hacia": id_base(f.get("Hacia ruta")),
            "hacia_texto": texto(f.get("Hacia ruta")),
            "hacia_liq": _liquidacion_en(f.get("Hacia ruta")),
            "activa": g.normalizar(f.get("Activa")) not in ("NO", "N"),
        })
    return reglas


# ---------------------------------------------------------------------------
# Cambios de ruta (sobre las filas de la SABANA, antes de agrupar)
# ---------------------------------------------------------------------------

def aplicar_movimientos(filas, reglas, fechas):
    """Aplica las reglas de CAMBIOS DE RUTA a la ruta que se esta armando.

    Es lo mismo que cambiar a mano el Chofer en la SABANA, pero sin tocar
    el maestro: las filas afectadas se copian con el Chofer de la ruta
    destino y '_movido_de' con la ruta de origen. Las SOLO ESE DÍA van
    antes que las SIEMPRE (y una SIEMPRE no deshace lo que hizo una SOLO
    ESE DÍA). Las reglas se encadenan: si un camion se pasa entero (TODOS)
    a otro, lo que otra regla le mandaba sigue hasta ese otro.
    Devuelve (filas, estados) con un estado (nivel, texto) por regla:
    nivel 'ok', 'info' o 'error'."""
    if not fechas:
        return filas, [("info", "") for _ in reglas]
    ctx = {"filas": filas, "fechas": fechas, "indices": [], "ruta_de": {}, "chofer_de": {},
           "liqs_de": {}, "nuevas": {}, "movido_por": {}, "vaciado_a": {}}
    for i, f in enumerate(filas):
        if f["Fecha Liq"] not in fechas:
            continue
        chofer = texto(f["Chofer"])
        ruta = g.id_camion(chofer)
        ctx["indices"].append(i)
        ctx["ruta_de"][i] = ruta
        ctx["chofer_de"].setdefault(ruta, chofer)
        ctx["liqs_de"].setdefault(ruta, set()).add(str(f["Numero Liq"]))
    ctx["ruta_actual"] = dict(ctx["ruta_de"])
    estados = [None] * len(reglas)
    orden = sorted(enumerate(reglas), key=lambda par: (par[1]["tipo"] != TIPO_DIA, par[0]))
    for posicion, r in orden:
        estados[posicion] = _aplicar_regla(r, ctx)
    nuevas = ctx["nuevas"]
    resultado = [nuevas.get(i, f) for i, f in enumerate(filas)] if nuevas else filas
    return resultado, estados


def _aplicar_regla(r, ctx):
    filas, fechas, indices = ctx["filas"], ctx["fechas"], ctx["indices"]
    ruta_de, ruta_actual, chofer_de = ctx["ruta_de"], ctx["ruta_actual"], ctx["chofer_de"]
    nuevas, movido_por = ctx["nuevas"], ctx["movido_por"]
    nombre_ruta = g.formato_fechas(fechas)
    es_dia = r["tipo"] == TIPO_DIA

    if not r["activa"]:
        return ("info", "Desactivada")
    if r["tipo"] not in (TIPO_DIA, TIPO_SIEMPRE):
        return ("error", f"Tipo debe ser {TIPO_DIA} o {TIPO_SIEMPRE}")
    if not r["codigo"]:
        return ("error", "Falta el código del cliente (o TODOS)")
    if not r["hacia"]:
        return ("error", "Falta 'Hacia ruta'")
    if not ruta_valida(r["hacia_texto"]):
        return ("error", f"No reconozco la ruta '{r['hacia_texto']}' en 'Hacia ruta'; usa el código del "
                         "camión (ej. F150)")
    if r["desde_texto"] and not ruta_valida(r["desde_texto"]):
        return ("error", f"No reconozco la ruta '{r['desde_texto']}' en 'Desde ruta'; usa el código del "
                         "camión (ej. F150)")
    if r["codigo"] == "TODOS" and not r["desde"]:
        return ("error", "Con TODOS hay que indicar 'Desde ruta'")
    if r["desde"] and r["desde"] == r["hacia"]:
        return ("error", "'Desde ruta' y 'Hacia ruta' son la misma")
    if es_dia:
        if not r["fecha"]:
            if r["fecha_texto"]:
                return ("error", f"No entiendo la fecha '{r['fecha_texto']}' (usa dd/mm/aaaa)")
            return ("error", f"Falta la fecha de la ruta (como sale en el mapa, ej. {max(fechas):%d/%m/%Y})")
        if r["fecha"] not in fechas:
            cuando = "ya pasó" if r["fecha"] < min(fechas) else "todavía no llega"
            return ("info", f"No aplica: es para la ruta del {g.formato_fecha(r['fecha'])} ({cuando}); "
                            f"la que se armó es del {nombre_ruta}")

    # Si el destino se paso entero a otro camion (TODOS), se sigue la cadena.
    hacia, vistos = r["hacia"], set()
    while hacia in ctx["vaciado_a"] and hacia not in vistos:
        vistos.add(hacia)
        hacia = ctx["vaciado_a"][hacia]
    nota = f" ({r['hacia']} se pasó entera a {hacia})" if hacia != r["hacia"] else ""
    if hacia not in chofer_de:
        return ("error" if es_dia else "info",
                f"La ruta {hacia} no sale el {nombre_ruta}; no se movió nada")
    cargue_destino = r["hacia_liq"] if hacia == r["hacia"] else ""
    if cargue_destino and cargue_destino not in ctx["liqs_de"].get(hacia, set()):
        return ("error", f"{hacia} no tiene la liquidación {cargue_destino} en la ruta del {nombre_ruta}")

    def coincide(i, rutas):
        f = filas[i]
        if r["codigo"] != "TODOS" and normalizar_codigo(f["CodigoCliente"]) != r["codigo"]:
            return False
        if r["desde"] and rutas[i] != r["desde"]:
            return False
        if r["desde_liq"] and str(f["Numero Liq"]) != r["desde_liq"]:
            return False
        if r["vendedor"] and g.normalizar(f.get("vendedor")) != r["vendedor"]:
            return False
        return True

    def puede_mover(i):
        # Algo que ya movio otra regla solo lo vuelve a mover una regla con
        # 'Desde ruta' explicito (cadena), y nunca una SIEMPRE lo que hizo
        # una SOLO ESE DÍA.
        if i not in movido_por:
            return True
        return bool(r["desde"]) and not (movido_por[i] == TIPO_DIA and not es_dia)

    candidatas = [i for i in indices if coincide(i, ruta_actual) and ruta_actual[i] != hacia]
    libres = [i for i in candidatas if puede_mover(i)]
    if not libres:
        ya_movido = [i for i in indices if coincide(i, ruta_de) and i in movido_por] or \
            [i for i in candidatas if i in movido_por]
        if ya_movido:
            return ("info", "Ya lo movió otra regla (las de SOLO ESE DÍA van primero)")
        if any(coincide(i, ruta_actual) for i in indices):
            return ("info", f"Ya estaba en {hacia}")
        donde = f"en {r['desde_texto'] or r['desde']} " if r["desde"] else ""
        return ("error" if es_dia else "info", f"El cliente no está {donde}en la ruta del {nombre_ruta}")
    origenes = sorted({ruta_actual[i] for i in libres})
    if len(origenes) > 1 and not r["desde"]:
        return ("error", f"El cliente está en varias rutas ({', '.join(origenes)}); indica 'Desde ruta'")

    for i in libres:
        ruta_actual[i] = hacia
        movido_por[i] = r["tipo"]
        if hacia == ruta_de[i]:
            nuevas.pop(i, None)          # volvio a su ruta original
            continue
        copia = dict(nuevas.get(i, filas[i]))
        copia["Chofer"] = chofer_de[hacia]
        copia["_movido_de"] = ruta_de[i]
        if cargue_destino:
            copia["_cargue_destino"] = cargue_destino
        else:
            copia.pop("_cargue_destino", None)
        nuevas[i] = copia
    if r["codigo"] == "TODOS" and not r["desde_liq"] and not r["vendedor"] \
            and not any(ruta_actual[i] == r["desde"] for i in indices):
        ctx["vaciado_a"][r["desde"]] = hacia
    clientes = len({normalizar_codigo(filas[i]["CodigoCliente"]) for i in libres})
    cuantos = f"{clientes} clientes" if clientes != 1 else "1 cliente"
    destino = f"{hacia} (cargue Liq {cargue_destino})" if cargue_destino else hacia
    return ("ok", f"Aplicado a la ruta del {nombre_ruta}: {cuantos} de {', '.join(origenes)} a {destino}{nota}")


# ---------------------------------------------------------------------------
# Ubicaciones (overrides.json): diferencias con la foto y aplicacion
# ---------------------------------------------------------------------------

def _mismo_punto(a, b) -> bool:
    if a is None or b is None:
        return a is None and b is None
    return abs(a[0] - b[0]) < 1e-7 and abs(a[1] - b[1]) < 1e-7


def _como_par(valor):
    if isinstance(valor, dict) and g.es_numero(valor.get("lat")) and g.es_numero(valor.get("lng")):
        return [float(valor["lat"]), float(valor["lng"])]
    return None


REPETIDO = "Código repetido en esta hoja: se usó la fila de arriba y esta se descarta"
FALTA_NUEVA = "Falta la 'Nueva coordenada' (columna amarilla); para quitar una corrección usa Acción = QUITAR"
YA_ESTABA = "Ya estaba así (no hubo que cambiar nada)"
RESUELTOS = ("Aplicada", "Quitada", "Ya estaba")


def es_fila_sin_geo(fila: dict) -> bool:
    return g.normalizar(fila.get("Estado")).startswith(MARCA_SIN_GEO)


def _motivo_de_fila(fila: dict) -> str:
    """El motivo (por que no tiene ubicacion) de una fila SIN GEORREFERENCIA.
    Si ahi se pego una coordenada o un enlace, eso no es el motivo."""
    if not es_fila_sin_geo(fila):
        return ""
    valor = texto(fila.get("Coordenada del sistema"))
    if es_enlace_corto(valor):
        return ""
    try:
        parsear_coordenada(valor, abrir_enlaces=False)
        return ""
    except ValueError:
        return valor


def _sin_georef_de_hoja(editables: dict):
    """Las filas SIN GEORREFERENCIA que ya estaban en la hoja y siguen sin
    coordenada, como lista de clientes (para conservarlas cuando no se pasa
    la lista nueva)."""
    salida = []
    for f in editables.get("ubicaciones", []):
        if es_fila_sin_geo(f) and not texto(f.get("Nueva coordenada")) and not texto(f.get("Acción")):
            salida.append({"codigo": normalizar_codigo(f.get("Código")), "nombre": texto(f.get("Cliente")),
                           "rutas": [r.strip() for r in texto(f.get("Ruta (última vez)")).split(",") if r.strip()],
                           "motivo": texto(f.get("Coordenada del sistema")).strip("()") or "sin coordenadas",
                           "direccion": texto(f.get("Dirección")),
                           "vendedores": [v.strip() for v in texto(f.get("Vendedor")).split(",") if v.strip()]})
    return salida


def _cerca(a, b, metros=1.0) -> bool:
    if a is None or b is None:
        return a is None and b is None
    return distancia_m(a[0], a[1], b[0], b[1]) < metros


def cambios_de_ubicacion(editables: dict):
    """Lo que hay que cambiar en overrides.json segun la hoja UBICACIONES.
    Devuelve (cambios, estados_por_fila). Borrar una fila nunca quita una
    correccion: para quitar se usa Accion = QUITAR. Si algo ya estaba asi
    no se decide aqui (con la foto, que puede ser vieja) sino al aplicar,
    contra lo ultimo de GitHub. No sale a internet (corre en la ventana): los
    enlaces cortos ya se abrieron al leer el Excel; si no se pudo, la fila
    queda con el aviso."""
    foto = editables.get("foto") or {}
    cambios, estados, vistos = [], {}, set()
    for i, f in enumerate(editables.get("ubicaciones", [])):
        codigo = normalizar_codigo(f.get("Código"))
        nueva = texto(f.get("Nueva coordenada"))
        accion = g.normalizar(f.get("Acción"))
        antes = foto.get(codigo)
        if not nueva and not accion:
            # Coordenada pegada por error en 'Corrección vigente' (la columna
            # de al lado): si es distinta a la correccion que habia, es nueva.
            # En una fila SIN GEORREFERENCIA tambien vale en 'Coordenada del
            # sistema' (ahi solo estaba el motivo).
            otras = [texto(f.get("Corrección vigente"))]
            if es_fila_sin_geo(f):
                otras.append(texto(f.get("Coordenada del sistema")))
            for vigente in otras:
                if not (codigo and vigente):
                    continue
                try:
                    par = list(parsear_coordenada(vigente, abrir_enlaces=False))
                except ValueError:
                    par = None
                    if es_enlace_corto(vigente):
                        # Enlace corto que no se pudo abrir: se trata como si
                        # estuviera en 'Nueva coordenada' para que salga su aviso
                        # (y al rearmar el Excel se pasa a esa columna).
                        nueva = vigente
                        break
                if par and not _cerca(antes, par):
                    nueva = vigente
                    break
            if not nueva:
                # Una fila SIN GEORREFERENCIA todavia sin coordenada no es un
                # error: es la lista de lo que falta (la ventana la muestra).
                if codigo and antes is None and not es_fila_sin_geo(f):
                    estados[i] = FALTA_NUEVA
                continue
        if not codigo:
            estados[i] = "Falta el código del cliente"
            continue
        if codigo in vistos:
            estados[i] = REPETIDO
            continue
        if accion and accion != QUITAR:
            estados[i] = f"Acción no válida (solo {QUITAR} o vacío)"
            continue
        if accion == QUITAR:
            vistos.add(codigo)
            cambios.append({"fila": i, "codigo": codigo, "accion": "quitar", "antes": antes})
            continue
        try:
            lat, lng = parsear_coordenada(nueva, abrir_enlaces=False)
        except ValueError as e:
            estados[i] = f"Nueva coordenada: {e}"
            continue
        vistos.add(codigo)
        cambios.append({"fila": i, "codigo": codigo, "accion": "poner", "lat": lat, "lng": lng, "antes": antes})
    return cambios, estados


def aplicar_ubicaciones_en_archivo(ruta: Path, cambios):
    """Aplica los cambios a overrides.json (ya traido de GitHub). Si un
    cliente se movio en el mapa despues de que se escribio la fila, no se
    pisa: queda como conflicto. Devuelve (estados_por_fila, hubo_cambios)."""
    try:
        actual = json.loads(ruta.read_text(encoding="utf-8")) if ruta.exists() else {}
    except ValueError:
        raise g.ErrorRuta("overrides.json del repo de datos está dañado; avisa.")
    ahora = datetime.now()
    estados, cambiado = {}, False
    for c in cambios:
        hoy = _como_par(actual.get(c["codigo"]))
        objetivo = None if c["accion"] == "quitar" else [round(c["lat"], 7), round(c["lng"], 7)]
        if _mismo_punto(hoy, objetivo):
            if _mismo_punto(c["antes"], objetivo):
                estados[c["fila"]] = YA_ESTABA
            else:
                # Ya estaba hecho (por ejemplo, un reintento despues de que
                # GitHub rechazo la primera subida).
                estados[c["fila"]] = (f"Quitada el {ahora:%d/%m %H:%M}" if objetivo is None
                                      else f"Aplicada el {ahora:%d/%m %H:%M}")
            continue
        if not _mismo_punto(hoy, c["antes"]):
            estados[c["fila"]] = ("Conflicto: esta ubicación se cambió en el mapa después de escribir esta fila. "
                                  "Revisa 'Corrección vigente'; si igual quieres la tuya, vuelve a aplicar.")
            continue
        if c["accion"] == "quitar":
            actual.pop(c["codigo"], None)
            estados[c["fila"]] = f"Quitada el {ahora:%d/%m %H:%M}"
        else:
            actual[c["codigo"]] = {"lat": round(c["lat"], 7), "lng": round(c["lng"], 7)}
            estados[c["fila"]] = f"Aplicada el {ahora:%d/%m %H:%M}"
        cambiado = True
    if cambiado:
        # Mismo formato que escribe el mapa (JSON.stringify compacto).
        ruta.write_text(json.dumps(actual, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    return estados, cambiado


# ---------------------------------------------------------------------------
# Armado del Excel
# ---------------------------------------------------------------------------

_INVALIDOS_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")


def _xml_texto(valor) -> str:
    return escape(_INVALIDOS_XML.sub("", str(valor)))


def _celda_xml(ref, valor) -> str:
    if valor is None or valor == "":
        return ""
    if isinstance(valor, bool):
        valor = "SI" if valor else ""
    if isinstance(valor, (int, float)):
        if isinstance(valor, float) and not math.isfinite(valor):
            return ""
        return f'<c r="{ref}"><v>{valor}</v></c>'
    s = str(valor)
    if s.startswith("=HYPERLINK("):
        return f'<c r="{ref}" t="str"><f>{_xml_texto(s[1:])}</f><v>Ver</v></c>'
    return f'<c r="{ref}" t="inlineStr"><is><t xml:space="preserve">{_xml_texto(s)}</t></is></c>'


def _inyectar_filas(ruta: Path, nombre_hoja: str, filas, columnas: int) -> None:
    """Escribe muchas filas en una hoja directo en el XML del .xlsx. Es
    ~10 veces mas rapido que openpyxl para las ~90 mil filas de RUTAS."""
    import zipfile
    letras = [get_column_letter(i + 1) for i in range(columnas)]
    partes = []
    for n, fila in enumerate(filas, start=2):
        celdas = "".join(_celda_xml(f"{letras[j]}{n}", v) for j, v in enumerate(fila))
        partes.append(f'<row r="{n}">{celdas}</row>')
    nuevas = "".join(partes)
    ultima = f"{letras[-1]}{len(filas) + 1}"
    temporal = ruta.with_name(ruta.stem + "_iny.xlsx")
    with zipfile.ZipFile(ruta) as z:
        hojas, _ = g.hojas_del_libro(z)
        parte = next(p for n, p in hojas if n == nombre_hoja)
        with zipfile.ZipFile(temporal, "w", zipfile.ZIP_DEFLATED, compresslevel=5) as salida:
            for info in z.infolist():
                contenido = z.read(info.filename)
                if info.filename == parte:
                    xml = contenido.decode("utf-8")
                    dimension = f'<dimension ref="A1:{ultima}"/>'
                    if "<dimension" in xml:
                        xml = re.sub(r'<dimension ref="[^"]*"\s*/>', dimension, xml, count=1)
                    else:
                        # openpyxl (modo rapido) no la escribe; sin ella, leer luego
                        # las hojas editables obliga a recorrer toda RUTAS.
                        for etiqueta in ("<sheetViews", "<sheetFormatPr", "<cols", "<sheetData"):
                            if etiqueta in xml:
                                xml = xml.replace(etiqueta, dimension + etiqueta, 1)
                                break
                    xml = xml.replace("</sheetData>", nuevas + "</sheetData>", 1)
                    contenido = xml.encode("utf-8", "replace")
                salida.writestr(info, contenido)
    os.replace(temporal, ruta)


def _encabezado(ws, columnas, editables=(), anchos=None):
    if anchos:
        for i, ancho in enumerate(anchos, start=1):
            ws.column_dimensions[get_column_letter(i)].width = ancho
    fila = []
    for c in columnas:
        celda = WriteOnlyCell(ws, value=c)
        celda.font = NEGRITA
        celda.fill = FONDO_EDITABLE if c in editables else FONDO_TITULO
        celda.alignment = Alignment(wrap_text=True, vertical="center")
        fila.append(celda)
    ws.append(fila)


def _ancho_auto(columnas, minimo=9, maximo=40):
    return [max(minimo, min(maximo, len(c) + 3)) for c in columnas]


def _dia(fecha_iso: str) -> str:
    try:
        return g.DIAS[date.fromisoformat(fecha_iso).weekday()]
    except ValueError:
        return ""


def _indice_clientes(rutas: dict):
    """codigo -> datos de la aparicion mas reciente (para UBICACIONES)."""
    indice = {}
    for fecha in sorted(rutas, reverse=True):
        for camion in rutas[fecha].get("camiones", []):
            for c in camion.get("clientes", []):
                cod = normalizar_codigo(c.get("codigo"))
                if cod and cod not in indice:
                    indice[cod] = {"nombre": c.get("nombre", ""), "ruta": camion.get("id", ""), "fecha": fecha,
                                   "lat": c.get("lat"), "lng": c.get("lng")}
                    if c.get("sin_coord_sistema"):
                        # El lat/lng publicado es la correccion, no del sistema.
                        indice[cod].update(lat=None, lng=None, motivo="sin coordenada válida en el sistema "
                                                                      "(NO quites la corrección)")
        for s in rutas[fecha].get("sin_ubicacion", []):
            cod = normalizar_codigo(s.get("codigo"))
            if cod and cod not in indice:
                indice[cod] = {"nombre": s.get("nombre", ""), "ruta": s.get("ruta", ""), "fecha": fecha,
                               "lat": None, "lng": None, "motivo": s.get("motivo", "")}
    return indice


def _filas_revisar(datos: dict):
    vigente = datos["rutas"].get(datos["vigente"]) if datos["vigente"] else None
    if not vigente:
        return []
    overrides = datos["overrides"]
    filas = []

    def agregar(motivo, ruta, c, lat, lng):
        cod = normalizar_codigo(c.get("codigo"))
        filas.append([motivo, ruta, cod, c.get("nombre", ""), c.get("direccion", ""), c.get("vendedor", ""),
                      texto_coord(lat, lng) if lat is not None else "",
                      formula_enlace(lat, lng) if lat is not None else "",
                      "SI" if cod in overrides else ""])

    for s in vigente.get("sin_ubicacion", []):
        agregar(f"No sale en el mapa: {s.get('motivo', 'sin ubicación válida')}", s.get("ruta", ""), s, None, None)

    anteriores = {}
    fechas_previas = sorted((f for f in datos["rutas"] if f < datos["vigente"]), reverse=True)[:DIAS_DETALLE]
    for fecha in fechas_previas:
        for camion in datos["rutas"][fecha].get("camiones", []):
            for c in camion.get("clientes", []):
                if not c.get("sin_coord_sistema"):
                    anteriores.setdefault(normalizar_codigo(c.get("codigo")), (fecha, c.get("lat"), c.get("lng")))

    por_coord = {}
    for camion in vigente.get("camiones", []):
        for c in camion.get("clientes", []):
            cod = normalizar_codigo(c.get("codigo"))
            ov = _como_par(overrides.get(cod))
            lat, lng = ov if ov else (c["lat"], c["lng"])
            por_coord.setdefault((round(lat, 6), round(lng, 6)), set()).add(cod)
    for camion in vigente.get("camiones", []):
        clientes = camion.get("clientes", [])
        efectivas = []
        for c in clientes:
            ov = _como_par(overrides.get(normalizar_codigo(c.get("codigo"))))
            efectivas.append(ov if ov else (c["lat"], c["lng"]))
        if len(clientes) >= 3:
            mlat = statistics.median(p[0] for p in efectivas)
            mlng = statistics.median(p[1] for p in efectivas)
        for c, (lat, lng) in zip(clientes, efectivas):
            cod = normalizar_codigo(c.get("codigo"))
            ruta = camion.get("id", "")
            if len(clientes) >= 3:
                km = distancia_m(lat, lng, mlat, mlng) / 1000
                if km > KM_LEJOS_DE_SU_RUTA:
                    agregar(f"Lejos del resto de su ruta ({km:.0f} km)", ruta, c, lat, lng)
            juntos = len(por_coord[(round(lat, 6), round(lng, 6))])
            if juntos >= 3:
                agregar(f"Misma coordenada que otros {juntos - 1} clientes", ruta, c, lat, lng)
            if c.get("sin_coord_sistema"):
                continue        # su lat/lng es la correccion: no hay coordenada del sistema que comparar
            previo = anteriores.get(cod)
            if previo and g.es_numero(previo[1]) and g.es_numero(previo[2]):
                m = distancia_m(c["lat"], c["lng"], previo[1], previo[2])
                if m > M_CAMBIO_DEL_SISTEMA:
                    agregar(f"La coordenada del sistema cambió {m:,.0f} m desde el "
                            f"{g.formato_fecha(date.fromisoformat(previo[0]))}", ruta, c, lat, lng)
            ov = _como_par(overrides.get(cod))
            if ov and distancia_m(ov[0], ov[1], c["lat"], c["lng"]) < M_SISTEMA_YA_CORRIGIO:
                agregar("El sistema ya trae esta ubicación: puedes QUITAR la corrección en UBICACIONES",
                        ruta, c, lat, lng)
    return filas


def _leeme(datos: dict):
    vigente = datos["vigente"]
    lineas = [
        "EXCEL DE CONTROL DE RUTAS",
        f"Armado el {datetime.now():%d/%m/%Y %H:%M}. Ruta vigente en el mapa: "
        f"{g.formato_fecha(date.fromisoformat(vigente)) if vigente else '—'}.",
        "Se arma solo con lo YA PUBLICADO en el mapa (no con el maestro). Se vuelve a armar después de cada "
        "publicación y con el botón 'Excel de control' de la ventana Actualizar Ruta.",
        "",
        "PARA REVISAR",
        f"• RUTAS: cada cliente de cada ruta de los últimos {DIAS_DETALLE} días. Usa los filtros del encabezado "
        "(fecha, ruta, vendedor, cliente...). 'Ver' abre el punto en Google Maps.",
        "• RESUMEN: clientes, peso y total por fecha y ruta, de todo el histórico.",
        "• REVISAR UBICACIÓN: clientes de la ruta vigente con ubicación sospechosa (lejos de su ruta, varios "
        "clientes en el mismo punto, coordenada del sistema que cambió, o que no salen en el mapa).",
        "",
        "PARA CAMBIAR (columnas en amarillo). Guarda y cierra el Excel antes de publicar.",
        "• UBICACIONES: todas las correcciones de ubicación (las mismas que haces arrastrando el punto en el "
        "mapa). Valen para siempre y para todas las rutas de ese cliente.",
        "   - Corregir o agregar: escribe el Código y en 'Nueva coordenada' pega la coordenada (en Google Maps, "
        "clic derecho sobre el punto exacto → clic en los números para copiarlos, ej. -2.900123, -79.004567).",
        "   - Quitar una corrección (volver a la ubicación del sistema): en 'Acción' elige QUITAR.",
        f"   - Filas en ROJO arriba ({MARCA_SIN_GEO}): clientes de la ruta que vas a publicar que no tienen "
        "ubicación (sin coordenadas o en 0,0). TODOS deben tenerla: pega su coordenada en 'Nueva coordenada', "
        "guarda y cierra; la ventana se actualiza sola y deja publicar cuando no falta ninguno.",
        "   - Se aplican al publicar o con el botón 'Aplicar ubicaciones' de la ventana. Un cliente que ya "
        "salía en el mapa se ve corregido al volver a abrir el mapa; uno que NO salía (sin ubicación) aparece "
        "recién al volver a 'Generar y publicar' (y en My Maps, al volver a subir).",
        "• CAMBIOS DE RUTA: pasar un cliente a la ruta de otro chofer.",
        f"   - Tipo: {TIPO_DIA} (solo para la ruta de esa fecha) o {TIPO_SIEMPRE} (todas las rutas que se "
        "publiquen mientras esté Activa = SI).",
        "   - Fecha de la ruta: obligatoria en SOLO ESE DÍA; es la fecha con la que sale la ruta en el mapa y "
        "en la ventana (fecha de liquidación).",
        "   - Código: el del cliente, o TODOS para pasar una ruta entera (por ejemplo, un camión dañado).",
        "   - Desde ruta: obligatoria si el cliente está en más de una ruta ese día (la columna 'Está hoy en' "
        "ayuda). Vendedor: opcional, para mover solo el pedido de ese vendedor.",
        "   - Desde/Hacia: el código del camión (F150). Si el chofer tiene 2 cargues separados, se puede "
        "indicar uno: 'F325 CARGUE 2 - Liq 239'.",
        "   - Se aplican al generar la ruta en la ventana (dale 'Volver a leer' después de guardar este Excel). "
        "El cliente sale en el mapa y en el MAPA Excel del nuevo chofer, con su número de liquidación original.",
        "   - Para quitar un cambio: borra la fila entera (clic derecho en el número de fila → Eliminar) "
        "o pon Activa = NO.",
        "• La columna Estado la llena el programa: dice si se aplicó o por qué no.",
        "",
        "Este archivo tiene datos de clientes: déjalo en esta PC (no lo subas ni lo compartas).",
    ]
    return lineas


def construir_control(log=None, editables=None, estados_cambios=None, estados_ubic=None,
                      ruta: Path = None, sin_georef=None) -> Path:
    """Arma (o vuelve a armar) el Excel de control. Conserva lo que el
    usuario escribio en las hojas editables y pone el Estado de cada fila.
    sin_georef: clientes de la ruta que se va a publicar sin georreferencia
    (generar_rutas.clientes_sin_georef); van arriba de UBICACIONES para que
    solo haya que pegar su coordenada."""
    ruta = ruta or ARCHIVO_CONTROL

    def avisar(m):
        if log:
            log(m)

    g.revisar_cerrados([ruta])
    # Nunca escribir encima filas mas nuevas que las que se tienen: si el
    # Excel se guardo despues de leerlo, se vuelve a leer (los Estados de
    # esta vuelta ya no corresponden a esas filas).
    if editables is None or editables.get("mtime") != mtime_control(ruta):
        if editables is not None:
            avisar("El Excel de control se cambió mientras tanto: se conservan tus filas de ahora.")
            estados_cambios = estados_ubic = None
        editables = leer_editables(ruta)
    datos = datos_del_repo(log)
    rutas, overrides, vigente = datos["rutas"], datos["overrides"], datos["vigente"]
    fechas = sorted(rutas, reverse=True)
    detalle = set(fechas[:DIAS_DETALLE])
    indice = _indice_clientes(rutas)

    # RUTAS y RESUMEN
    filas_rutas, filas_resumen = [], []
    for fecha in fechas:
        estado = "VIGENTE" if fecha == vigente else "HISTÓRICO"
        for camion in rutas[fecha].get("camiones", []):
            clientes = camion.get("clientes", [])
            corregidos = movidos = 0
            for c in clientes:
                cod = normalizar_codigo(c.get("codigo"))
                ov = _como_par(overrides.get(cod))
                corregidos += bool(ov)
                movidos += bool(c.get("movido_de"))
                if fecha in detalle:
                    lat, lng = ov if ov else (c.get("lat"), c.get("lng"))
                    # Si el sistema no traia coordenada valida, el lat/lng
                    # publicado ES la correccion: no se muestra como del sistema.
                    sistema = (None, None) if c.get("sin_coord_sistema") else (c.get("lat"), c.get("lng"))
                    filas_rutas.append([
                        fecha, _dia(fecha), estado, camion.get("id", ""), camion.get("chofer", ""),
                        c.get("vendedor", ""), cod, c.get("nombre", ""), c.get("direccion", ""), c.get("canal", ""),
                        round(c.get("peso") or 0, 2), round(c.get("venta") or 0, 2), round(c.get("total") or 0, 2),
                        sistema[0], sistema[1], lat, lng, "SI" if ov else "", c.get("movido_de", ""),
                        formula_enlace(lat, lng) if g.es_numero(lat) and g.es_numero(lng) else ""])
            filas_resumen.append([
                fecha, _dia(fecha), estado, camion.get("id", ""), camion.get("chofer", ""), len(clientes),
                len({c.get("vendedor") for c in clientes}),
                round(sum(c.get("peso") or 0 for c in clientes), 2),
                round(sum(c.get("venta") or 0 for c in clientes), 2),
                round(sum(c.get("total") or 0 for c in clientes), 2), movidos, corregidos])

    # UBICACIONES: una fila por correccion vigente + las filas pendientes
    estados_ubic = estados_ubic or {}
    cambios_pendientes, estados_previos = cambios_de_ubicacion(editables)
    cambio_de_fila = {c["fila"]: c for c in cambios_pendientes}
    notas = {}
    for f in editables.get("ubicaciones", []):
        cod = normalizar_codigo(f.get("Código"))
        if cod and texto(f.get("Nota")) and cod not in notas:
            notas[cod] = f.get("Nota")
    # Datos de clientes que aun no salen en lo publicado (un cliente nuevo sin
    # ubicacion nunca llega a publicarse): de su fila SIN GEORREFERENCIA o de
    # la lista de la ventana. Asi su fila conserva nombre, ruta, direccion y
    # vendedor, y no se marca como "codigo mal escrito".
    extra = {}
    for c in sin_georef or []:
        cod = normalizar_codigo(c.get("codigo"))
        if cod:
            extra.setdefault(cod, {"nombre": c.get("nombre", ""), "ruta": ", ".join(c.get("rutas", [])),
                                   "motivo": f"({c.get('motivo', 'sin coordenadas')})",
                                   "direccion": c.get("direccion", ""), "vendedor": ", ".join(c.get("vendedores", []))})
    for f in editables.get("ubicaciones", []):
        cod = normalizar_codigo(f.get("Código"))
        if cod and (es_fila_sin_geo(f) or texto(f.get("Dirección")) or texto(f.get("Vendedor"))):
            extra.setdefault(cod, {"nombre": texto(f.get("Cliente")), "ruta": texto(f.get("Ruta (última vez)")),
                                   "motivo": _motivo_de_fila(f),
                                   "direccion": texto(f.get("Dirección")), "vendedor": texto(f.get("Vendedor"))})
    for cod, info in extra.items():
        if cod not in indice and info["nombre"]:
            indice[cod] = {"nombre": info["nombre"], "ruta": info["ruta"], "lat": None, "lng": None,
                           "motivo": info["motivo"] or "(sin coordenadas en el sistema)"}
    pendientes = []
    for i, f in enumerate(editables.get("ubicaciones", [])):
        cod = normalizar_codigo(f.get("Código"))
        estado = estados_ubic.get(i)
        previo = estados_previos.get(i)
        if (estado and estado.startswith(RESUELTOS)) or previo == REPETIDO:
            continue
        nueva = texto(f.get("Nueva coordenada"))
        if not nueva and cambio_de_fila.get(i, {}).get("accion") == "poner":
            # Pegada en otra columna (la de al lado): se pasa a 'Nueva coordenada'.
            nueva = texto_coord(cambio_de_fila[i]["lat"], cambio_de_fila[i]["lng"])
        if not nueva and i not in cambio_de_fila:
            # Enlace corto pegado en la columna de al lado que aun no se pudo
            # abrir: se pasa a 'Nueva coordenada' (si no, se perderia).
            otras = ["Corrección vigente"] + (["Coordenada del sistema"] if es_fila_sin_geo(f) else [])
            nueva = next((texto(f.get(c)) for c in otras if es_enlace_corto(f.get(c))), "")
        accion = texto(f.get("Acción"))
        if not (nueva or accion or previo == FALTA_NUEVA):
            continue        # (incluye las SIN GEORREFERENCIA aun vacias: se regeneran abajo)
        estado_hoja = "" if es_fila_sin_geo(f) else texto(f.get("Estado"))
        pendientes.append((cod, nueva or None, accion or None,
                           estado or previo or estado_hoja
                           or "Pendiente: se aplica al publicar o con 'Aplicar ubicaciones'",
                           f.get("Nota") or notas.get(cod)))
    filas_ubic = []
    codigos_pendientes = {p[0] for p in pendientes if p[0]}
    # Arriba de todo, los clientes sin georreferencia que no tengan ya una
    # fila del usuario ni una correccion valida.
    filas_sin_geo = []
    vistos_sin_geo = set()
    if sin_georef is None:
        # Sin lista nueva (por ejemplo, armado desde otro lado): se conservan
        # las que ya estaban y siguen sin coordenada.
        sin_georef = _sin_georef_de_hoja(editables)
    for c in sin_georef:
        cod = normalizar_codigo(c.get("codigo"))
        if not cod or cod in codigos_pendientes or cod in vistos_sin_geo or g.correccion_valida(overrides, cod):
            continue
        vistos_sin_geo.add(cod)
        filas_sin_geo.append(_fila_sin_georef(cod, c))
    for cod, nueva, accion, estado, nota in pendientes:
        filas_ubic.append(_fila_ubicacion(cod, indice, overrides, nueva, accion, estado, nota, extra.get(cod)))
    aplicadas = {}
    for i, e in estados_ubic.items():
        if e.startswith("Aplicada") and i < len(editables.get("ubicaciones", [])):
            cod = normalizar_codigo(editables["ubicaciones"][i].get("Código"))
            aplicadas[cod] = e + ("" if cod in indice else " — OJO: este código no sale en las rutas "
                                                           "publicadas; revisa que esté bien escrito")
    for cod in sorted(overrides, key=lambda k: (indice.get(k, {}).get("ruta", ""), k)):
        if cod in codigos_pendientes:
            continue
        filas_ubic.append(_fila_ubicacion(cod, indice, overrides, None, None, aplicadas.get(cod, ""),
                                          notas.get(cod), extra.get(cod)))
    # Foto con que se compara cada fila al aplicar: la correccion de hoy, salvo
    # en las filas pendientes, que guardan la que habia cuando se escribieron
    # (asi, si el mapa la cambia mientras tanto, se detecta el conflicto). Tras
    # un conflicto se toma la de hoy, para que "volver a aplicar" funcione.
    foto_vieja = editables.get("foto") or {}
    foto_nueva = {cod: par for cod, par in ((k, _como_par(v)) for k, v in overrides.items()) if par}
    for cod, _, _, estado, _ in pendientes:
        if cod and not str(estado).startswith("Conflicto"):
            if cod in foto_vieja:
                foto_nueva[cod] = foto_vieja[cod]
            else:
                foto_nueva.pop(cod, None)

    # CAMBIOS DE RUTA: las filas del usuario, en su orden, con su Estado
    estados_cambios = estados_cambios or {}
    hoy_en = {}
    if vigente and vigente in rutas:
        for camion in rutas[vigente].get("camiones", []):
            for c in camion.get("clientes", []):
                hoy_en.setdefault(normalizar_codigo(c.get("codigo")), set()).add(camion.get("id", ""))
    filas_cambios = []
    for i, f in enumerate(editables.get("cambios", [])):
        cod = normalizar_codigo(f.get("Código"))
        estado = estados_cambios.get(i) or texto(f.get("Estado"))
        filas_cambios.append([f.get("Tipo"), a_fecha(f.get("Fecha de la ruta")) or f.get("Fecha de la ruta"),
                              f.get("Código"), f.get("Desde ruta"), f.get("Vendedor (opcional)"), f.get("Hacia ruta"),
                              f.get("Activa") or "SI", f.get("Nota"), estado,
                              ", ".join(sorted(hoy_en.get(cod, []))) if cod and cod != "TODOS" else ""])

    # Listas para los desplegables
    rutas_recientes = set()
    for fecha in fechas[:7]:
        for camion in rutas[fecha].get("camiones", []):
            rutas_recientes.add(id_base(camion.get("id", "")))
    rutas_recientes = sorted(r for r in rutas_recientes if r)

    temporal = ruta.with_name("_control_tmp.xlsx")
    wb = openpyxl.Workbook(write_only=True)

    ws = wb.create_sheet(HOJA_LEEME)
    ws.column_dimensions["A"].width = 150
    for i, linea in enumerate(_leeme(datos)):
        celda = WriteOnlyCell(ws, value=linea)
        if i == 0 or linea in ("PARA REVISAR",) or linea.startswith("PARA CAMBIAR"):
            celda.font = NEGRITA
        ws.append([celda])

    ws = wb.create_sheet(HOJA_RUTAS)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS_RUTAS))}{max(2, len(filas_rutas) + 1)}"
    _encabezado(ws, COLS_RUTAS, anchos=ANCHOS_RUTAS)

    ws = wb.create_sheet(HOJA_RESUMEN)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS_RESUMEN))}{max(2, len(filas_resumen) + 1)}"
    _encabezado(ws, COLS_RESUMEN, anchos=_ancho_auto(COLS_RESUMEN))
    for fila in filas_resumen:
        ws.append(fila)

    ws = wb.create_sheet(HOJA_REVISAR)
    filas_rev = _filas_revisar(datos)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS_REVISAR))}{max(2, len(filas_rev) + 1)}"
    _encabezado(ws, COLS_REVISAR, anchos=[60, 22, 11, 34, 40, 10, 24, 9, 10])
    for fila in filas_rev:
        ws.append(fila)

    ws = wb.create_sheet(HOJA_UBIC)
    ws.freeze_panes = "B2"
    _encabezado(ws, COLS_UBIC, EDITABLES_UBIC, ANCHOS_UBIC)
    dv = DataValidation(type="list", formula1=f"{HOJA_LISTAS}!$C$1:$C$1", allow_blank=True,
                        showErrorMessage=True, errorTitle="Acción", error="Deja vacío o elige QUITAR")
    dv.add(f"G2:G{FILAS_EDITABLES}")
    ws.data_validations.append(dv)
    for fila in filas_sin_geo:
        celdas = []
        for i, valor in enumerate(fila):
            celda = WriteOnlyCell(ws, value=valor)
            if i in (0, 1, 2, 7):
                celda.fill = FONDO_SIN_GEO
            elif i == 5:
                celda.fill = FONDO_EDITABLE
            celdas.append(celda)
        ws.append(celdas)
    for fila in filas_ubic:
        ws.append(fila)

    ws = wb.create_sheet(HOJA_CAMBIOS)
    ws.freeze_panes = "A2"
    _encabezado(ws, COLS_CAMBIOS, EDITABLES_CAMBIOS, [15, 14, 11, 14, 12, 14, 8, 30, 70, 22])
    for formula, rango, titulo in ((f"{HOJA_LISTAS}!$B$1:$B$2", f"A2:A{FILAS_EDITABLES}", "Tipo"),
                                   (f"{HOJA_LISTAS}!$D$1:$D$2", f"G2:G{FILAS_EDITABLES}", "Activa")):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True,
                            errorTitle=titulo, error="Elige un valor de la lista")
        dv.add(rango)
        ws.data_validations.append(dv)
    if rutas_recientes:
        for col in ("D", "F"):
            dv = DataValidation(type="list", formula1=f"{HOJA_LISTAS}!$A$1:$A${len(rutas_recientes)}",
                                allow_blank=True, showErrorMessage=False)
            dv.add(f"{col}2:{col}{FILAS_EDITABLES}")
            ws.data_validations.append(dv)
    for fila in filas_cambios:
        celdas = list(fila)
        if isinstance(celdas[1], (date, datetime)):
            c = WriteOnlyCell(ws, value=celdas[1])
            c.number_format = "dd/mm/yyyy"
            celdas[1] = c
        ws.append(celdas)

    ws = wb.create_sheet(HOJA_LISTAS)
    ws.sheet_state = "hidden"
    extras = [[TIPO_DIA, QUITAR, "SI"], [TIPO_SIEMPRE, None, "NO"]]
    for i in range(max(len(rutas_recientes), 2)):
        b, c, d = extras[i] if i < 2 else (None, None, None)
        ws.append([rutas_recientes[i] if i < len(rutas_recientes) else None, b, c, d])

    ws = wb.create_sheet(HOJA_FOTO)
    ws.sheet_state = "hidden"
    ws.append(["codigo", "lat", "lng"])
    for cod, par in sorted(foto_nueva.items()):
        ws.append([cod, par[0], par[1]])

    if filas_sin_geo:
        # Que abra directo en UBICACIONES, donde estan las filas por llenar.
        wb.active = wb.sheetnames.index(HOJA_UBIC)
    wb.save(temporal)
    _inyectar_filas(temporal, HOJA_RUTAS, filas_rutas, len(COLS_RUTAS))
    import time as reloj
    for intento in range(4):
        try:
            os.replace(temporal, ruta)
            break
        except PermissionError:
            # Puede ser un instante (la ventana leyendolo justo ahora); si
            # sigue, es que esta abierto en Excel.
            if intento < 3:
                reloj.sleep(0.4)
                continue
            try:
                temporal.unlink()
            except OSError:
                pass
            raise g.ErrorRuta("Cierra el Excel de control (CONTROL DE RUTAS.xlsx) para que se pueda actualizar.")
    guardar_respaldo(editables)
    avisar(f"Excel de control actualizado: {len(fechas)} días publicados ({len(detalle)} con detalle), "
           f"{len(overrides)} ubicaciones corregidas, {len(filas_rev)} para revisar"
           + (f", {len(filas_sin_geo)} clientes SIN georreferencia arriba en UBICACIONES." if filas_sin_geo else "."))
    return ruta


def _fila_ubicacion(cod, indice, overrides, nueva, accion, estado, nota, extra=None):
    info = indice.get(cod, {})
    lat_s, lng_s = info.get("lat"), info.get("lng")
    sistema = texto_coord(lat_s, lng_s) if g.es_numero(lat_s) and g.es_numero(lng_s) else \
        (info.get("motivo") or ("" if info else "(no sale en las rutas publicadas)"))
    ov = _como_par(overrides.get(cod))
    vigente = texto_coord(*ov) if ov else ""
    distancia = round(distancia_m(ov[0], ov[1], lat_s, lng_s)) if ov and g.es_numero(lat_s) and \
        g.es_numero(lng_s) else None
    return [cod, info.get("nombre", ""), info.get("ruta", ""), sistema, vigente, nueva, accion, estado,
            formula_enlace(lat_s, lng_s) if g.es_numero(lat_s) and g.es_numero(lng_s) else None,
            formula_enlace(*ov) if ov else None, distancia, nota,
            (extra or {}).get("direccion") or None, (extra or {}).get("vendedor") or None]


def _fila_sin_georef(cod, c):
    """Fila de UBICACIONES para un cliente sin georreferencia (ver ESTADO_SIN_GEO).
    En 'Ver sistema' va un enlace que busca el local en Google Maps."""
    from urllib.parse import quote_plus
    consulta = ", ".join(x for x in (c.get("nombre", ""), c.get("direccion", ""), "Ecuador") if x)
    prefijo = "https://www.google.com/maps/search/?api=1&query="
    # Excel no acepta textos de mas de 255 caracteres en una formula (sale
    # #¡VALOR!): se acorta la busqueda hasta que la URL ya codificada quepa.
    while consulta and len(prefijo + quote_plus(consulta, safe=",")) > 250:
        consulta = consulta[:-5]
    buscar = f'=HYPERLINK("{prefijo}{quote_plus(consulta, safe=",")}","Buscar")'
    return [cod, c.get("nombre", ""), ", ".join(c.get("rutas", [])), f"({c.get('motivo', 'sin coordenadas')})",
            "", None, None, ESTADO_SIN_GEO, buscar, None, None, None, c.get("direccion", ""),
            ", ".join(c.get("vendedores", []))]
