"""
Convierte los archivos "MAPA F1xx.xlsx" de una carpeta de ruta diaria
(las que ya se generan hoy para subir a Google My Maps) en un unico
archivo data/latest.json liviano para el visor web del mapa.

Uso:
    python convertir.py "C:\\RUTAS DIARIAS\\15MARTES 15"

No modifica los archivos originales, solo los lee.
"""
import json
import re
import sys
from datetime import datetime, date
from pathlib import Path

import openpyxl

COLUMNAS_ESPERADAS = [
    "Fecha Liq", "Chofer", "Numero Liq", "vendedor", "CodigoCliente",
    "Razon", "Direccion", "canal", "longitud_x", "latitud_y",
    "Suma de Peso", "Suma de VentaNeta", "Suma de Total",
]

NOMBRE_MASTER = "RUTAS DIARIAS CAMIONES - MY MAPS.xlsx"


def id_camion_desde_archivo(nombre_archivo: str) -> str:
    nombre = nombre_archivo.replace(".xlsx", "")
    nombre = re.sub(r"^MAPA\s+", "", nombre, flags=re.IGNORECASE)
    return nombre.strip()


def leer_archivo_camion(ruta: Path):
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    filas = list(ws.iter_rows(values_only=True))
    wb.close()

    if not filas:
        return None

    encabezado = [str(c).strip() if c else "" for c in filas[0]]
    idx = {nombre: encabezado.index(nombre) for nombre in COLUMNAS_ESPERADAS if nombre in encabezado}

    faltantes = [c for c in COLUMNAS_ESPERADAS if c not in idx]
    if faltantes:
        print(f"  ! {ruta.name}: faltan columnas {faltantes}, se omite")
        return None

    clientes = []
    chofer = ""
    fecha_liq = None

    for fila in filas[1:]:
        if fila is None or all(v is None for v in fila):
            continue

        lat = fila[idx["latitud_y"]]
        lng = fila[idx["longitud_x"]]
        if lat is None or lng is None:
            continue

        chofer = fila[idx["Chofer"]] or chofer
        if fecha_liq is None and fila[idx["Fecha Liq"]]:
            fecha_liq = fila[idx["Fecha Liq"]]

        clientes.append({
            "codigo": fila[idx["CodigoCliente"]],
            "nombre": (fila[idx["Razon"]] or "").strip(),
            "direccion": (fila[idx["Direccion"]] or "").strip(),
            "canal": (fila[idx["canal"]] or "").strip(),
            "lat": float(lat),
            "lng": float(lng),
            "peso": fila[idx["Suma de Peso"]],
            "venta": fila[idx["Suma de VentaNeta"]],
            "total": fila[idx["Suma de Total"]],
        })

    if not clientes:
        return None

    return {
        "id": id_camion_desde_archivo(ruta.name),
        "chofer": chofer,
        "fecha_liq": fecha_liq,
        "clientes": clientes,
    }


def convertir_carpeta(carpeta: Path) -> dict:
    archivos = sorted(
        f for f in carpeta.glob("MAPA *.xlsx")
        if f.name != NOMBRE_MASTER
    )

    if not archivos:
        raise SystemExit(f"No se encontraron archivos 'MAPA *.xlsx' en {carpeta}")

    camiones = []
    fechas_vistas = []

    for archivo in archivos:
        print(f"Leyendo {archivo.name} ...")
        resultado = leer_archivo_camion(archivo)
        if resultado is None:
            continue

        fecha_liq = resultado.pop("fecha_liq")
        if isinstance(fecha_liq, (datetime, date)):
            fechas_vistas.append(fecha_liq)

        camiones.append(resultado)
        print(f"  -> {resultado['id']}: {len(resultado['clientes'])} clientes")

    if not camiones:
        raise SystemExit("Ningun archivo tenia datos utilizables (revisa columnas o coordenadas vacias)")

    fecha_ruta = max(fechas_vistas).strftime("%Y-%m-%d") if fechas_vistas else carpeta.name

    camiones.sort(key=lambda c: c["id"])

    return {
        "fecha": fecha_ruta,
        "generado": datetime.now().isoformat(timespec="seconds"),
        "camiones": camiones,
    }


def main():
    if len(sys.argv) != 2:
        print(__doc__)
        raise SystemExit(1)

    carpeta = Path(sys.argv[1])
    if not carpeta.is_dir():
        raise SystemExit(f"No existe la carpeta: {carpeta}")

    datos = convertir_carpeta(carpeta)

    salida = Path(__file__).parent / "data" / "latest.json"
    salida.parent.mkdir(exist_ok=True)
    salida.write_text(json.dumps(datos, ensure_ascii=False, indent=None), encoding="utf-8")

    total_clientes = sum(len(c["clientes"]) for c in datos["camiones"])
    print()
    print(f"Listo: {salida}")
    print(f"Fecha de ruta: {datos['fecha']}")
    print(f"Camiones: {len(datos['camiones'])} | Clientes totales: {total_clientes}")


if __name__ == "__main__":
    main()
